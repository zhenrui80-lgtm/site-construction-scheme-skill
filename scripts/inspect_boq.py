#!/usr/bin/env python3
import argparse
import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return value


def sheet_payload(sheet, max_rows, cached_sheet=None):
    rows = []
    numeric_cells = 0
    text_cells = 0
    formula_cells = 0
    formulas_without_cache = []
    for row_index, row in enumerate(sheet.iter_rows(), start=1):
        if max_rows and row_index > max_rows:
            break
        values = [clean(cell.value) for cell in row]
        cached_values = [
            clean(cached_sheet.cell(row=row_index, column=cell.column).value) if cached_sheet is not None else None
            for cell in row
        ]
        while values and values[-1] is None and cached_values[-1] is None:
            values.pop()
            cached_values.pop()
        if not any(value is not None for value in values) and not any(value is not None for value in cached_values):
            continue
        for column_index, value in enumerate(values, start=1):
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numeric_cells += 1
            elif value is not None:
                text_cells += 1
            if isinstance(value, str) and value.startswith("="):
                formula_cells += 1
                cached_value = cached_values[column_index - 1] if column_index <= len(cached_values) else None
                if cached_value is None:
                    formulas_without_cache.append(f"{sheet.cell(row=row_index, column=column_index).coordinate}")
        rows.append({"row": row_index, "values": values, "cached_values": cached_values})
    return {
        "title": sheet.title,
        "state": sheet.sheet_state,
        "max_row": sheet.max_row,
        "max_column": sheet.max_column,
        "nonempty_rows": len(rows),
        "numeric_cells": numeric_cells,
        "text_cells": text_cells,
        "formula_cells": formula_cells,
        "formulas_without_cache": formulas_without_cache,
        "hidden_rows": [index for index, dimension in sheet.row_dimensions.items() if dimension.hidden],
        "hidden_columns": [key for key, dimension in sheet.column_dimensions.items() if dimension.hidden],
        "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
        "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
        "rows": rows,
    }


def main():
    parser = argparse.ArgumentParser(description="Extract a BOQ workbook into reviewable JSON.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--max-rows", type=int, default=0, help="0 means all rows")
    args = parser.parse_args()

    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    formulas = load_workbook(workbook_path, data_only=False, read_only=False)
    cached = load_workbook(workbook_path, data_only=True, read_only=False)
    sheets = []
    for formula_sheet in formulas.worksheets:
        cached_sheet = cached[formula_sheet.title]
        payload = sheet_payload(formula_sheet, args.max_rows, cached_sheet)
        payload["cached_nonempty_rows"] = sum(
            1 for row in payload["rows"] if any(value is not None for value in row["cached_values"])
        )
        sheets.append(payload)

    result = {
        "source": str(workbook_path),
        "sheet_count": len(sheets),
        "defined_names": sorted(str(name) for name in formulas.defined_names),
        "sheets": sheets,
    }
    output = args.output or workbook_path.with_suffix(".boq-analysis.json")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "sheets": len(sheets), "rows": sum(s["nonempty_rows"] for s in sheets)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
