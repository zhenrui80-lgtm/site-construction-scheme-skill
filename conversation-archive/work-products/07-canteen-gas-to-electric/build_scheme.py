from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "template-reference.docx"
OUTPUT_DIR = BASE / "output"
OUTPUT = OUTPUT_DIR / "质检中心食堂气改电施工方案.docx"
ASSETS = BASE / "assets"
NORMS = BASE / "norms"
SHEET_RENDERS = BASE / "spreadsheet" / "renders"

PROJECT = "（实事工程）质检中心食堂气改电"
COMPANY = "武汉兴达建筑安装工程有限责任公司"
START = "2026年8月10日"
FINISH = "2026年12月31日"
DURATION = "144日历天"
COMPILE_DATE = "2026年7月31日"

BODY_FONT = "仿宋_GB2312"
TITLE_FONT = "黑体"
TABLE_FONT = "宋体"
ACCENT = "1F4E78"
LIGHT_HEADER = "D9EAF7"
LIGHT_GRAY = "F2F2F2"


def set_run_font(run, name=BODY_FONT, size=12, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, outer_color="7F7F7F", inner_color="B7B7B7"):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8" if edge in ("top", "left", "bottom", "right") else "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), outer_color if edge in ("top", "left", "bottom", "right") else inner_color)
        borders.append(border)
    tbl_pr.append(borders)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, bold=False, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    set_run_font(r, TABLE_FONT, size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, TABLE_FONT, 9)


def configure_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
    else:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.175)
        section.right_margin = Cm(3.175)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.2)


def clear_header_footer(section, header_text=""):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for part in (section.header, section.footer):
        for p in part.paragraphs:
            p._element.getparent().remove(p._element)
        part.add_paragraph()
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header_text:
        r = hp.add_run(header_text)
        set_run_font(r, TABLE_FONT, 8.5, False, (100, 100, 100))
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("第 ")
    set_run_font(r, TABLE_FONT, 9)
    add_page_field(fp)
    r = fp.add_run(" 页")
    set_run_font(r, TABLE_FONT, 9)


def add_paragraph(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.keep_together = keep
    r = p.add_run(text)
    set_run_font(r, BODY_FONT, size, bold)
    return p


def add_paragraphs(doc, items):
    for item in items:
        add_paragraph(doc, item)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, TITLE_FONT, 16 if level == 1 else 14 if level == 2 else 12, True)
    if level == 1:
        p_pr = p._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), ACCENT)
        p_bdr.append(bottom)
        p_pr.append(p_bdr)
    return p


def add_bullet(doc, text, number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    prefix = f"{number}、" if number is not None else "（1）"
    r = p.add_run(prefix + text)
    set_run_font(r, BODY_FONT, 12)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2, header_fill=LIGHT_HEADER):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, text in enumerate(headers):
        set_cell_text(header.cells[i], text, True, font_size)
        set_cell_shading(header.cells[i], header_fill)
    for values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(row.cells[i], value, False, font_size, align)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_picture(doc, path, caption, width_cm=14.5, source=None, page_break_before=False):
    if page_break_before:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run()
    r.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = bool(source)
    cap.paragraph_format.space_after = Pt(0)
    cr = cap.add_run(caption)
    set_run_font(cr, TABLE_FONT, 9.5, True)
    if source:
        src = doc.add_paragraph()
        src.alignment = WD_ALIGN_PARAGRAPH.CENTER
        src.paragraph_format.space_after = Pt(2)
        sr = src.add_run("查新来源：" + source)
        set_run_font(sr, TABLE_FONT, 8, False, (100, 100, 100))
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, "D6B656", "D6B656")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF2CC")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run(title + "：")
    set_run_font(r, TITLE_FONT, 11, True)
    r = p.add_run(text)
    set_run_font(r, BODY_FONT, 11)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def remove_existing_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def copy_template_footer_field(doc):
    # The source template remains the formatting authority. New content receives
    # an explicit PAGE field because all original body sections are replaced.
    for section in doc.sections:
        clear_header_footer(section, "")


def add_cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(PROJECT)
    set_run_font(r, TITLE_FONT, 18, True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run("施\n工\n方\n案")
    set_run_font(r, "方正仿宋简体", 26, True)

    for _ in range(2):
        doc.add_paragraph()
    info = doc.add_table(rows=5, cols=2)
    set_table_borders(info)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("工程名称", PROJECT),
        ("建设地点", "武汉钢铁有限公司质检中心食堂（以现场交底为准）"),
        ("施工单位", COMPANY),
        ("计划工期", f"{START}开工，{FINISH}完工（{DURATION}）"),
        ("编制日期", COMPILE_DATE),
    ]
    for row, (name, value) in zip(info.rows, values):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.9)
        set_cell_text(row.cells[0], name, True, 11)
        set_cell_text(row.cells[1], value, False, 11, WD_ALIGN_PARAGRAPH.LEFT)
        row.cells[0].width = Cm(3.2)
        row.cells[1].width = Cm(10.3)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(COMPANY)
    set_run_font(r, TITLE_FONT, 14, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("（盖章）")
    set_run_font(r, BODY_FONT, 12)


def add_approval_page(doc):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("施工方案审批表")
    set_run_font(r, TITLE_FONT, 18, True)
    doc.add_paragraph()

    table = doc.add_table(rows=8, cols=4)
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.3, 4.4, 2.3, 4.4]
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            set_cell_margins(cell, 100, 100, 100, 100)
    fields = [
        ("工程名称", PROJECT, "施工单位", COMPANY),
        ("编制人", "", "编制日期", COMPILE_DATE),
        ("项目负责人", "", "技术负责人", ""),
    ]
    for ridx, row_values in enumerate(fields):
        table.rows[ridx].height = Cm(0.85)
        table.rows[ridx].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for cidx, value in enumerate(row_values):
            set_cell_text(table.rows[ridx].cells[cidx], value, cidx % 2 == 0, 10, WD_ALIGN_PARAGRAPH.LEFT)
            if cidx % 2 == 0:
                set_cell_shading(table.rows[ridx].cells[cidx], LIGHT_GRAY)

    labels = [
        ("施工单位审核意见", "审核人：                    年    月    日"),
        ("施工单位审批意见", "批准人：                    年    月    日"),
        ("监理单位审查意见", "专业监理工程师：            年    月    日"),
        ("建设单位审批意见", "负责人：                    年    月    日"),
        ("产权/使用单位意见", "负责人：                    年    月    日"),
    ]
    for offset, (label, sign) in enumerate(labels, start=3):
        row = table.rows[offset]
        row.cells[0].merge(row.cells[3])
        cell = row.cells[0]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label + "\n\n")
        set_run_font(r, TITLE_FONT, 11, True)
        r = p.add_run(sign)
        set_run_font(r, BODY_FONT, 10.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        row.height = Cm(3.0)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        prevent_row_split(row)


def add_toc(doc):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目  录")
    set_run_font(r, TITLE_FONT, 18, True)
    items = [
        "一、编制依据及编制说明",
        "二、工程概况及清单响应",
        "三、施工组织与总体部署",
        "四、施工准备",
        "五、总体施工程序及施工分区",
        "六、原燃气系统停用、置换与拆除",
        "七、电气系统安装工程",
        "八、甲供厨房设备安装与接电",
        "九、土建、给排水及装饰恢复工程",
        "十、施工进度计划与工期保证",
        "十一、质量保证体系及检验计划",
        "十二、安全生产及文明施工",
        "十三、主要危险源辨识与控制",
        "十四、环境保护、食堂卫生与成品保护",
        "十五、劳动力、机械和材料计划",
        "十六、试验、调试、资料与验收",
        "十七、应急预案与事故处置",
        "十八、雨季、高温及冬季施工措施",
        "十九、工程移交与保修",
        "附表：施工进度横道图",
    ]
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        p.paragraph_format.right_indent = Cm(1.0)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(item)
        set_run_font(r, TABLE_FONT, 11.5, idx <= 19 and item.startswith(("一", "二", "三", "四", "五", "六", "七", "八", "九", "十")))
    doc.add_page_break()


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    remove_existing_body(doc)

    first = doc.sections[0]
    configure_section(first, landscape=False)
    first.different_first_page_header_footer = True
    clear_header_footer(first, "")

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    add_cover(doc)
    add_approval_page(doc)
    add_toc(doc)

    add_heading(doc, "一、编制依据及编制说明", 1)
    add_heading(doc, "1.1 编制依据", 2)
    add_paragraphs(doc, [
        f"本施工方案依据《{PROJECT}工程量清单》、建设单位现场管理制度、甲供设备技术资料、经会审确认的施工图或现场深化成果以及现行国家、行业施工质量验收和安全技术标准编制。清单共49项，覆盖电气、燃气停用拆除、混凝土与砌体、屋面防水、建筑地面、排水沟及管道、脚手和机械台班等内容。",
        "规范查新以2026年7月31日为基准。强制性工程建设规范与配套验收标准规定不一致时，执行强制性工程建设规范；国家、行业标准在施工期间发布新版本并已经实施的，项目部应组织二次查新，经技术负责人确认后按新标准执行，不得继续引用已废止版本。",
        "本方案是施工组织和技术管理文件，不代替燃气停用专项操作票、盲板抽堵作业票、动火作业票、停送电操作票、吊装及高处作业审批。涉及业主能源系统的阀门操作、停气、送气、停电和送电，必须由有权限的人员按建设单位制度实施或监护，施工单位不得擅自操作。",
    ])

    standards = [
        ["1", "GB 55024-2022", "建筑电气与智能化通用规范", "电气安装、保护接地和验收总控"],
        ["2", "GB 50303-2015", "建筑电气工程施工质量验收规范", "盘柜、母线、桥架、电缆及接地验收；与GB 55024不一致时从其规定"],
        ["3", "GB 50168-2018", "电气装置安装工程 电缆线路施工及验收标准", "电缆敷设、终端制作和试验"],
        ["4", "GB 50169-2016", "电气装置安装工程 接地装置施工及验收规范", "接地及等电位"],
        ["5", "GB 50254-2014", "电气装置安装工程 低压电器施工及验收规范", "低压设备和电器安装"],
        ["6", "JGJ/T 46-2024", "建筑与市政工程施工现场临时用电安全技术标准", "施工临时用电；JGJ 46-2005已废止"],
        ["7", "GB 55009-2021", "燃气工程项目规范", "介质确认为天然气或城镇燃气时执行"],
        ["8", "GB 6222-2025", "工业企业煤气安全规范", "介质确认为工业煤气时执行；2026-05-01实施"],
        ["9", "GB 30871-2022", "危险化学品企业特殊作业安全规范", "适用范围满足时用于动火、盲板抽堵、高处、吊装等特殊作业"],
        ["10", "GB 55037-2022", "建筑防火通用规范", "防火及施工消防总控"],
        ["11", "GB 50204-2015", "混凝土结构工程施工质量验收规范", "C30基础和结构"],
        ["12", "GB 50203-2011", "砌体结构工程施工质量验收规范", "墙体拆除后恢复"],
        ["13", "GB 50207-2012", "屋面工程质量验收规范", "SBS屋面防水恢复"],
        ["14", "GB 50209-2010", "建筑地面工程施工质量验收规范", "不发火防静电地面和地砖恢复"],
        ["15", "GB 50210-2018", "建筑装饰装修工程质量验收标准", "墙面、吊顶、门窗和饰面"],
        ["16", "GB 50242-2002", "建筑给水排水及采暖工程施工质量验收规范", "排水管及沟槽工程"],
        ["17", "GB 50300-2013", "建筑工程施工质量验收统一标准", "分项、分部和单位工程验收"],
        ["18", "JGJ 80-2016", "建筑施工高处作业安全技术规范", "脚手架、登高平台和临边作业"],
        ["19", "GB 6095-2021", "坠落防护 安全带", "高处作业个人防护"],
        ["20", "GB/T 29639-2020", "生产经营单位生产安全事故应急预案编制导则", "应急预案和现场处置"],
    ]
    add_table(doc, ["序号", "标准编号", "标准名称", "本项目主要适用内容"], standards, [1.1, 3.0, 5.3, 5.3], 8.8)
    add_callout(doc, "版本控制提示", "GB 6222-2005已于2026年5月1日废止，工业煤气相关作业应执行GB 6222-2025；JGJ 46-2005已由JGJ/T 46-2024替代。本方案不得在报审或交底时回退使用旧版编号。")

    add_picture(
        doc,
        NORMS / "GB55024-2022_effective.png",
        "图1-1  GB 55024-2022发布及实施信息网络查新截图",
        14.4,
        "住房和城乡建设部国家标准公告（国标电子书库页面）",
    )
    add_picture(
        doc,
        NORMS / "GB6222-2025_status.png",
        "图1-2  GB 6222-2025现行状态网络查新截图",
        14.4,
        "全国标准信息公共服务平台",
    )

    add_heading(doc, "1.2 编制原则", 2)
    add_paragraphs(doc, [
        "施工组织坚持“先确认介质和能源边界、后实施隔断；先拆除和土建、后机电安装；先隐蔽试验、后封闭恢复；先单机调试、后负荷联调；施工记录与实体同步”的原则。任何赶工措施不得压缩置换检测、混凝土养护、防水检验、电气试验和安全审批的必要时间。",
        "食堂属于人员密集和食品加工场所，施工区与运营区必须硬隔离。粉尘、噪声、异味、飞溅和临时断电可能影响食品加工时，应与使用单位确定停业或错峰窗口，完成彻底清洁、消毒和安全确认后方可恢复使用。",
    ])

    add_heading(doc, "二、工程概况及清单响应", 1)
    add_heading(doc, "2.1 工程基本信息", 2)
    project_info = [
        ["工程名称", PROJECT],
        ["施工地点", "武汉钢铁有限公司质检中心食堂，准确边界以现场交底为准"],
        ["施工单位", COMPANY],
        ["计划开工", START],
        ["计划完工", FINISH],
        ["理论总工期", DURATION],
        ["质量目标", "符合现行国家、行业标准和建设单位要求，一次验收合格"],
        ["安全目标", "杜绝燃气中毒窒息、火灾爆炸、触电、高处坠落、起重伤害及食品污染事故"],
    ]
    add_table(doc, ["项目", "内容"], project_info, [3.2, 11.0], 10)

    add_heading(doc, "2.2 工程范围", 2)
    add_paragraphs(doc, [
        "本项目总体目标是将食堂原有燃气用能系统安全退出，拆除不再使用的DN120燃气管道，并新增低压配电、封闭母线、桥架、电缆、配管、接地及等电位系统，为甲供成套厨房电气设备提供可靠电源；同步完成设备基础、结构、墙洞、屋面、地面、墙面、排水及相关装饰恢复。",
        "电气工程包括GGD低压配电柜2台、JXF3户外不锈钢动力箱1台、截面240mm²及以下动力电缆280m、截面10mm²及以下动力电缆450m、封闭铜母线TMY-3[2(125×10)]+1(125×10)] 20m、带接头金属软管Φ100为15m、Φ32为20m、镀锌钢管5.262t、镀锌钢构件0.503t、接地用镀锌钢材0.038t、BVR 1×6黄绿导线100m、200×100带盖镀锌桥架1.364t和TD28等电位联结端子箱1台。",
        "建筑及给排水工程包括墙体开洞600×300和200×200各1处、甲供厨房成套设备1套、C30钢筋混凝土基础4.02m³、C30钢筋混凝土结构10.56m³、200mm墙体拆除15.96m²、屋面防水层拆除35.28m²、室内地砖地坪开挖14.4m²、室外电缆沟3.6m、屋面3+3厚SBS防水及保温保护层52.92m²、墙体和内外墙恢复、80mm地砖踢脚12.24m、不发火防静电地面196.02m²、吊顶14.59m²、60系列铝合金推拉窗6.48m²、C20混凝土散水1.15m³、PVC雨水管DN110为4.2m和800×800地砖恢复18m²。",
        "排水及道路恢复包括DN120不锈钢管回装18m、预制钢筋混凝土电缆沟盖板更换4.8m³、新建净宽净深300×300排水沟40m、200mm混凝土道路拆除5m²、人行道恢复5m²、PVC-U DN110排水管更换15m、垃圾外运21.61m³及余土处置23.771t。措施项目包括施工围挡350m²、脚手架220m²、12m登高平台5台班、25t汽车吊2台班和10t平板车2台班。",
    ])

    add_heading(doc, "2.3 工程特点和难点", 2)
    challenges = [
        ["1", "燃气介质及边界必须先确认", "清单使用“燃气”表述，可能涉及天然气或工业煤气。两类介质适用规范和检测要求不同，必须由建设单位能源管理部门书面确认。"],
        ["2", "不停产/不停业条件下交叉施工", "施工区邻近食堂运营区域，人员流动、食品卫生、噪声粉尘和临时断电均需精细协调。"],
        ["3", "甲供设备接口风险", "厨房设备型号、功率、进线方向、保护电器和固定孔位如未提前核对，易造成基础、母线和电缆返工。"],
        ["4", "多专业恢复量大", "电气改造伴随混凝土、砌体、屋面防水、防静电不发火地面、排水沟和装饰恢复，工序衔接是工期控制重点。"],
        ["5", "送电和负荷试运行风险", "厨房设备集中用电，必须完成容量、相序、保护、接地、绝缘和温升检查，严禁直接带负荷试送电。"],
    ]
    add_table(doc, ["序号", "工程特点", "主要控制要求"], challenges, [1.0, 4.0, 9.0], 9.2)

    add_heading(doc, "三、施工组织与总体部署", 1)
    add_heading(doc, "3.1 项目组织机构", 2)
    add_paragraphs(doc, [
        "项目实行项目经理负责制，设置项目经理、技术负责人、施工员、安全员、质量员、材料员、资料员及电气、燃气拆除、土建装饰、设备安装、检测调试班组。项目经理对工期、质量、安全、协调和资源投入负责；技术负责人组织图纸会审、专项交底、危险作业技术复核和试验调试；安全员对停气、盲板、置换、动火、临电、吊装、高处和食堂隔离进行全过程检查。",
        "燃气停用作业实行建设单位能源管理部门、食堂使用单位、施工项目部三方交接。电气停送电实行工作票或等效书面制度，送电前由电气技术负责人、质量员、业主电气管理人员共同确认。甲供设备安装实行开箱、基础、就位、接口、单机、负荷六个检查点，未经上一检查点确认不得进入下一工序。",
    ])

    organization = [
        ["岗位", "主要职责"],
        ["项目经理", "统筹工期、资源、业主协调和危险作业审批，签发项目内部开工条件确认。"],
        ["技术负责人", "方案交底、深化复核、技术核定、燃气置换计算、电气试验和竣工技术资料审核。"],
        ["施工员", "作业面移交、测量放线、日计划、工序交接、隐蔽验收申请和现场签证。"],
        ["安全员", "作业票、能源隔离、气体检测、动火监护、高处吊装、临电、围挡和食堂卫生隔离检查。"],
        ["质量员", "材料报验、样板、实测实量、试验仪器有效性、质量记录和问题销项。"],
        ["材料员", "甲供设备交接、材料计划、进场验收、储存防护和不合格品隔离。"],
        ["资料员", "同步收集审批、交底、报验、隐蔽、试验、调试、照片、竣工图和移交资料。"],
    ]
    add_table(doc, organization[0], organization[1:], [3.2, 10.8], 9.4)

    add_heading(doc, "3.2 总体部署", 2)
    add_paragraphs(doc, [
        "现场划分为燃气停用拆除区、配电柜及母线安装区、桥架配管与电缆敷设区、厨房设备安装区、土建恢复区和材料暂存区。各区采用硬质围挡、警示标识和独立通道；燃气置换放散点、电气试验区、吊装半径及屋面作业区设置专项警戒。",
        "总体施工顺序为：进场准备和深化核对→食堂分区隔离及临时设施→燃气停用审批、可靠隔断、氮气置换和检测→原管道分段拆除→墙洞、沟槽及基础施工→桥架、配管、接地→盘柜和母线→电缆敷设及终端→甲供设备就位接线→土建防水和装饰恢复→电气试验和单机调试→负荷试运行→清洁消毒、联合验收和移交。",
        "工期组织采用燃气拆除和电气深化先行、土建与机电平行穿插、装饰恢复分区封闭的方式。设备基础和主要沟槽完成后，盘柜母线、桥架配管、屋面墙面恢复分别形成独立流水；不得在防水、地坪和洁净恢复完成面上继续进行无保护的电焊、切割或重物拖运。",
    ])

    add_heading(doc, "四、施工准备", 1)
    add_heading(doc, "4.1 技术准备", 2)
    for idx, text in enumerate([
        "组织清单、施工图、设备资料和现场条件四方核对，形成49项清单响应台账。重点确认两台GGD柜型号及进出线、母线额定电流及短路耐受、主电缆型号和回路、设备总功率、需用系数、保护电器整定、桥架路由、墙洞位置、基础标高、设备固定孔和检修空间。",
        "由建设单位书面确认原燃气介质、运行压力、上游和下游隔断点、放散点、停气时间、阀门操作责任人、检测项目及合格判据。未取得确认文件，不进行盲板、置换、拆除和动火。",
        "复测配电间、厨房和屋面现状，标注既有电缆、给排水、消防、燃气和结构构件。墙体钻孔、地坪切割和沟槽开挖前采用图纸核对、仪器探测和人工试探相结合，严禁盲目机械开挖。",
        "编制临时用电方案、吊装作业计划、燃气置换操作卡和送电调试方案，完成班组书面安全技术交底。涉及危险性较大的分部分项工程时，按法规及建设单位制度另行编制专项方案并履行审批。",
    ], 1):
        add_bullet(doc, text, idx)

    add_heading(doc, "4.2 现场及资源准备", 2)
    add_paragraphs(doc, [
        "施工围挡按清单350m²统筹布置，食堂运营界面采用封闭式防尘隔断，设置带门禁的施工通道。材料卸车、25t汽车吊、10t平板车和12m登高平台行走路线提前确认，吊装支腿位置复核地面承载和地下设施，必要时铺设钢板分散荷载。",
        "临时配电采用专用配电箱、分级保护和保护导体，潮湿和金属容器附近使用适宜安全电压照明。电焊机、切割机、卷扬或牵引设备、手持电动工具逐台检查，漏电保护器按制度试验并记录。",
        "气体检测仪、绝缘电阻表、接地电阻表、相序表、钳形电流表、扭矩扳手和压力试验器具均应在有效检定或校准期。可燃气体检测仪的传感器类型必须与实际介质匹配，不得用不适配仪器代替。",
    ])

    add_heading(doc, "五、总体施工程序及施工分区", 1)
    add_picture(doc, ASSETS / "总体施工示意图.png", "图5-1  质检中心食堂气改电总体施工示意图", 14.5)
    add_paragraphs(doc, [
        "图示表达施工组织关系，不替代经审查的电气系统图、设备接线图和燃气隔断点图。正式施工前应把盘柜编号、母线段、桥架分支、设备回路、PE导体及等电位端子逐项编码，并在现场标识和竣工图中保持一致。",
        "施工过程中实行红、黄、绿三色状态管理：红色表示燃气未置换或电气未隔离，禁止施工；黄色表示已隔离但正在检测、试验或待验收，限制进入；绿色表示交接完成且具备下一工序条件。状态转换必须以书面记录为依据。",
    ])

    add_heading(doc, "六、原燃气系统停用、置换与拆除", 1)
    add_heading(doc, "6.1 适用规范及介质确认", 2)
    add_paragraphs(doc, [
        "原管道在清单中表述为DN120钢制燃气管道，拆除量3.6t，估算规格φ133×4.5、长度约252m。施工前必须由建设单位明确介质：若为天然气或城镇燃气，执行GB 55009-2021及燃气经营/管理单位的停气规定；若为工业企业煤气，执行GB 6222-2025。两类介质不得只凭经验判断，也不得混用检测指标。",
        "阀门关闭不作为唯一的可靠隔断。隔断方式、盲板材质厚度、安装位置、法兰等级和密封垫由技术人员根据管道参数及业主制度确定；盲板应编号、挂牌并纳入台账，抽堵作业由具备资格和经验的人员实施，设置专人监护。",
    ])
    add_picture(
        doc,
        NORMS / "GB55009-2021_cover.png",
        "图6-1  GB 55009-2021《燃气工程项目规范》封面网络检索截图",
        9.0,
        "网络检索标准文本封面，正式执行版本以建设单位标准库为准",
    )
    add_picture(doc, ASSETS / "燃气停用置换示意图.png", "图6-2  原燃气管道停气、隔断与氮气置换工序示意图", 14.5)

    add_heading(doc, "6.2 停气隔断与置换准备", 2)
    steps = [
        ["1", "作业申请", "提出停气、盲板、置换、拆除和可能的动火申请，明确时间、范围、责任人和应急措施。"],
        ["2", "系统识别", "核对管线走向、阀门、支路、末端及与运行系统的连接，现场挂牌并绘制隔断示意图。"],
        ["3", "停止用气", "由授权人员按操作规程关闭气源，确认相关设备停止，防止误启。"],
        ["4", "可靠隔断", "安装经确认的盲板或采用物理断开方式，阀门挂牌上锁，登记隔断状态。"],
        ["5", "置换准备", "氮气瓶组可靠固定，配置汇流排、减压器、导静电软管和防回流措施；放散引至安全地点。"],
        ["6", "警戒监护", "划定警戒区，停止无关动火和车辆进入，准备灭火器材、通讯和应急救援器材。"],
    ]
    add_table(doc, ["步骤", "工作", "控制要求"], steps, [1.1, 2.8, 10.0], 9.2)

    add_heading(doc, "6.3 氮气置换计算与操作", 2)
    add_paragraphs(doc, [
        "按清单估算φ133×4.5钢管，内径约124mm，长度约252m，管内容积约为π÷4×0.124²×252≈3.04m³。清单暂按不少于5倍管内容积考虑，理论置换量约15.2m³，另计管路死角、泄漏和操作余量，清单暂列氮气40m³。实际氮气量不得仅以理论倍数判定结束，必须以批准的操作程序和连续气体检测结果为准。",
        "置换宜从低点缓慢送入、从高点或末端安全放散，控制流量，防止静电和压力冲击。放散口远离人员、火源、进风口和低洼积聚处，设置警戒和专人监护。氮气具有窒息危险，严禁在室内或通风不良处直接放散，作业人员不得进入可能形成缺氧的区域。",
        "检测仪应先自检或标定，取样位置覆盖管道首端、末端、支路和可能积聚部位。检测项目、频次、合格值及连续稳定时间按实际介质标准和建设单位制度执行，结果由检测人、监护人和业主接口人员签字。严禁用明火试漏或判断气体是否存在。",
    ])

    add_heading(doc, "6.4 管道拆除", 2)
    add_paragraphs(doc, [
        "置换检测合格并完成书面交接后，优先采用冷切割和无火花工具。确需动火时，重新办理动火审批，作业前、作业中和间断后复测气体，清理周边可燃物，配置接火和灭火措施，监护人不得离岗。",
        "拆除采用“先支撑、后切割；先支管、后主管；分段固定、分段落地”的顺序。高处管段设置可靠吊点和导向绳，禁止直接抛掷。每段拆除前确认内部无积液、无残余压力，切口和遗留管口及时采用金属盲板或经确认方式封闭并挂牌。",
        "拆除钢管分类堆放，按建设单位要求办理废旧物资交接。地面、墙面和屋面穿越处清理后进行防水、防火及装饰恢复；遗留阀门、支架或接口若不在拆除范围内，应做永久标识和防腐处理。",
    ])

    add_heading(doc, "七、电气系统安装工程", 1)
    add_heading(doc, "7.1 施工前电气技术核对", 2)
    add_paragraphs(doc, [
        "设备采购和施工前完成负荷表复核，包括每台厨房设备额定功率、电压、相数、启动特性、功率因数、同时使用系数和保护要求。根据核对结果校验变压器及上级电源余量、两台GGD柜和动力箱额定参数、母线和电缆载流量、短路耐受、保护选择性及电压降。若甲供设备参数与设计不一致，须形成设计或建设单位书面确认。",
        "建立电气回路编号表，包含上级电源、柜号、开关号、母线段、桥架段、电缆型号、起讫点、设备名称、PE连接点和试验结果。编号在图纸、标签、试验记录和竣工资料中一致。",
    ])

    add_heading(doc, "7.2 配电柜及动力箱安装", 2)
    add_paragraphs(doc, [
        "GGD柜和动力箱到场后进行开箱验收，核对型号、尺寸、元件、母排、出线方式、防护等级、铭牌、合格证及随箱资料。柜体无变形、掉漆和受潮，元件固定可靠，母线绝缘支持完整。甲供设备的缺件、损伤或参数偏差形成开箱记录，不得自行拆改。",
        "基础槽钢和混凝土基础轴线、标高、水平度及预留孔经验收后安装柜体。柜列排列整齐，柜间连接紧固，门开启灵活，检修通道满足要求。柜体及可开启金属门设置可靠保护接地，接地导体截面和连接方式符合设计。",
        "二次回路端子排列清晰，导线留有适当余量，线号与图纸一致。断路器、隔离电器和保护装置在送电前检查整定值、机械操作、联锁和分合闸状态，未经授权不得改变整定参数。",
    ])

    add_heading(doc, "7.3 封闭母线安装", 2)
    add_paragraphs(doc, [
        "封闭铜母线20m安装前复核相序、规格、分段、转角、穿墙、伸缩节及与盘柜连接尺寸。支吊架由厂家资料和设计确定，生根可靠，禁止利用工艺管道或不明承载构件作为支撑。母线外壳连接连续，穿越墙体处设置套管及防火封堵。",
        "母线连接面保持清洁干燥，螺栓按厂家规定顺序和扭矩紧固，使用校验合格的扭矩扳手并作标记。不同金属连接按设计设置过渡措施。安装后检查绝缘距离、外壳跨接、相色和方向，完成绝缘电阻及规定试验后方可连接负荷。",
    ])

    add_heading(doc, "7.4 桥架、钢管及支架安装", 2)
    add_paragraphs(doc, [
        "200×100镀锌桥架带盖板，清单量1.364t。路由深化应避开高温、蒸汽、给排水易漏点和设备检修空间，跨越通道处保持安全净高。支架间距、转弯支撑、伸缩补偿和抗震要求按设计及产品文件执行，现场不得仅凭经验任意加大间距。",
        "桥架切割使用机械方法，切口平整并防腐补口；桥架接头板螺栓齐全，盖板固定可靠。金属桥架全长保护接地连续，连接板不能当然视为唯一接地跨接时，应按设计设置跨接导体。桥架内电缆按回路和电压等级分层分区，填充率和弯曲半径满足要求。",
        "镀锌钢管5.262t及金属软管Φ100、Φ32施工前复核管径、弯曲半径和接线口。钢管连接可靠，螺纹和切口作防腐处理，管口设护口；竖向和设备末端固定牢固。金属软管只用于允许活动或减振连接，不得作为永久保护接地导体，也不得过长、扭曲或受拉。",
    ])

    add_heading(doc, "7.5 电缆敷设和终端", 2)
    add_paragraphs(doc, [
        "动力电缆共730m，其中截面240mm²及以下280m、10mm²及以下450m。敷设前核对盘号、长度、起讫点和路径，检查外观并测量绝缘电阻。大截面电缆采用放线架和适宜牵引工具，控制牵引力、侧压力和弯曲半径，严禁在地面锐角、积水和污染物上拖拽。",
        "电缆在桥架转弯、垂直段、进柜和设备端固定，排列整齐，不交叉挤压。两端和中间可识别位置设置耐久标牌。穿墙、穿楼板和进柜孔洞完成防火封堵、防水及防鼠措施，封堵前拍照并办理隐蔽验收。",
        "终端制作由熟练电工实施，剥切不得损伤导体，铜端子与导体截面匹配，采用合格压接模具，压接后检查外观并做防松、绝缘和相色标识。连接螺栓按设备要求紧固，送电前复测绝缘、核对相序及回路。",
    ])

    add_heading(doc, "7.6 接地及等电位", 2)
    add_paragraphs(doc, [
        "接地用镀锌钢材0.038t、BVR 1×6黄绿导线100m和TD28等电位端子箱1台按设计施工。配电柜、动力箱、封闭母线外壳、桥架、钢管、厨房设备金属外壳及其他可导电部分均应可靠接入PE系统或等电位系统。",
        "接地连接位置清理至金属本体，螺栓连接设置防松和防腐，焊接长度及焊缝质量符合规范。黄绿双色导体不得挪作相线或中性线。隐蔽前检查连接连续性并拍照，系统完成后测量接地电阻和保护导体连续性，结果符合设计和规范。",
    ])

    add_heading(doc, "八、甲供厨房设备安装与接电", 1)
    add_paragraphs(doc, [
        "甲供成套厨房设备1套到场后，施工、使用、供货和监理/建设单位共同开箱验收，核对型号、数量、额定功率、额定电流、相数、进线口、接地端子、随机配件、安装说明和质保文件。设备缺陷及接口变化必须在就位前关闭。",
        "设备基础和地面完成面达到安装条件后放线定位。设备运输采用适宜叉运、滚杠、吊装或搬运设备，保护不锈钢外壳、控制面板和门体，严禁在防静电不发火地面上无保护拖拉。设备就位后调整水平、垂直和相互间距，固定牢靠并保留检修空间。",
        "电气接线前确认上级开关断开并执行挂牌上锁。电缆型号、端子和保护电器与设备铭牌一致，PE导体先接后拆，软连接长度适宜且有机械保护。可能受水、油和蒸汽影响的接线盒、隔离开关和电缆引入口应具备相应防护。",
        "单机试运行按设备说明书进行，依次检查绝缘、接地、相序、控制回路、急停、门联锁、温控、风机、加热元件和异常报警。空载合格后在使用单位配合下进行负荷试运行，记录工作电流、温升、噪声、振动和保护动作。发现异味、异常升温或跳闸立即停机检查。",
    ])

    add_heading(doc, "九、土建、给排水及装饰恢复工程", 1)
    add_heading(doc, "9.1 拆除及洞口施工", 2)
    add_paragraphs(doc, [
        "拆除200mm墙体15.96m²、屋面防水35.28m²、室内地砖地坪14.4m²及局部道路前，先划定边界、切割分离并保护保留结构和设施。墙洞600×300和200×200各1处，开洞前确认不影响主体结构及暗埋管线；如涉及结构构件，必须取得设计确认。",
        "拆除采用小型机具分段施工，湿法降尘，垃圾及时装袋或覆盖外运。高处和屋面拆除设置防坠和防物体打击措施，洞口形成后立即加设防护或临时封闭。拆除面应切边整齐，便于后续恢复。",
    ])

    add_heading(doc, "9.2 混凝土基础及结构", 2)
    add_paragraphs(doc, [
        "C30钢筋混凝土基础4.02m³和C30钢筋混凝土结构10.56m³施工前复核设备荷载、基础尺寸、标高、预埋件和地脚螺栓。模板刚度、钢筋保护层、预留孔和接地预埋经验收后浇筑。",
        "混凝土浇筑连续进行，分层振捣密实，防止漏振、涨模和预埋件位移。浇筑后及时养护，按规定制作试件。设备安装前确认混凝土强度满足设计及设备安装要求，基础表面平整，地脚孔清洁。",
    ])

    add_heading(doc, "9.3 电缆沟、排水沟及管道", 2)
    add_paragraphs(doc, [
        "室外电缆沟3.6m按200×200及钢筋混凝土盖板要求施工，既有预制沟盖板更换约4.8m³。沟槽开挖前探查地下管线，设置临边和排水措施；沟底标高、坡度和承载检查合格后施工垫层、沟体及盖板。",
        "新建排水沟40m，净宽净深300×300，底和壁采用100mm厚C20混凝土；PVC-U DN110排水管更换15m，壁厚不小于3.2mm。沟体模板顺直，坡向正确，接口严密。完成后进行通水检查，确认无倒坡、渗漏和积水，再恢复地面。",
        "DN120不锈钢管回装18m应确认用途和连接方式，管道切割、坡口、焊接或机械连接符合相应工艺，安装后进行外观和功能检查。PVC雨水管DN110为4.2m，卡箍间距、伸缩和出水位置合理，避免污染墙面。",
    ])

    add_heading(doc, "9.4 屋面防水、墙面及地面恢复", 2)
    add_paragraphs(doc, [
        "屋面恢复52.92m²采用3+3厚SBS防水、70mm挤塑聚苯板保温和20mm水泥砂浆保护层。基层坚实干燥、阴阳角圆顺，管根、洞口和收头先做附加层；卷材搭接、热熔和收头密封符合规范。施工完成进行淋水或蓄水检查，合格后做保护层。",
        "墙体恢复按240mm砖墙及抹灰、腻子、防霉涂料等清单做法实施，内外墙饰面分别为85.68m²和67.32m²。新旧墙连接设置拉结或加强措施，抹灰分层施工，防止空鼓开裂。涂料颜色先做样板并由使用单位确认。",
        "不发火防静电地面196.02m²按23J909-3-106 KM1及设计要求施工。基层强度、含水率、平整度和接地网络验收后进行面层施工，材料配合比和厚度受控；完成后检查不发火、防静电性能、平整度、空鼓和裂缝。800×800地砖恢复18m²，排版、标高、缝宽及与既有地面的色差受控。",
        "吊顶14.59m²、60系列5+9+5铝合金推拉窗6.48m²、80mm地砖踢脚12.24m和C20散水1.15m³按样板施工。窗框固定、防水密封和开启功能检查合格，吊顶封板前完成上方机电隐蔽验收。",
    ])

    add_heading(doc, "十、施工进度计划与工期保证", 1)
    add_paragraphs(doc, [
        f"计划工期为{START}至{FINISH}，共{DURATION}。进度计划以燃气安全退出、设备基础和沟槽、盘柜母线、桥架电缆、甲供设备接电、建筑恢复、电气试验、负荷试运行为主要控制节点。完整横道图见本方案附表及配套Excel文件。",
        "材料采购和深化设计必须前置。两台GGD柜、封闭母线、动力电缆、桥架和甲供厨房设备均可能影响关键线路，项目部在开工后7日内形成到货计划和技术问题清单，每周滚动更新。设备参数未确认造成的返工不得通过压缩试验和安全工序弥补。",
    ])
    schedule_rows = [
        ["1", "8月10日-8月16日", "进场、会审、清单和接口确认"],
        ["2", "8月12日-9月5日", "深化、材料报审、甲供设备到货策划"],
        ["3", "8月17日-8月28日", "围挡、临电、脚手、墙洞准备"],
        ["4", "8月24日-9月10日", "停气隔断、置换检测及燃气管拆除"],
        ["5", "8月24日-9月30日", "拆除、沟槽、基础及混凝土结构"],
        ["6", "9月12日-10月20日", "桥架配管接地、盘柜和母线安装"],
        ["7", "10月8日-11月5日", "电缆敷设、终端及防火封堵"],
        ["8", "10月25日-11月20日", "厨房设备就位和电气末端连接"],
        ["9", "9月20日-11月30日", "屋面、墙面、地面、排水和装饰恢复"],
        ["10", "11月21日-12月22日", "电气试验、单机调试及负荷试运行"],
        ["11", "12月23日-12月31日", "整改、资料、联合验收和移交"],
    ]
    add_table(doc, ["序号", "计划时间", "主要工作"], schedule_rows, [1.2, 4.2, 9.6], 9.4)
    add_paragraphs(doc, [
        "项目部实行周计划、三日滚动计划和日协调。周计划检查关键材料、审批窗口和前置条件；三日计划落实人员机械和作业面；日协调解决食堂运营、停送电、噪声粉尘、吊装通道和交叉施工。关键节点延误时，优先增加平行作业面、调整班次和提前预制，不盲目延长高风险作业时间。",
        "进度考核不以单纯完成工程量为准，必须同时具备验收记录和可移交条件。未完成隐蔽验收、防水检查、电气试验或清洁恢复的工作不得计作实质完成。",
    ])

    add_heading(doc, "十一、质量保证体系及检验计划", 1)
    add_heading(doc, "11.1 质量管理制度", 2)
    add_paragraphs(doc, [
        "项目执行材料进场验收、样板先行、首件认可、工序交接、隐蔽验收、实测实量和质量问题销项制度。对燃气隔断、混凝土预埋、接地、桥架封闭前、电缆终端、防火封堵、屋面防水、吊顶封板和送电试验设置停检点，未经确认不得转序。",
        "质量记录与实体同步。检测仪器和扭矩工具在有效期，试验人员具备相应能力。甲供设备不因甲供而免除到货验收，施工单位对保管、搬运、安装和接线质量负责。",
    ])
    quality_rows = [
        ["燃气隔断", "隔断点、盲板编号、挂牌上锁、交接状态", "停气交接及盲板台账"],
        ["气体置换", "介质匹配、放散、检测点、结果和签字", "置换检测记录"],
        ["基础混凝土", "轴线、标高、钢筋、预埋、强度", "隐蔽及试块报告"],
        ["配电柜", "基础、垂直度、连接、接地、元件整定", "盘柜安装检查记录"],
        ["封闭母线", "支架、连接扭矩、相序、绝缘和外壳接地", "母线安装及试验记录"],
        ["桥架配管", "路由、支撑、弯曲、跨接、防腐补口", "隐蔽验收记录"],
        ["电缆", "型号、弯曲半径、固定、标牌、绝缘和终端", "电缆敷设/绝缘记录"],
        ["接地等电位", "连接连续、黄绿标识、接地电阻", "接地测试记录"],
        ["屋面防水", "基层、附加层、搭接、收头及蓄淋水", "防水隐蔽及试验记录"],
        ["防静电不发火地面", "基层、厚度、平整、接地及性能", "分项验收/检测报告"],
        ["厨房设备", "定位、固定、接线、保护、单机和负荷", "设备安装调试记录"],
        ["排水", "坡向、接口、通水、无渗漏和积水", "通水试验记录"],
    ]
    add_table(doc, ["检验对象", "主要检查内容", "形成记录"], quality_rows, [3.0, 7.4, 3.6], 8.9)

    add_heading(doc, "十二、安全生产及文明施工", 1)
    add_heading(doc, "12.1 安全管理总要求", 2)
    add_paragraphs(doc, [
        "所有人员进场接受厂区和项目安全教育，特种作业人员持有效证件。每日班前会说明当日能源状态、作业票、交叉作业、天气和应急撤离路线。安全员对作业许可有效期、监护人到位、隔离状态和防护设施进行确认。",
        "危险作业实行“作业前确认、作业中监护、作业后清场”。燃气、动火、吊装、高处、临时用电、动土和盲板抽堵不得以同一张泛化作业票代替专项审批。作业条件变化、停工后复工或作业超出批准时间时重新确认。",
    ])
    add_picture(
        doc,
        NORMS / "GB30871-2022_status.png",
        "图12-1  GB 30871-2022现行状态网络查新截图",
        14.4,
        "全国标准信息公共服务平台；仅在适用范围满足时作为强制执行依据",
    )
    add_picture(
        doc,
        NORMS / "JGJT46-2024_effective.png",
        "图12-2  JGJ/T 46-2024实施信息网络查新截图",
        14.4,
        "网络查新页面，显示2025年1月1日实施并废止JGJ 46-2005",
    )

    add_heading(doc, "12.2 专项安全措施", 2)
    safety_items = [
        "燃气作业区禁止无关火源、车辆和人员进入，设置风向标和撤离方向。检测和监护连续进行，氮气瓶组防倾倒，放散区防止缺氧。",
        "停送电由授权人员执行，开关断开后验电、挂牌、上锁并建立锁具登记。多班组作业时采用个人锁或等效联锁管理，未清点人员和工具不得送电。",
        "临时用电按JGJ/T 46-2024组织设计、验收、检查和拆除。配电箱防雨上锁，电缆避免机械损伤，手持工具和潮湿场所保护措施完整。",
        "脚手架220m²搭设后验收挂牌，12m登高平台由合格人员操作，作业人员系挂安全带，工具材料防坠。不得站在桥架、管道或设备顶部作业。",
        "25t汽车吊作业前核实吊物重量、幅度、地面和吊点，支腿全伸并垫实，设置警戒和专职指挥。恶劣天气或视线不良时停止吊装。",
        "切割、焊接和打磨办理动火审批，清理可燃物，设置接火盘和灭火器，作业后留守检查。燃气管道即使已置换，也须按批准条件复测后方可动火。",
        "沟槽和洞口设置盖板、护栏、警示和夜间照明。开挖土方距沟边保持安全距离，雨后检查边坡和积水，禁止人员在无防护沟内作业。",
    ]
    for idx, text in enumerate(safety_items, 1):
        add_bullet(doc, text, idx)

    add_heading(doc, "十三、主要危险源辨识与控制", 1)
    risks = [
        ["燃气泄漏/中毒窒息", "停气、隔断、置换和拆除", "介质确认、可靠隔断、连续检测、通风警戒、专人监护、应急撤离"],
        ["火灾爆炸", "燃气残留、动火、切割", "冷切优先、动火审批、复测、接火、灭火器、清场留守"],
        ["氮气窒息", "瓶组、放散和管内置换", "室外安全放散、通风、氧气检测、警戒、禁止进入低洼或封闭区"],
        ["触电/电弧", "临电、盘柜、试送电", "验电上锁、绝缘防护、分级保护、授权送电、试验区警戒"],
        ["高处坠落", "脚手、屋面、登高平台", "合格平台、护栏、安全带、监护、恶劣天气停工"],
        ["物体打击", "高处拆除、桥架配管", "工具防坠、下方警戒、材料限量、禁止上下交叉"],
        ["起重伤害", "柜体、母线、设备吊运", "吊装计划、合格吊具、专职指挥、禁入吊臂下方"],
        ["坍塌/跌落", "沟槽和洞口", "探查、放坡支护、临边防护、排水和复查"],
        ["粉尘和噪声", "墙地面拆除、切割", "硬隔离、湿法、集尘、错峰、个人防护"],
        ["食品污染", "施工粉尘、异味、杂物", "运营隔离、停业窗口、覆盖清洁、消毒确认"],
        ["烫伤/机械伤害", "厨房设备试运行", "厂家指导、护罩、急停、分级试车、禁止带故障运行"],
        ["滑倒/绊倒", "电缆敷设、积水、材料", "通道管理、即时清理、桥架放线区警戒、防滑"],
    ]
    add_table(doc, ["主要危险源", "可能环节", "控制措施"], risks, [3.0, 4.0, 7.0], 8.9)

    add_heading(doc, "十四、环境保护、食堂卫生与成品保护", 1)
    add_paragraphs(doc, [
        "拆除和切割采用湿法或集尘设备，建筑垃圾分类袋装或覆盖，运输车辆不带泥污染道路。废油、化学品容器、焊条头和含油抹布按建设单位分类要求收集，不得进入排水沟。施工废水经沉淀或收集处理，禁止直接排入厨房或雨水系统。",
        "食堂运营区与施工区设置连续硬隔离和封闭门，施工人员不得进入食品加工和储存区。影响食品安全的拆除、涂料、防水、切割和大面积粉尘作业安排在停业或批准窗口。恢复营业前清除粉尘异物，完成清洁、消毒、通风和设备安全检查，由使用单位确认。",
        "甲供设备在开箱后覆盖保护，控制面板、玻璃、门体和不锈钢面不得贴附难清除胶带。已完成防静电不发火地面、地砖、墙面、吊顶和防水区域设置专用通道和垫板，重物搬运不得直接碾压。电缆端头临时密封，盘柜和母线防潮防尘。",
        "屋面防水完成后限制上人，必须交叉作业时铺设保护板。墙洞、管根和电缆穿墙处完成防水防火封堵后挂牌，后续施工不得擅自拆除。",
    ])

    add_heading(doc, "十五、劳动力、机械和材料计划", 1)
    resources = [
        ["项目管理", "项目经理、技术负责人、施工员、安全员、质量员、材料员、资料员", "7人", "全工期"],
        ["燃气停用拆除", "能源接口、持证管道人员、检测和监护人员", "约6人", "8月24日-9月10日"],
        ["电气安装", "盘柜母线、桥架配管、电缆和接地人员", "8-10人", "9月12日-11月20日"],
        ["土建装饰", "拆除、钢筋模板、混凝土、砌筑、防水、地坪和装饰人员", "10-12人", "8月17日-11月30日"],
        ["厨房设备", "设备安装及末端接线人员", "约6人", "10月25日-11月20日"],
        ["检测调试", "电气试验、设备调试及使用单位配合人员", "4-6人", "11月21日-12月22日"],
        ["起重运输", "25t汽车吊、10t平板车、12m登高平台", "2/2/5台班", "按吊装计划"],
        ["燃气器材", "氮气瓶组、汇流排、减压器、软管、盲板、气体检测仪", "专项配置", "置换拆除阶段"],
        ["电气仪表", "绝缘表、接地表、相序表、钳表、扭矩扳手、压接工具", "两作业面", "安装调试阶段"],
        ["安全环保", "围挡、灭火器、防火毯、集尘、警示和个人防护用品", "全场配置", "全工期"],
    ]
    add_table(doc, ["类别", "资源内容", "计划投入", "主要阶段"], resources, [2.4, 7.4, 2.2, 3.0], 9)
    add_paragraphs(doc, [
        "主要材料和甲供设备实行到货里程碑管理。配电柜、母线、电缆和厨房设备技术确认后再排产或供货；桥架、钢管、接地材料、混凝土、防水、地面和装饰材料分区分批到场，减少食堂周边堆放和二次搬运。",
        "材料进场提供合格证、检验报告和规格清单。防水、密封、防火封堵和不发火防静电地面材料按规范及合同要求复验或提供性能证明。易受潮的电气设备和材料存放在干燥封闭区域，盘柜及母线端口封堵。",
    ])

    add_heading(doc, "十六、试验、调试、资料与验收", 1)
    add_heading(doc, "16.1 电气试验及送电条件", 2)
    add_paragraphs(doc, [
        "送电前完成盘柜和母线外观、机械操作、接线、相序、绝缘电阻、接地电阻、保护导体连续性、回路核对和保护整定检查。所有施工人员撤离，临时接地和工具清点，柜门盖板复位，送电区设警戒。",
        "送电按上级电源、主柜、分柜、支路和设备逐级进行，每级观察电压、相序、异常声响、气味和温升。发生保护动作不得强行合闸，查明原因并完成记录后再决定。",
    ])
    add_heading(doc, "16.2 资料清单", 2)
    records = [
        ["技术管理", "施工方案、图纸会审、技术核定、施工交底、作业面交接"],
        ["安全管理", "安全教育、专项交底、停气、盲板、置换、动火、停送电、吊装、高处和临电记录"],
        ["材料设备", "材料报验、合格证、复验、甲供设备交接、开箱记录"],
        ["燃气作业", "介质确认、隔断图、盲板台账、置换计算、检测记录、拆除交接"],
        ["电气隐蔽", "桥架配管、接地、穿墙、防火封堵和电缆敷设记录"],
        ["电气试验", "绝缘、接地、相序、连续性、保护动作、盘柜母线检查"],
        ["土建装饰", "钢筋模板、混凝土、屋面防水、地面、墙面、吊顶、门窗和排水验收"],
        ["设备调试", "设备安装、单机试运、负荷试运、培训和问题销项"],
        ["竣工移交", "竣工图、工程照片、验收记录、操作维护资料和移交清单"],
    ]
    add_table(doc, ["资料类别", "主要文件"], records, [3.0, 11.0], 9.3)

    add_heading(doc, "十七、应急预案与事故处置", 1)
    add_paragraphs(doc, [
        "项目建立现场应急组织，由项目经理任现场总指挥，安全员负责报警、警戒和人员疏散，技术负责人负责能源状态和技术处置，各班组长清点人员。建设单位应急电话、消防、医疗、能源调度和项目管理人员电话在开工前填写并张贴在现场。",
        "发生燃气报警、异味或人员不适时，立即停止作业，严禁启闭可能产生火花的电器，人员沿上风向撤离，扩大警戒并报告建设单位能源管理部门。未经检测和批准不得返回。人员疑似中毒或窒息时，救援人员不得无防护盲目进入，须在专业救援条件下施救。",
        "发生火灾时立即停止能源作业、报警并使用适用灭火器材扑救初起火灾；火势无法控制时迅速撤离。电气设备未确认断电前不得用水直接扑救。发生触电时先切断电源或使用绝缘工具脱离电源，再进行心肺复苏并呼叫医疗救援。",
        "发生高处坠落、起重伤害或坍塌时，立即停止相关设备，保护现场，防止二次伤害。未经评估不得移动重伤人员或擅自进入不稳定区域。事故后按建设单位和法规要求报告、调查和整改。",
    ])
    emergency = [
        ["建设单位应急/调度", "待开工前填写", "燃气、电力及厂区应急协调"],
        ["消防报警", "119", "火灾报警"],
        ["医疗急救", "120", "人员伤害急救"],
        ["项目经理", "待填写", "现场总指挥"],
        ["安全员", "待填写", "警戒、疏散和报告"],
        ["技术负责人", "待填写", "能源状态和技术处置"],
    ]
    add_table(doc, ["联系人/单位", "电话", "职责"], emergency, [4.0, 3.2, 6.8], 9.5)

    add_heading(doc, "十八、雨季、高温及冬季施工措施", 1)
    add_paragraphs(doc, [
        "施工跨越武汉夏季、秋季和初冬。8月至9月重点防高温、暴雨和雷电，屋面、沟槽、吊装和登高作业关注天气预警。雨前覆盖电气设备和防水材料，检查沟槽排水、临电防雨和脚手基础；暴雨雷电时停止露天高处、吊装和电气试验。",
        "高温时调整作息，配置饮水、防暑药品和休息点。燃气监护、盘柜内接线、屋面防水等作业避免长时间连续暴露，发现中暑症状立即转移处理。易燃材料远离热源，不在车厢或阳光下密闭存放。",
        "11月至12月气温下降后，混凝土、砂浆、涂料、防水和地面材料施工温度符合产品和规范要求。低温条件下采取保温养护，禁止基层结冰或受冻后直接施工。送电前检查电气设备是否受潮凝露，必要时干燥后复测绝缘。",
    ])

    add_heading(doc, "十九、工程移交与保修", 1)
    add_paragraphs(doc, [
        "工程移交前组织燃气遗留接口、电气系统、厨房设备、屋面防水、地面、墙面、排水和现场卫生全专业联合巡查。燃气管道遗留封堵和标识、电气回路编号、盘柜警示、等电位端子、设备操作说明及应急停机方式向使用单位现场交底。",
        "竣工验收完成后移交竣工图、试验报告、设备资料、保修清单和问题整改记录。保修期内对安装质量问题及时响应；设备本体质量、甲供范围和使用维护责任按合同界面执行。未经许可的改接、超负荷使用或拆除保护接地不属于正常保修处置，应及时报告管理单位。",
        f"{COMPANY}将按照本方案和经批准的施工文件组织实施，确保燃气退出过程可控、电气系统安全可靠、土建恢复完整、食堂环境清洁，力争于{FINISH}前完成联合验收和移交。",
    ])

    # Landscape appendix with a document-optimized monthly control chart. The
    # companion XLSX retains the detailed 21-week view.
    # WPS automatically begins a new page when a continuous section changes
    # paper orientation; an extra manual page break would create a blank page.
    landscape = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(landscape, landscape=True)
    clear_header_footer(landscape, PROJECT + "施工进度横道图")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run()
    r.add_picture(str(SHEET_RENDERS / "文档版横道图.png"), width=Cm(22.2))

    doc.core_properties.title = PROJECT + "施工方案"
    doc.core_properties.subject = "燃气停用拆除、低压配电及厨房设备气改电施工"
    doc.core_properties.author = COMPANY
    doc.core_properties.comments = "依据工程量清单及2026年7月31日规范查新编制"
    doc.save(OUTPUT)
    return doc


if __name__ == "__main__":
    build_document()
    with zipfile.ZipFile(OUTPUT) as zf:
        assert "word/document.xml" in zf.namelist()
        assert "word/styles.xml" in zf.namelist()
    print(OUTPUT)
