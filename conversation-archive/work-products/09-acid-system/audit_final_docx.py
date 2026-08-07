from collections import Counter
from pathlib import Path

from docx import Document


DOCX = Path(r"C:\Users\Administrator\Desktop\101机组酸循环系统和酸雾系统综合改造\101机组酸循环系统和酸雾系统综合改造施工方案.docx")
doc = Document(DOCX)

paragraph_text = "\n".join(p.text for p in doc.paragraphs)
table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
all_text = paragraph_text + "\n" + table_text

required = [
    "101机组酸循环系统和酸雾系统综合改造",
    "2026年8月7日",
    "2026年12月31日",
    "PPH",
    "危险源辨识",
    "应急预案",
    "施工进度计划",
]
legacy = ["质检中心食堂", "气改电", "燃气管道", "厨房设备", "GGD"]

font_names = Counter()
font_sizes = Counter()
run_count = 0
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if not run.text:
            continue
        run_count += 1
        rpr = run._r.rPr
        if rpr is not None and rpr.rFonts is not None:
            names = tuple(
                rpr.rFonts.get(key)
                for key in (
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii",
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
                )
            )
            font_names[names] += 1
        if run.font.size:
            font_sizes[round(run.font.size.pt, 2)] += 1

print(f"path={DOCX}")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)} inline_shapes={len(doc.inline_shapes)}")
print(f"characters_no_space={sum(1 for c in all_text if not c.isspace())}")
print("required=" + repr({term: term in all_text for term in required}))
print("legacy=" + repr({term: term in all_text for term in legacy}))
print(f"nonempty_runs={run_count}")
print("font_names=" + repr(font_names.most_common(8)))
print("font_sizes=" + repr(sorted(font_sizes.items())))
for i, section in enumerate(doc.sections, 1):
    print(
        f"section_{i}=page({section.page_width.inches:.2f}x{section.page_height.inches:.2f}) "
        f"margins({section.left_margin.inches:.2f},{section.right_margin.inches:.2f},"
        f"{section.top_margin.inches:.2f},{section.bottom_margin.inches:.2f})"
    )
