#!/usr/bin/env python3
"""Validate the paired DOCX/WPS and five-sheet XLSX construction-bid delivery."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


REQUIRED_SHEETS = ["横道进度计划", "网络计划逻辑", "资源计划", "评标响应索引", "清单工程量摘要"]
FORMULA_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def valid_zip(path: Path) -> str | None:
    if not path.is_file():
        return f"file not found: {path}"
    if not zipfile.is_zipfile(path):
        return f"not a valid OOXML ZIP package: {path}"
    with zipfile.ZipFile(path) as package:
        bad = package.testzip()
        if bad:
            return f"corrupt ZIP member {bad}: {path}"
    return None


def doc_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def scan_residues(text: str, stale_terms: list[str]) -> dict[str, int]:
    return {term: text.count(term) for term in stale_terms if term and term in text}


def numeric_day(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer() and 1 <= int(value) <= 10000:
        return int(value)
    if isinstance(value, str):
        match = re.fullmatch(r"\s*(?:D|第)?\s*(\d+)\s*(?:天|日)?\s*", value, re.IGNORECASE)
        if match:
            return int(match.group(1))
    return None


def all_values(workbook) -> list[str]:
    values: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value))
    return values


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", required=True)
    parser.add_argument("--days", required=True, type=int)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--xlsx", required=True, type=Path)
    parser.add_argument("--wps", type=Path, help="Optional WPS-format copy; only existence/size is checked")
    parser.add_argument("--template", type=Path)
    parser.add_argument("--expected-template-sha256")
    parser.add_argument("--expected-boq-items", type=int)
    parser.add_argument("--stale", action="append", default=[])
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    failures: list[str] = []
    warnings: list[str] = []
    report: dict = {"project": args.project, "days": args.days}

    docx = args.docx.expanduser().resolve()
    xlsx = args.xlsx.expanduser().resolve()
    for package in (docx, xlsx):
        error = valid_zip(package)
        if error:
            failures.append(error)

    document = None
    if not any(str(docx) in failure for failure in failures):
        try:
            document = Document(docx)
        except Exception as exc:  # pragma: no cover - diagnostic boundary
            failures.append(f"DOCX open failed: {exc}")
    if document is not None:
        text = doc_text(document)
        report["docx"] = {
            "path": str(docx),
            "sha256": sha256(docx),
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "sections": len(document.sections),
            "inline_shapes": len(document.inline_shapes),
            "stale_hits": scan_residues(text, args.stale),
        }
        if args.project not in text:
            failures.append("project name not found in DOCX")
        if report["docx"]["stale_hits"]:
            failures.append(f"stale terms found in DOCX: {report['docx']['stale_hits']}")
        duration_patterns = [rf"(?<!\d){args.days}\s*日历天", rf"(?<!\d){args.days}\s*天(?!\d)", rf"D\s*{args.days}(?!\d)"]
        if not any(re.search(pattern, text, re.IGNORECASE) for pattern in duration_patterns):
            failures.append(f"duration/final day {args.days} not found in DOCX")

        if args.template:
            template = args.template.expanduser().resolve()
            if not template.is_file():
                failures.append(f"template not found: {template}")
            else:
                template_hash = sha256(template)
                template_doc = Document(template)
                report["template"] = {
                    "path": str(template),
                    "sha256": template_hash,
                    "paragraphs": len(template_doc.paragraphs),
                    "tables": len(template_doc.tables),
                    "sections": len(template_doc.sections),
                }
                if args.expected_template_sha256 and template_hash != args.expected_template_sha256.upper():
                    failures.append("template SHA-256 differs from the expected fixed-template fingerprint")
                for key in ("paragraphs", "tables", "sections"):
                    if report["docx"][key] != report["template"][key]:
                        failures.append(
                            f"DOCX {key} {report['docx'][key]} != template {report['template'][key]}"
                        )

    workbook = None
    if not any(str(xlsx) in failure for failure in failures):
        try:
            workbook = load_workbook(xlsx, data_only=False, read_only=False)
        except Exception as exc:  # pragma: no cover - diagnostic boundary
            failures.append(f"XLSX open failed: {exc}")
    if workbook is not None:
        report["xlsx"] = {"path": str(xlsx), "sha256": sha256(xlsx), "sheets": workbook.sheetnames}
        if workbook.sheetnames != REQUIRED_SHEETS:
            failures.append(f"sheet names/order mismatch: {workbook.sheetnames}")
        hidden = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state != "visible"]
        if hidden:
            failures.append(f"hidden sheets found: {hidden}")
        values = all_values(workbook)
        joined = "\n".join(values)
        if args.project not in joined:
            failures.append("project name not found in XLSX")
        stale_hits = scan_residues(joined, args.stale)
        report["xlsx"]["stale_hits"] = stale_hits
        if stale_hits:
            failures.append(f"stale terms found in XLSX: {stale_hits}")
        formula_errors = sorted({value for value in values if value in FORMULA_ERRORS})
        report["xlsx"]["formula_errors"] = formula_errors
        if formula_errors:
            failures.append(f"formula errors found: {formula_errors}")

        if "横道进度计划" in workbook.sheetnames:
            gantt = workbook["横道进度计划"]
            header_days = []
            for row in gantt.iter_rows(min_row=1, max_row=min(15, gantt.max_row)):
                for cell in row:
                    day = numeric_day(cell.value)
                    if day is not None:
                        header_days.append(day)
            last_day = max(header_days) if header_days else None
            milestone_days = []
            for row in gantt.iter_rows():
                row_values = [cell.value for cell in row]
                if any("里程碑" in str(value) or "竣工" in str(value) or "完成验收" in str(value) for value in row_values if value is not None):
                    milestone_days.extend(day for day in (numeric_day(value) for value in row_values) if day is not None)
            report["xlsx"]["gantt_last_day"] = last_day
            report["xlsx"]["milestone_days"] = sorted(set(milestone_days))
            if last_day != args.days:
                failures.append(f"Gantt last day {last_day} != expected {args.days}")
            if args.days not in milestone_days:
                failures.append(f"completion milestone does not include day {args.days}")

        if args.expected_boq_items is not None and "清单工程量摘要" in workbook.sheetnames:
            boq = workbook["清单工程量摘要"]
            serials = []
            for row in boq.iter_rows():
                for cell in row[:4]:
                    if isinstance(cell.value, int) and 1 <= cell.value <= args.expected_boq_items * 5:
                        serials.append(cell.value)
                        break
            unique_serials = sorted(set(serials))
            report["xlsx"]["boq_serials_detected"] = unique_serials
            if len(unique_serials) != args.expected_boq_items:
                failures.append(
                    f"BOQ item count {len(unique_serials)} != expected {args.expected_boq_items}; review serial-column layout"
                )

    if args.wps:
        wps = args.wps.expanduser().resolve()
        report["wps"] = {"path": str(wps), "exists": wps.is_file(), "bytes": wps.stat().st_size if wps.is_file() else 0}
        if not wps.is_file() or wps.stat().st_size == 0:
            failures.append(f"WPS copy missing or empty: {wps}")

    report["warnings"] = warnings
    report["failures"] = sorted(set(failures))
    report["passed"] = not failures
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    raise SystemExit(0 if report["passed"] else 1)


if __name__ == "__main__":
    main()
