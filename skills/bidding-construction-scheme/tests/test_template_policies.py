"""Behavioral regression tests using generated fixtures, not LFS assets."""

import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import Workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
SPEC = importlib.util.spec_from_file_location("bidding_template_fidelity", SCRIPTS / "audit_template_fidelity.py")
FIDELITY = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(FIDELITY)


def run_script(script, *arguments):
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    return subprocess.run(
        [sys.executable, str(SCRIPTS / script), *(str(argument) for argument in arguments)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=env,
        check=False,
    )


class FigureAnchorPolicyTests(unittest.TestCase):
    def test_legacy_fingerprint_enables_legacy_anchors(self):
        result = FIDELITY.resolve_allowed_body(FIDELITY.LEGACY_TEMPLATE_SHA256.lower(), {7})
        self.assertEqual(result, FIDELITY.DEFAULT_FIGURE_PARAGRAPHS | {7})

    def test_unknown_fingerprint_keeps_only_explicit_anchors(self):
        explicit = {7}
        self.assertEqual(FIDELITY.resolve_allowed_body("0" * 64, explicit), {7})
        self.assertEqual(explicit, {7})

    def test_legacy_defaults_can_be_disabled_without_losing_explicit_anchors(self):
        result = FIDELITY.resolve_allowed_body(FIDELITY.LEGACY_TEMPLATE_SHA256, {7}, True)
        self.assertEqual(result, {7})

    def test_unknown_template_has_no_silent_legacy_format_exemption(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "template.docx"
            output = Path(directory) / "output.docx"
            document = Document()
            for index in range(212):
                document.add_paragraph().add_run(f"Paragraph {index}")
            document.save(template)
            document.paragraphs[210].runs[0].bold = True
            document.save(output)

            result = run_script("audit_template_fidelity.py", template, output)
            self.assertEqual(result.returncode, 1, result.stderr)
            report = json.loads(result.stdout)
            self.assertFalse(report["passed"])
            self.assertFalse(report["default_figure_anchors_applied"])
            self.assertEqual(report["allowed_body_paragraphs"], [])
            self.assertEqual(report["template_sha256"], hashlib.sha256(template.read_bytes()).hexdigest().upper())
            self.assertTrue(any(item["location"] == "body paragraph 210" for item in report["differences"]))

            result = run_script("audit_template_fidelity.py", template, output, "--allow-body-paragraph", 210)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(result.stdout)
            self.assertTrue(report["passed"])
            self.assertEqual(report["allowed_body_paragraphs"], [210])
            self.assertEqual(report["explicit_allowed_body_paragraphs"], [210])
            self.assertFalse(report["default_figure_anchors_applied"])


class PairedStructurePolicyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.template = self.directory / "template.docx"
        self.output = self.directory / "scheme.docx"
        self.xlsx = self.directory / "schedule.xlsx"
        self.project = "格式回归测试项目"
        document = Document()
        document.add_paragraph().add_run(f"{self.project}，工期76天。")
        document.save(self.template)
        document.add_paragraph().add_run("新增施工工艺。")
        document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("新增附表")
        document.add_section()
        document.save(self.output)
        self.template_hash = hashlib.sha256(self.template.read_bytes()).hexdigest().upper()
        self.write_workbook(76)

    def tearDown(self):
        self.temp.cleanup()

    def write_workbook(self, last_day):
        workbook = Workbook()
        workbook.active.title = "横道进度计划"
        for name in ("网络计划逻辑", "资源计划", "评标响应索引", "清单工程量摘要"):
            workbook.create_sheet(name)
        gantt = workbook["横道进度计划"]
        gantt["A1"] = self.project
        gantt["A2"] = "D1"
        gantt["B2"] = f"D{last_day}"
        gantt["A3"] = "竣工里程碑"
        gantt["B3"] = last_day
        workbook.save(self.xlsx)
        workbook.close()

    def run_validator(self, *options, with_template=True, docx=None):
        arguments = ["--project", self.project, "--days", 76, "--docx", docx or self.output, "--xlsx", self.xlsx]
        if with_template:
            arguments.extend(["--template", self.template])
        return run_script("validate_paired_outputs.py", *arguments, *options)

    def test_default_fixed_policy_still_rejects_count_changes(self):
        result = self.run_validator()
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report["structure_policy"], "fixed")
        self.assertFalse(report["passed"])
        for key in ("paragraphs", "tables", "sections"):
            self.assertTrue(any(f"DOCX {key}" in failure for failure in report["failures"]))

    def test_fixed_policy_accepts_unchanged_structure(self):
        result = self.run_validator("--structure-policy", "fixed", docx=self.template)
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertTrue(report["passed"])
        self.assertEqual(report["template"]["count_deltas"], {"paragraphs": 0, "tables": 0, "sections": 0})

    def test_expandable_records_deltas_and_requires_separate_format_audit(self):
        result = self.run_validator("--structure-policy", "expandable", "--expected-template-sha256", self.template_hash.lower())
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertTrue(report["passed"])
        self.assertTrue(report["profile_format_audit_required"])
        self.assertTrue(report["template"]["fingerprint_verified"])
        self.assertEqual(report["template"]["sha256"], self.template_hash)
        self.assertEqual(report["template"]["count_deltas"], {"paragraphs": 2, "tables": 1, "sections": 1})
        self.assertTrue(any("not a formatting approval" in warning for warning in report["warnings"]))

    def test_expandable_rejects_missing_template_argument(self):
        result = self.run_validator("--structure-policy", "expandable", with_template=False)
        self.assertEqual(result.returncode, 2)
        self.assertIn("requires --template", result.stderr)

    def test_expandable_rejects_incorrect_template_hash(self):
        result = self.run_validator("--structure-policy", "expandable", "--expected-template-sha256", "0" * 64)
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads(result.stdout)
        self.assertFalse(report["template"]["fingerprint_verified"])
        self.assertTrue(any("SHA-256 differs" in failure for failure in report["failures"]))

    def test_expandable_reports_unverified_fingerprint_when_expected_hash_omitted(self):
        result = self.run_validator("--structure-policy", "expandable")
        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertFalse(report["template"]["fingerprint_verified"])
        self.assertTrue(any("not verified" in warning for warning in report["warnings"]))

    def test_expandable_does_not_bypass_schedule_checks(self):
        self.write_workbook(75)
        result = self.run_validator("--structure-policy", "expandable", "--expected-template-sha256", self.template_hash)
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads(result.stdout)
        self.assertTrue(any("Gantt last day 75 != expected 76" in failure for failure in report["failures"]))


if __name__ == "__main__":
    unittest.main()
