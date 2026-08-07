from collections import Counter
from pathlib import Path
import os

from docx import Document


doc = Document(Path(os.environ["SOURCE_DOCX"]))
size_counter = Counter()
font_counter = Counter()
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.text:
            size_counter[str(run.font.size.pt if run.font.size else None)] += len(run.text)
            font_counter[str(run.font.name)] += len(run.text)

print("sizes", size_counter.most_common(12))
print("fonts", font_counter.most_common(12))
for index in [0, 1, 4, 8, 15, 36, 37, 60, 61, 100, 108, 132, 257, 365, 437, 579, 596, 603, 605]:
    p = doc.paragraphs[index]
    run_evidence = [(r.text[:30], r.font.name, r.font.size.pt if r.font.size else None, r.bold) for r in p.runs if r.text]
    print(index, p.style.name, repr(p.text), run_evidence)
