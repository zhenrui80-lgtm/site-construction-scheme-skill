#!/usr/bin/env python3
import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def run_font_names(run):
    rpr = run._r.rPr
    if rpr is None or rpr.rFonts is None:
        return []
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return [rpr.rFonts.get(namespace + key) for key in ("ascii", "hAnsi", "eastAsia", "cs") if rpr.rFonts.get(namespace + key)]


def style_font_names(paragraph):
    style = paragraph.style
    if style is None:
        return []
    rpr = style.element.rPr
    if rpr is None or rpr.rFonts is None:
        return [style.font.name] if style.font.name else []
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    names = [rpr.rFonts.get(namespace + key) for key in ("ascii", "hAnsi", "eastAsia", "cs") if rpr.rFonts.get(namespace + key)]
    return names or ([style.font.name] if style.font.name else [])


def ordered_body_paragraphs(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs


def paragraphs_after_marker(document, marker, occurrence):
    hits = 0
    found = False
    result = []
    for paragraph in ordered_body_paragraphs(document):
        if marker in paragraph.text:
            hits += 1
            if hits == occurrence:
                found = True
        if found:
            result.append(paragraph)
    if not found:
        raise ValueError(f"Body marker not found at occurrence {occurrence}: {marker}")
    return result


def main():
    parser = argparse.ArgumentParser(description="Audit content, fonts, sections, images, and DOCX integrity.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--required", action="append", default=[])
    parser.add_argument("--forbidden", action="append", default=[])
    parser.add_argument("--expected-font")
    parser.add_argument("--expected-body-font")
    parser.add_argument("--allowed-font", action="append", default=[])
    parser.add_argument("--body-start")
    parser.add_argument("--occurrence", type=int, default=2)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    path = args.docx.expanduser().resolve()
    document = Document(path)
    paragraphs = list(all_paragraphs(document))
    if args.expected_body_font:
        if not args.body_start:
            parser.error("--expected-body-font requires --body-start")
        font_paragraphs = paragraphs_after_marker(document, args.body_start, args.occurrence)
        expected_font = args.expected_body_font
        font_scope = "body"
    else:
        font_paragraphs = paragraphs
        expected_font = args.expected_font
        font_scope = "whole-document"
    text = "\n".join(paragraph.text for paragraph in paragraphs)

    font_names = Counter()
    unexpected_fonts = Counter()
    sizes = Counter()
    nonempty_runs = 0
    unresolved_font_runs = 0
    allowed_fonts = set(args.allowed_font)
    if expected_font:
        allowed_fonts.add(expected_font)
    for paragraph in font_paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            nonempty_runs += 1
            names = run_font_names(run) or style_font_names(paragraph)
            if not names:
                unresolved_font_runs += 1
            for name in names:
                font_names[name] += 1
                if expected_font and name not in allowed_fonts:
                    unexpected_fonts[name] += 1
            if run.font.size:
                sizes[round(run.font.size.pt, 2)] += 1

    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()

    required = {term: term in text for term in args.required}
    forbidden = {term: term in text for term in args.forbidden}
    result = {
        "path": str(path),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "sections": len(document.sections),
        "inline_shapes": len(document.inline_shapes),
        "characters_no_space": sum(1 for character in text if not character.isspace()),
        "nonempty_runs": nonempty_runs,
        "font_scope": font_scope,
        "expected_font": expected_font,
        "allowed_fonts": sorted(allowed_fonts),
        "unresolved_font_runs": unresolved_font_runs,
        "required": required,
        "forbidden": forbidden,
        "font_names": dict(font_names),
        "font_sizes_pt": {str(size): count for size, count in sorted(sizes.items())},
        "unexpected_fonts": dict(unexpected_fonts),
        "zip_bad_member": bad_member,
        "sections_page_inches": [
            {
                "width": round(section.page_width.inches, 2),
                "height": round(section.page_height.inches, 2),
                "left": round(section.left_margin.inches, 2),
                "right": round(section.right_margin.inches, 2),
                "top": round(section.top_margin.inches, 2),
                "bottom": round(section.bottom_margin.inches, 2),
            }
            for section in document.sections
        ],
    }
    failures = []
    failures.extend(f"missing required term: {term}" for term, present in required.items() if not present)
    failures.extend(f"found forbidden term: {term}" for term, present in forbidden.items() if present)
    if unexpected_fonts:
        failures.append(f"unexpected explicit fonts: {dict(unexpected_fonts)}")
    if bad_member:
        failures.append(f"corrupt DOCX member: {bad_member}")
    result["failures"] = failures

    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding="utf-8")
    print(rendered)
    raise SystemExit(1 if failures else 0)


if __name__ == "__main__":
    main()
