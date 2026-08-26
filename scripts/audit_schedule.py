#!/usr/bin/env python3
"""Check that the schedule basis and duration appear consistently in the DOCX."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from docx import Document


def document_text(path):
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return "\n".join(texts)


def date_variants(value):
    parsed = date.fromisoformat(value)
    return {
        value,
        f"{parsed.year}.{parsed.month}.{parsed.day}",
        f"{parsed.year}年{parsed.month}月{parsed.day}日",
        f"{parsed.month}月{parsed.day}日",
    }


def audit(path, basis, duration_days, start_date=None, finish_date=None):
    text = document_text(path)
    failures = []
    result = {"basis": basis, "duration_days": duration_days}
    if duration_days < 1:
        failures.append("duration_days must be positive")
    duration_patterns = [
        rf"(?<!\d){duration_days}\s*日历天",
        rf"(?<!\d){duration_days}\s*天(?!\d)",
        rf"D\s*{duration_days}(?!\d)",
        rf"第\s*{duration_days}\s*日",
    ]
    duration_hits = [pattern for pattern in duration_patterns if re.search(pattern, text, re.IGNORECASE)]
    if not duration_hits:
        failures.append(f"duration or final day {duration_days} not found in DOCX")
    result["duration_patterns_found"] = duration_hits
    if basis == "fixed_dates":
        if not start_date or not finish_date:
            failures.append("fixed_dates requires start_date and finish_date")
        else:
            start = date.fromisoformat(start_date)
            finish = date.fromisoformat(finish_date)
            computed = (finish - start).days + 1
            result["computed_inclusive_days"] = computed
            if computed != duration_days:
                failures.append(f"inclusive duration {computed} != configured {duration_days}")
            start_hits = sorted(value for value in date_variants(start_date) if value in text)
            finish_hits = sorted(value for value in date_variants(finish_date) if value in text)
            result["start_date_hits"] = start_hits
            result["finish_date_hits"] = finish_hits
            if not start_hits:
                failures.append(f"start date {start_date} not found in DOCX")
            if not finish_hits:
                failures.append(f"finish date {finish_date} not found in DOCX")
    elif basis == "day_index":
        day_index_hits = [value for value in (f"D{duration_days}", f"D {duration_days}") if value in text]
        result["final_day_hits"] = day_index_hits
        if not day_index_hits:
            failures.append(f"day-index final day D{duration_days} not found in DOCX")
    else:
        failures.append(f"unsupported schedule basis: {basis}")
    result["failures"] = failures
    return result


def main():
    parser = argparse.ArgumentParser(description="Audit construction-scheme schedule dates and duration.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--basis", choices=("fixed_dates", "day_index"), required=True)
    parser.add_argument("--duration-days", type=int, required=True)
    parser.add_argument("--start-date")
    parser.add_argument("--finish-date")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = audit(
        args.docx.expanduser().resolve(),
        args.basis,
        args.duration_days,
        args.start_date,
        args.finish_date,
    )
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.expanduser().resolve().write_text(rendered, encoding="utf-8")
    print(rendered)
    raise SystemExit(1 if result["failures"] else 0)


if __name__ == "__main__":
    main()
