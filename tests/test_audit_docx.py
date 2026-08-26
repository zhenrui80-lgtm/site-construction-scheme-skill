import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "audit_docx.py"


def set_font(run, name):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


class DocxFontAuditTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / "mixed-font.docx"
        document = Document()
        set_font(document.add_paragraph().add_run("施工方案"), "宋体")
        document.add_paragraph("一、编制依据及工程概况")
        document.add_paragraph("目录占位")
        document.add_paragraph("一、编制依据及工程概况")
        set_font(document.add_paragraph().add_run("正文内容"), "方正仿宋_GB2312")
        document.save(self.path)

    def tearDown(self):
        self.directory.cleanup()

    def run_audit(self, *extra):
        output = Path(self.directory.name) / "result.json"
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(self.path), *extra, "--output", str(output)],
            capture_output=True,
            text=True,
        )
        return result.returncode, json.loads(output.read_text(encoding="utf-8"))

    def test_body_font_ignores_frozen_cover_font(self):
        code, result = self.run_audit(
            "--body-start", "一、编制依据及工程概况",
            "--occurrence", "2",
            "--expected-body-font", "方正仿宋_GB2312",
        )
        self.assertEqual(code, 0, result)
        self.assertEqual(result["font_scope"], "body")

    def test_legacy_whole_document_font_check_still_fails(self):
        code, result = self.run_audit("--expected-font", "方正仿宋_GB2312")
        self.assertNotEqual(code, 0)
        self.assertIn("宋体", result["unexpected_fonts"])


if __name__ == "__main__":
    unittest.main()
