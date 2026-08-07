from pathlib import Path
import os

from docx import Document


source = Path(os.environ["SOURCE_DOCX"])
output = Path(__file__).with_name("source_content.md")
doc = Document(source)

lines = [f"# {source.name}", "", "## Paragraphs", ""]
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.replace("\t", " ").strip()
    if text:
        lines.append(f"- P{index} [{paragraph.style.name}]: {text}")

lines.extend(["", "## Tables", ""])
for table_index, table in enumerate(doc.tables):
    lines.append(f"### Table {table_index}")
    for row_index, row in enumerate(table.rows):
        values = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
        lines.append(f"- R{row_index}: " + " | ".join(values))
    lines.append("")

output.write_text("\n".join(lines), encoding="utf-8")
print(output)
