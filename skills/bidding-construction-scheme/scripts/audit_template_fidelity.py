#!/usr/bin/env python3
"""Compare DOCX formatting structure while allowing run-level text replacement."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


DEFAULT_FIGURE_PARAGRAPHS = {210, 211, 247, 248, 258, 259, 270, 271, 279, 280, 288, 289}


def scalar(value):
    if value is None:
        return None
    if hasattr(value, "emu"):
        return int(value.emu)
    if hasattr(value, "value"):
        return value.value
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def paragraph_format_signature(paragraph) -> dict:
    fmt = paragraph.paragraph_format
    tabs = [
        (scalar(stop.position), scalar(stop.alignment), scalar(stop.leader))
        for stop in fmt.tab_stops
    ]
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    return {
        "style": paragraph.style.style_id if paragraph.style is not None else None,
        "alignment": scalar(paragraph.alignment),
        "left_indent": scalar(fmt.left_indent),
        "right_indent": scalar(fmt.right_indent),
        "first_line_indent": scalar(fmt.first_line_indent),
        "space_before": scalar(fmt.space_before),
        "space_after": scalar(fmt.space_after),
        "line_spacing": scalar(fmt.line_spacing),
        "line_spacing_rule": scalar(fmt.line_spacing_rule),
        "keep_together": fmt.keep_together,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "widow_control": fmt.widow_control,
        "tabs": tabs,
        "num_id": scalar(num_pr.numId.val) if num_pr is not None and num_pr.numId is not None else None,
        "num_level": scalar(num_pr.ilvl.val) if num_pr is not None and num_pr.ilvl is not None else None,
    }


def run_format_signature(run) -> tuple:
    font = run.font
    r_fonts = run._r.rPr.rFonts if run._r.rPr is not None else None
    explicit_fonts = tuple(
        r_fonts.get(qn(name)) if r_fonts is not None else None
        for name in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs")
    )
    return (
        run.style.style_id if run.style is not None else None,
        font.name,
        scalar(font.size),
        font.bold,
        font.italic,
        scalar(font.underline),
        font.strike,
        font.double_strike,
        font.all_caps,
        font.small_caps,
        font.subscript,
        font.superscript,
        font.hidden,
        str(font.color.rgb) if font.color is not None and font.color.rgb is not None else None,
        scalar(font.color.theme_color) if font.color is not None else None,
        scalar(font.highlight_color),
        explicit_fonts,
    )


def run_formats_are_preserved(source, final) -> bool:
    source_formats = set(run_format_signature(run) for run in source.runs)
    final_formats = set(run_format_signature(run) for run in final.runs)
    return final_formats.issubset(source_formats)


def table_format_signature(table) -> dict:
    return {
        "style": table.style.style_id if table.style is not None else None,
        "alignment": scalar(table.alignment),
        "autofit": table.autofit,
        "column_widths": [scalar(column.width) for column in table.columns],
        "row_heights": [scalar(row.height) for row in table.rows],
        "row_height_rules": [scalar(row.height_rule) for row in table.rows],
        "cell_widths": [[scalar(cell.width) for cell in row.cells] for row in table.rows],
        "vertical_alignment": [
            [scalar(cell.vertical_alignment) for cell in row.cells] for row in table.rows
        ],
    }


def section_format_signature(section) -> dict:
    return {
        "start_type": scalar(section.start_type),
        "orientation": scalar(section.orientation),
        "page_width": scalar(section.page_width),
        "page_height": scalar(section.page_height),
        "top_margin": scalar(section.top_margin),
        "bottom_margin": scalar(section.bottom_margin),
        "left_margin": scalar(section.left_margin),
        "right_margin": scalar(section.right_margin),
        "header_distance": scalar(section.header_distance),
        "footer_distance": scalar(section.footer_distance),
        "gutter": scalar(section.gutter),
        "different_first_page_header_footer": section.different_first_page_header_footer,
    }


def table_paragraphs(table):
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            for paragraph_index, paragraph in enumerate(cell.paragraphs):
                yield (row_index, column_index, paragraph_index), paragraph


def compare(template_path: Path, output_path: Path, allowed_body: set[int]) -> dict:
    template = Document(template_path)
    output = Document(output_path)
    failures: list[str] = []
    differences: list[dict] = []

    counts = {
        "template": {
            "paragraphs": len(template.paragraphs),
            "tables": len(template.tables),
            "sections": len(template.sections),
        },
        "output": {
            "paragraphs": len(output.paragraphs),
            "tables": len(output.tables),
            "sections": len(output.sections),
        },
    }
    if counts["template"] != counts["output"]:
        failures.append(f"top-level structure mismatch: {counts}")

    for index, (source, final) in enumerate(zip(template.paragraphs, output.paragraphs)):
        if index not in allowed_body and paragraph_format_signature(source) != paragraph_format_signature(final):
            differences.append({"location": f"body paragraph {index}", "property": "paragraph format"})
        if index not in allowed_body and not run_formats_are_preserved(source, final):
            differences.append({"location": f"body paragraph {index}", "property": "new run formatting"})

    for table_index, (source_table, final_table) in enumerate(zip(template.tables, output.tables)):
        if len(source_table.rows) != len(final_table.rows) or len(source_table.columns) != len(final_table.columns):
            failures.append(
                f"table {table_index} shape mismatch: "
                f"{len(source_table.rows)}x{len(source_table.columns)} != "
                f"{len(final_table.rows)}x{len(final_table.columns)}"
            )
            continue
        if table_format_signature(source_table) != table_format_signature(final_table):
            differences.append({"location": f"table {table_index}", "property": "table format"})
        source_paragraphs = list(table_paragraphs(source_table))
        final_paragraphs = list(table_paragraphs(final_table))
        if len(source_paragraphs) != len(final_paragraphs):
            failures.append(f"table {table_index} paragraph count mismatch")
            continue
        for (location, source_paragraph), (_, final_paragraph) in zip(source_paragraphs, final_paragraphs):
            paragraph_changed = paragraph_format_signature(source_paragraph) != paragraph_format_signature(final_paragraph)
            run_changed = not run_formats_are_preserved(source_paragraph, final_paragraph)
            if paragraph_changed or run_changed:
                differences.append(
                    {
                        "location": f"table {table_index} cell {location}",
                        "property": "paragraph or run format",
                    }
                )

    for index, (source_section, final_section) in enumerate(zip(template.sections, output.sections)):
        if section_format_signature(source_section) != section_format_signature(final_section):
            differences.append({"location": f"section {index}", "property": "section format"})

    if differences:
        failures.append(f"format/property differences found: {len(differences)}")
    return {
        "template": str(template_path),
        "output": str(output_path),
        "allowed_body_paragraphs": sorted(allowed_body),
        "counts": counts,
        "differences": differences,
        "failures": failures,
        "passed": not failures,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--allow-body-paragraph",
        action="append",
        type=int,
        default=[],
        help="Zero-based body paragraph allowed to change paragraph/run formatting, normally a figure or caption slot",
    )
    parser.add_argument("--no-default-figure-anchors", action="store_true")
    parser.add_argument("--output", type=Path, dest="report_path")
    args = parser.parse_args()

    allowed = set(args.allow_body_paragraph)
    if not args.no_default_figure_anchors:
        allowed.update(DEFAULT_FIGURE_PARAGRAPHS)
    template = args.template.expanduser().resolve()
    output = args.output.expanduser().resolve()
    if not template.is_file() or not output.is_file():
        raise SystemExit("template and output DOCX files must exist")
    result = compare(template, output, allowed)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report_path:
        args.report_path.parent.mkdir(parents=True, exist_ok=True)
        args.report_path.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    raise SystemExit(0 if result["passed"] else 1)


if __name__ == "__main__":
    main()
