#!/usr/bin/env python3
import argparse
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document


PROFILES = {
    "small": {
        "body_chars": 22000,
        "narrative_chars": 18000,
        "narrative_paragraphs": 180,
        "heading_candidates": 50,
        "major_chapters": 10,
        "management_topics": 10,
        "project_terms": 4,
        "method_terms": 5,
        "narrative_ratio": 0.78,
        "tables_per_10k_chars": 5.0,
        "table_rows_per_10k_chars": 45.0,
        "prose_heavy_tables": 1,
        "chars_per_media": 1500,
        "chars_per_standard_excerpt": 3500,
        "method_detail_ratio": 0.60,
        "project_chapter_spread": 1.5,
        "duplicate_ratio": 0.05,
    },
    "standard": {
        "body_chars": 35000,
        "narrative_chars": 30000,
        "narrative_paragraphs": 280,
        "heading_candidates": 85,
        "major_chapters": 14,
        "management_topics": 12,
        "project_terms": 6,
        "method_terms": 8,
        "narrative_ratio": 0.82,
        "tables_per_10k_chars": 4.0,
        "table_rows_per_10k_chars": 35.0,
        "prose_heavy_tables": 2,
        "chars_per_media": 1800,
        "chars_per_standard_excerpt": 4000,
        "method_detail_ratio": 0.70,
        "project_chapter_spread": 2.0,
        "duplicate_ratio": 0.03,
    },
    "complex": {
        "body_chars": 48000,
        "narrative_chars": 40000,
        "narrative_paragraphs": 360,
        "heading_candidates": 120,
        "major_chapters": 15,
        "management_topics": 14,
        "project_terms": 8,
        "method_terms": 12,
        "narrative_ratio": 0.82,
        "tables_per_10k_chars": 4.0,
        "table_rows_per_10k_chars": 35.0,
        "prose_heavy_tables": 3,
        "chars_per_media": 1800,
        "chars_per_standard_excerpt": 4000,
        "method_detail_ratio": 0.80,
        "project_chapter_spread": 2.5,
        "duplicate_ratio": 0.02,
    },
}

MANAGEMENT_TOPICS = {
    "编制依据": ("编制依据",),
    "工程概况": ("工程概况", "项目概况"),
    "施工组织": ("施工组织", "组织机构", "项目组织"),
    "施工准备": ("施工准备", "开工准备"),
    "总平面临设": ("施工总平面", "总平面布置", "临时设施", "临时用电", "临时用水"),
    "进度计划": ("进度计划", "施工进度", "横道图"),
    "工期保证": ("工期保证", "进度保证", "赶工措施"),
    "专项工法": ("施工方法", "施工工艺", "专项施工", "主要施工工序"),
    "质量管理": ("质量保证", "质量管理", "技术质量", "质量控制"),
    "安全危险源": ("安全措施", "安全管理", "危险源", "风险辨识"),
    "文明环保职业健康": ("文明施工", "环境保护", "职业健康", "环保措施"),
    "资源计划": ("资源需求", "劳动力计划", "机械设备", "材料计划"),
    "季节特殊条件": ("季节性施工", "雨季施工", "高温施工", "冬季施工", "特殊条件施工"),
    "成品生产保护": ("成品保护", "既有设施保护", "生产系统保护"),
    "应急预案": ("应急预案", "应急处置", "应急响应"),
    "协调配合": ("协调配合", "交叉施工", "各方协调", "施工配合"),
    "试验调试移交": ("试验与调试", "试验、调试", "调试方案", "竣工验收", "资料移交"),
    "保修服务": ("保修服务", "质量保修", "履约承诺"),
}

METHOD_CONTEXT = {
    "preconditions": ("前提", "条件", "交接", "确认", "作业票", "许可", "停机", "隔离"),
    "sequence": ("工序", "步骤", "顺序", "施工", "安装", "拆除", "制作", "连接"),
    "parameters": ("参数", "尺寸", "标高", "偏差", "扭矩", "压力", "厚度", "强度", "图纸", "厂家"),
    "quality": ("质量", "检查", "检验", "试验", "检测", "验收", "见证", "停检点"),
    "safety": ("安全", "防护", "监护", "隔离", "风险", "危险", "停工", "警戒"),
    "records": ("记录", "签字", "报验", "台账", "资料", "追溯", "移交"),
}

HEADING_RE = re.compile(
    r"^\s*(?:第[一二三四五六七八九十百0-9]+[章节]|[一二三四五六七八九十百]{1,3}[、.．]|[0-9]{1,2}(?:\.[0-9]{1,2}){0,3}[.．]?\s+|[（(][一二三四五六七八九十0-9]+[）)])\s*[^。；]{1,80}$"
)
CAPTION_RE = re.compile(r"^\s*(?:图|表)\s*[0-9一二三四五六七八九十]+")
STANDARD_EXCERPT_RE = re.compile(r"(?:GB|GB/T|JGJ|JGJ/T|CJJ|TSG|DL/T|YB/T)\s*\d[^\n]{0,40}(?:原文|摘录)", re.I)


def compact_length(text):
    return sum(1 for character in text if not character.isspace())


def normalized(text):
    return re.sub(r"[\s，。；：、,.!?！？:;（）()\-—_]+", "", text).lower()


def is_heading(paragraph):
    text = paragraph.text.strip()
    style = (paragraph.style.name if paragraph.style else "").lower()
    return bool(text and ("heading" in style or "标题" in style or HEADING_RE.match(text)))


def chinese_number(value):
    digits = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
    if value == "十":
        return 10
    if "十" in value:
        left, right = value.split("十", 1)
        return (digits.get(left, 1) * 10) + digits.get(right, 0)
    return digits.get(value)


def major_heading_number(paragraph):
    text = paragraph.text.strip()
    style = (paragraph.style.name if paragraph.style else "").lower()
    if not text or len(text) > 140:
        return None
    match = re.match(r"^第([一二三四五六七八九十百]+|[0-9]{1,2})章", text)
    if match:
        value = match.group(1)
        return int(value) if value.isdigit() else chinese_number(value)
    match = re.match(r"^([0-9]{1,2})[.．]\s+(?![0-9])", text)
    if match:
        return int(match.group(1))
    match = re.match(r"^([一二三四五六七八九十]{1,3})[、.．]", text)
    if match:
        return chinese_number(match.group(1))
    if style in {"heading 1", "标题 1", "标题1"}:
        match = re.match(r"^([0-9]{1,2})", text)
        return int(match.group(1)) if match else None
    return None


def table_text(document):
    values = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                values.extend(paragraph.text for paragraph in cell.paragraphs if paragraph.text.strip())
    return "\n".join(values)


def table_metrics(document, body_chars):
    row_count = sum(len(table.rows) for table in document.tables)
    prose_heavy = []
    for table_index, table in enumerate(document.tables, start=1):
        cells = []
        for row in table.rows:
            for cell in row.cells:
                value = compact_length(cell.text)
                if value:
                    cells.append(value)
        long_cells = sum(value >= 120 for value in cells)
        average = sum(cells) / max(1, len(cells))
        if long_cells >= 2 or average >= 80:
            prose_heavy.append(
                {
                    "table": table_index,
                    "rows": len(table.rows),
                    "nonempty_cells": len(cells),
                    "long_cells": long_cells,
                    "average_cell_characters": round(average, 1),
                }
            )
    units = max(body_chars / 10000, 0.01)
    return {
        "tables": len(document.tables),
        "rows": row_count,
        "tables_per_10k_body_characters": round(len(document.tables) / units, 2),
        "rows_per_10k_body_characters": round(row_count / units, 2),
        "prose_heavy_tables": prose_heavy,
    }


def media_count(path):
    with zipfile.ZipFile(path) as archive:
        return len({name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")})


def chapter_bins(paragraphs):
    current = 0
    bins = []
    for paragraph in paragraphs:
        number = major_heading_number(paragraph)
        if number:
            current = number
        bins.append(current)
    return bins


def nearby_text(paragraphs, index, radius=5):
    start = max(0, index - radius)
    end = min(len(paragraphs), index + radius + 1)
    return "\n".join(item.text for item in paragraphs[start:end])


def main():
    parser = argparse.ArgumentParser(description="Block shallow or generic construction-scheme DOCX deliverables.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--profile", choices=sorted(PROFILES), default="standard")
    parser.add_argument("--project-term", action="append", default=[])
    parser.add_argument("--method-term", action="append", default=[])
    parser.add_argument("--legacy-term", action="append", default=[])
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    path = args.docx.expanduser().resolve()
    document = Document(path)
    paragraphs = list(document.paragraphs)
    nonempty = [paragraph for paragraph in paragraphs if paragraph.text.strip()]
    heading_flags = [is_heading(paragraph) for paragraph in paragraphs]
    major_numbers = [major_heading_number(paragraph) for paragraph in paragraphs]
    narrative_paragraphs = [
        paragraph.text.strip()
        for index, paragraph in enumerate(paragraphs)
        if paragraph.text.strip() and not heading_flags[index] and not CAPTION_RE.match(paragraph.text.strip())
    ]
    narrative_text = "\n".join(narrative_paragraphs)
    all_body_text = "\n".join(paragraph.text for paragraph in paragraphs)
    tabular_text = table_text(document)
    narrative_chars = compact_length(narrative_text)
    body_chars = compact_length(all_body_text)
    table_chars = compact_length(tabular_text)
    narrative_ratio = narrative_chars / max(1, narrative_chars + table_chars)
    tables = table_metrics(document, body_chars)
    headings = sum(heading_flags)
    major_chapters = max((number for number in major_numbers if number), default=0)
    media = media_count(path)
    effective_media = max(0, media - 1)
    chars_per_media = narrative_chars / max(1, effective_media) if effective_media else None
    standard_excerpts = len(STANDARD_EXCERPT_RE.findall(all_body_text))
    chars_per_standard_excerpt = narrative_chars / max(1, standard_excerpts) if standard_excerpts else None

    topic_hits = {
        topic: [term for term in terms if term in all_body_text or term in tabular_text]
        for topic, terms in MANAGEMENT_TOPICS.items()
    }
    present_topics = {topic: hits for topic, hits in topic_hits.items() if hits}

    bins = chapter_bins(paragraphs)
    project_term_stats = {}
    for term in args.project_term:
        indexes = [index for index, paragraph in enumerate(paragraphs) if term in paragraph.text]
        project_term_stats[term] = {
            "occurrences": all_body_text.count(term) + tabular_text.count(term),
            "chapter_spread": len({bins[index] for index in indexes if bins[index] > 0}),
        }

    method_term_stats = {}
    detailed_methods = 0
    for term in args.method_term:
        indexes = [index for index, paragraph in enumerate(paragraphs) if term in paragraph.text]
        context = "\n".join(nearby_text(paragraphs, index) for index in indexes[:8])
        categories = {
            category: any(keyword in context for keyword in keywords)
            for category, keywords in METHOD_CONTEXT.items()
        }
        detailed = bool(indexes) and sum(categories.values()) >= 4
        detailed_methods += int(detailed)
        method_term_stats[term] = {
            "occurrences": all_body_text.count(term) + tabular_text.count(term),
            "context_categories": categories,
            "detailed": detailed,
        }

    long_paragraphs = [normalized(text) for text in narrative_paragraphs if compact_length(text) >= 80]
    duplicate_counts = Counter(long_paragraphs)
    duplicate_instances = sum(count - 1 for count in duplicate_counts.values() if count > 1)
    duplicate_ratio = duplicate_instances / max(1, len(long_paragraphs))

    thresholds = PROFILES[args.profile]
    average_project_spread = (
        sum(item["chapter_spread"] for item in project_term_stats.values()) / max(1, len(project_term_stats))
    )
    method_detail_ratio = detailed_methods / max(1, len(args.method_term))
    legacy_hits = {term: all_body_text.count(term) + tabular_text.count(term) for term in args.legacy_term if term}

    failures = []

    def minimum(name, actual, expected):
        if actual < expected:
            failures.append(f"{name} below minimum: {actual} < {expected}")

    minimum("body characters", body_chars, thresholds["body_chars"])
    minimum("narrative characters", narrative_chars, thresholds["narrative_chars"])
    minimum("narrative paragraphs", len(narrative_paragraphs), thresholds["narrative_paragraphs"])
    minimum("heading candidates", headings, thresholds["heading_candidates"])
    minimum("major chapters", major_chapters, thresholds["major_chapters"])
    minimum("management topics", len(present_topics), thresholds["management_topics"])
    minimum("project-term inputs", len(args.project_term), thresholds["project_terms"])
    minimum("method-term inputs", len(args.method_term), thresholds["method_terms"])
    minimum("narrative/table character ratio", round(narrative_ratio, 4), thresholds["narrative_ratio"])
    minimum("method detail ratio", round(method_detail_ratio, 4), thresholds["method_detail_ratio"])
    minimum("average project-term chapter spread", round(average_project_spread, 2), thresholds["project_chapter_spread"])

    missing_project_terms = [term for term, item in project_term_stats.items() if item["occurrences"] == 0]
    missing_method_terms = [term for term, item in method_term_stats.items() if item["occurrences"] == 0]
    if missing_project_terms:
        failures.append(f"missing project terms: {missing_project_terms}")
    if missing_method_terms:
        failures.append(f"missing method terms: {missing_method_terms}")
    if chars_per_media is not None and chars_per_media < thresholds["chars_per_media"]:
        failures.append(
            f"image/text balance below minimum: {chars_per_media:.1f} narrative chars per effective media < {thresholds['chars_per_media']}"
        )
    if chars_per_standard_excerpt is not None and chars_per_standard_excerpt < thresholds["chars_per_standard_excerpt"]:
        failures.append(
            "standards-excerpt/text balance below minimum: "
            f"{chars_per_standard_excerpt:.1f} narrative chars per excerpt < {thresholds['chars_per_standard_excerpt']}"
        )
    if duplicate_ratio > thresholds["duplicate_ratio"]:
        failures.append(f"duplicate long-paragraph ratio too high: {duplicate_ratio:.4f} > {thresholds['duplicate_ratio']}")
    if tables["tables_per_10k_body_characters"] > thresholds["tables_per_10k_chars"]:
        failures.append(
            "table density too high: "
            f"{tables['tables_per_10k_body_characters']} tables per 10k body chars > {thresholds['tables_per_10k_chars']}"
        )
    if tables["rows_per_10k_body_characters"] > thresholds["table_rows_per_10k_chars"]:
        failures.append(
            "table-row density too high: "
            f"{tables['rows_per_10k_body_characters']} rows per 10k body chars > {thresholds['table_rows_per_10k_chars']}"
        )
    if len(tables["prose_heavy_tables"]) > thresholds["prose_heavy_tables"]:
        failures.append(
            "too many prose-heavy tables: "
            f"{len(tables['prose_heavy_tables'])} > {thresholds['prose_heavy_tables']}"
        )
    found_legacy = {term: count for term, count in legacy_hits.items() if count}
    if found_legacy:
        failures.append(f"legacy terms found: {found_legacy}")

    result = {
        "path": str(path),
        "profile": args.profile,
        "thresholds": thresholds,
        "metrics": {
            "body_paragraphs": len(paragraphs),
            "nonempty_body_paragraphs": len(nonempty),
            "body_characters": body_chars,
            "narrative_paragraphs": len(narrative_paragraphs),
            "narrative_characters": narrative_chars,
            "table_characters": table_chars,
            "narrative_table_character_ratio": round(narrative_ratio, 4),
            "tables": tables["tables"],
            "table_rows": tables["rows"],
            "tables_per_10k_body_characters": tables["tables_per_10k_body_characters"],
            "table_rows_per_10k_body_characters": tables["rows_per_10k_body_characters"],
            "prose_heavy_table_count": len(tables["prose_heavy_tables"]),
            "heading_candidates": headings,
            "major_chapters": major_chapters,
            "management_topics": len(present_topics),
            "media_files": media,
            "effective_media_files": effective_media,
            "narrative_characters_per_effective_media": round(chars_per_media, 1) if chars_per_media is not None else None,
            "standard_excerpt_captions": standard_excerpts,
            "narrative_characters_per_standard_excerpt": round(chars_per_standard_excerpt, 1) if chars_per_standard_excerpt is not None else None,
            "average_project_term_chapter_spread": round(average_project_spread, 2),
            "method_detail_ratio": round(method_detail_ratio, 4),
            "duplicate_long_paragraph_ratio": round(duplicate_ratio, 4),
        },
        "management_topic_hits": present_topics,
        "prose_heavy_tables": tables["prose_heavy_tables"],
        "project_terms": project_term_stats,
        "method_terms": method_term_stats,
        "legacy_hits": legacy_hits,
        "failures": failures,
    }

    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding="utf-8")
    print(rendered)
    raise SystemExit(1 if failures else 0)


if __name__ == "__main__":
    main()
