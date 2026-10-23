#!/usr/bin/env python3
"""One-time, source-bound template preparation; never a project format cleaner.

Only accepts the inspected blank format reference (49 paragraphs, 4 blank tables).
Preserves source table XML, page geometry, and original table typography. Creates
new role styles and explicit placeholder slots; never rewrites a finished scheme.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
PROFILE = ROOT / "references/latest-format-20261023.json"
CN = "一二三四五六七八九十"


def chinese(number):
    return CN[number-1] if number <= 10 else "十" + CN[number-11]


def sanitize_reference(source: Path, target: Path):
    """Remove private author/custom metadata only; document content stays exact."""
    with ZipFile(source) as src, ZipFile(target, "w", ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                root = etree.fromstring(data)
                for node in list(root):
                    root.remove(node)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item.filename == "docProps/custom.xml":
                root = etree.fromstring(data)
                for node in list(root):
                    root.remove(node)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(item, data)


def add_style(doc, name, size, bold=False, outline=None, alignment=None):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(size)
    style.font.bold = bold
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    fmt = style.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if alignment is not None:
        fmt.alignment = alignment
    if outline is not None:
        node = OxmlElement("w:outlineLvl")
        node.set(qn("w:val"), str(outline))
        style._element.get_or_add_pPr().append(node)
        fmt.keep_with_next = True
    return style


def add_text(doc, value, style, break_before=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(value)
    if break_before:
        paragraph.paragraph_format.page_break_before = True
    return paragraph


def field(paragraph, instruction, result):
    for kind, text in (("begin", None), ("instruction", instruction), ("separate", None), ("result", result), ("end", None)):
        run = paragraph.add_run()
        if kind in {"begin", "separate", "end"}:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif kind == "instruction":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        else:
            run.text = text


def prepare(source: Path, out_dir: Path, expected_source_sha256: str | None = None):
    source = source.resolve()
    doc = Document(source)
    profile = json.loads(PROFILE.read_text(encoding="utf-8"))
    expected_hash = expected_source_sha256 or profile["converted_source_sha256"]
    actual_hash = hashlib.sha256(source.read_bytes()).hexdigest().upper()
    if actual_hash != expected_hash.upper():
        raise ValueError("Source fingerprint differs from inspected blank reference; re-inspect before authorizing a new hash")
    if len(doc.paragraphs) != 49 or len(doc.tables) != 4 or doc.paragraphs[0].text != "施工组织设计":
        raise ValueError("Not the inspected blank 2026-10-23 format reference; refusing project rewrite")
    if [len(t.columns) for t in doc.tables] != [10, 9, 8, 4]:
        raise ValueError("Source annex schemas changed; inspect new authority before preparation")
    for table in doc.tables:
        if any(cell.text.strip() for row in table.rows[1:] for cell in row.cells):
            raise ValueError("Reference has filled data rows; refuse to treat a project as a blank template")
    out_dir.mkdir(parents=True, exist_ok=True)
    reference = out_dir / "format_reference_20261023.docx"
    output = out_dir / "template_construction_organization_20261023.docx"
    if source in {reference.resolve(), output.resolve()}:
        raise ValueError("Source must be distinct from outputs")
    sanitize_reference(source, reference)
    doc = Document(reference)
    tables = {n: deepcopy(t._tbl) for n, t in zip((1, 2, 3, 6), doc.tables)}
    first_section = deepcopy(doc.sections[0]._sectPr)
    body = doc._element.body
    for child in list(body):
        body.remove(child)
    body.append(first_section)
    # Intentional baseline changes are isolated to new semantic styles.
    body_style = add_style(doc, "BiddingBody", 12)
    body_style.paragraph_format.first_line_indent = Pt(24)
    add_style(doc, "BiddingHeading1", 16, True, 0)
    add_style(doc, "BiddingHeading2", 14, True, 1)
    add_style(doc, "BiddingAppendix", 12, True, 1)
    add_style(doc, "BiddingCaption", 10.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_style(doc, "BiddingCover", 16, True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_style(doc, "BiddingTOC", 12)
    add_style(doc, "BiddingFooter", 10.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    # Explicit TOC role, not an inherited heading or body paragraph.
    for name in ("TOC 1", "TOC 2"):
        add_style(doc, name, 12)
    add_text(doc, "【项目名称】", "BiddingCover")
    add_text(doc, "施工组织设计", "BiddingCover")
    add_text(doc, "【工期】日历天", "BiddingCover")
    add_text(doc, "目录", "BiddingCover", True)
    toc = doc.add_paragraph(style="BiddingTOC")
    field(toc, ' TOC \\t "BiddingHeading1,1,BiddingHeading2,2" \\h \\z ', "【填写：完成正文后更新目录】")
    for chapter in profile["chapters"]:
        number = chapter["number"]
        add_text(doc, f"第{chinese(number)}章  {chapter['title']}", "BiddingHeading1", number == 1)
        add_text(doc, f"【填写：依据本项目清单及招标要求编写{chapter['title']}，见 latest-format.md 内容覆盖要求】", "BiddingBody")
        for appendix in profile["appendices"]:
            if appendix["chapter"] != number:
                continue
            n = appendix["number"]
            add_text(doc, f"附表{chinese(n)}：{appendix['title']}", "BiddingAppendix", True)
            if n == 3:
                add_text(doc, "单位：人", "BiddingCaption")
            if n in tables:
                # Verbatim OOXML preserves source widths, borders, merge and typography.
                body.insert(len(body)-1, tables[n])
            else:
                hint = "插入与 Excel 同源的横道或关键线路网络图，标明 D1 至 Dn" if n == 4 else "插入临设、道路、仓储、水电、消防及生产隔离布置图，缺尺寸时为非比例示意"
                add_text(doc, "【填写：" + hint + "】", "BiddingBody")
    # Dynamic page field: used for page-count verification; no company identifiers.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.style = doc.styles["BiddingFooter"]
    field(footer, " PAGE ", "1")
    setting = OxmlElement("w:updateFields")
    setting.set(qn("w:val"), "true")
    doc.settings._element.append(setting)
    doc.core_properties.title = "施工组织设计空白模板 2026-10-23"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(output)
    # Prove all four original table XML trees survived unchanged.
    final = Document(output)
    if [etree.tostring(t._tbl) for t in final.tables] != [etree.tostring(t._tbl) for t in Document(reference).tables]:
        # Namespace declarations may differ after round-trip; compare canonical nodes.
        def stripped(el):
            return etree.tostring(etree.fromstring(etree.tostring(el)), method="c14n", exclusive=True)
        if [stripped(t._tbl) for t in final.tables] != [stripped(t._tbl) for t in Document(reference).tables]:
            raise ValueError("Source annex XML changed during preparation")
    manifest = {"profile": profile["profile"], "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest().upper(), "artifacts": {}}
    for path in (reference, output):
        manifest["artifacts"][path.name] = {"sha256": hashlib.sha256(path.read_bytes()).hexdigest().upper(), "bytes": path.stat().st_size}
    (out_dir / "latest-template-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="WPS/Word-converted blank format reference, not a project scheme")
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--expected-source-sha256", help="Override only after manually inspecting an equivalent fresh conversion; not a bypass for project documents")
    args = parser.parse_args()
    print(json.dumps(prepare(args.source, args.out_dir, args.expected_source_sha256), ensure_ascii=False, indent=2))
