from collections import Counter
from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\Administrator\Desktop\Metallurgical_BOQ_AgentSkill_Fixed\agents\_work_qc_canteen_gas_to_electric\output\质检中心食堂气改电施工方案_修订版.docx")
OUTPUT = Path(r"C:\Users\Administrator\Desktop\Metallurgical_BOQ_AgentSkill_Fixed\agents\_work_quote_scheme\template_structure.txt")
doc = Document(SOURCE)

lines = []
lines.append(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} shapes={len(doc.inline_shapes)} sections={len(doc.sections)}")
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip().replace("\n", " | ")
    if not text:
        continue
    sizes = sorted({round(run.font.size.pt, 1) for run in paragraph.runs if run.text.strip() and run.font.size})
    fonts = sorted({run.font.name for run in paragraph.runs if run.text.strip() and run.font.name})
    bold = any(run.bold for run in paragraph.runs if run.text.strip())
    lines.append(f"P{index:03d}\tstyle={paragraph.style.name}\tsizes={sizes}\tfonts={fonts}\tbold={bold}\t{text}")

for index, table in enumerate(doc.tables):
    lines.append(f"TABLE {index}: rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows[:6]):
        lines.append(f"  R{row_index}: " + " || ".join(cell.text.strip().replace("\n", " / ")[:120] for cell in row.cells))

font_counts = Counter()
size_counts = Counter()
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        font_counts[run.font.name] += len(run.text)
        size_counts[round(run.font.size.pt, 1) if run.font.size else None] += len(run.text)
lines.append(f"font_counts={font_counts}")
lines.append(f"size_counts={size_counts}")
OUTPUT.write_text("\n".join(lines), encoding="utf-8")
print(OUTPUT)
