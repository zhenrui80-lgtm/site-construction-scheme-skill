"""Test the one-time blank-template migration with small dynamic fixtures."""

import hashlib
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "prepare_latest_template.py"
SPEC = importlib.util.spec_from_file_location("bidding_prepare_latest_template", SCRIPT)
PREPARE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(PREPARE)


def canonical(element):
    return etree.tostring(element, method="c14n", exclusive=True)


class PrepareLatestTemplateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.source = self.directory / "source.docx"
        self.output_dir = self.directory / "prepared"
        self.make_source()

    def tearDown(self):
        self.temp.cleanup()

    def make_source(self, *, paragraphs=49, title="施工组织设计", columns=(10, 9, 8, 4), filled=False):
        document = Document()
        document.add_paragraph().add_run(title)
        for _ in range(paragraphs - 1):
            document.add_paragraph()
        for index, width in enumerate(columns):
            table = document.add_table(rows=2, cols=width)
            table.style = "Table Grid"
            table.autofit = False
            for column, cell in enumerate(table.rows[0].cells):
                run = cell.paragraphs[0].add_run(f"表{index + 1}列{column + 1}")
                run.font.name = "仿宋"
                run.font.size = Pt(10.5)
            if index == 2 and width == 8:
                table.cell(0, 1).merge(table.cell(0, 7))
            if filled and index == 0:
                table.cell(1, 0).paragraphs[0].add_run("已填项目设备数据")
        document.sections[0].top_margin = Cm(2.3)
        document.sections[0].bottom_margin = Cm(2.1)
        document.sections[0].left_margin = Cm(2.4)
        document.sections[0].right_margin = Cm(2.2)
        document.core_properties.author = "private-test-author"
        document.core_properties.last_modified_by = "private-test-editor"
        document.save(self.source)

    def source_hash(self):
        return hashlib.sha256(self.source.read_bytes()).hexdigest()

    def prepare_reviewed_fixture(self):
        return PREPARE.prepare(self.source, self.output_dir, expected_source_sha256=self.source_hash())

    def test_unknown_same_shape_is_rejected_by_default_fingerprint_guard(self):
        original = self.source.read_bytes()
        with self.assertRaises(ValueError):
            PREPARE.prepare(self.source, self.output_dir)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertFalse(self.output_dir.exists())

    def test_explicit_wrong_fingerprint_is_rejected_before_write(self):
        original = self.source.read_bytes()
        with self.assertRaises(ValueError):
            PREPARE.prepare(self.source, self.output_dir, expected_source_sha256="0" * 64)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertFalse(self.output_dir.exists())

    def test_filled_data_row_is_not_accepted_as_blank_template(self):
        self.make_source(filled=True)
        original = self.source.read_bytes()
        with self.assertRaisesRegex(ValueError, "filled data rows"):
            self.prepare_reviewed_fixture()
        self.assertEqual(self.source.read_bytes(), original)
        self.assertFalse(self.output_dir.exists())

    def test_unexpected_paragraph_count_title_and_table_shape_are_rejected(self):
        cases = (
            {"paragraphs": 48},
            {"title": "某项目施工组织设计"},
            {"columns": (10, 9, 8)},
            {"columns": (10, 9, 7, 4)},
        )
        for changes in cases:
            with self.subTest(changes=changes):
                self.make_source(**changes)
                original = self.source.read_bytes()
                with self.assertRaises(ValueError):
                    self.prepare_reviewed_fixture()
                self.assertEqual(self.source.read_bytes(), original)
                self.assertFalse(self.output_dir.exists())

    def test_source_cannot_be_one_of_the_output_files(self):
        self.output_dir.mkdir()
        named_source = self.output_dir / "format_reference_20261023.docx"
        document = Document(self.source)
        document.save(named_source)
        original = named_source.read_bytes()
        with self.assertRaisesRegex(ValueError, "distinct"):
            PREPARE.prepare(
                named_source,
                self.output_dir,
                expected_source_sha256=hashlib.sha256(original).hexdigest(),
            )
        self.assertEqual(named_source.read_bytes(), original)
        self.assertFalse((self.output_dir / "template_construction_organization_20261023.docx").exists())

    def test_preparation_preserves_source_and_all_four_table_xml_trees(self):
        original = self.source.read_bytes()
        source_doc = Document(self.source)
        expected_tables = [canonical(table._tbl) for table in source_doc.tables]
        manifest = self.prepare_reviewed_fixture()
        reference = self.output_dir / "format_reference_20261023.docx"
        output = self.output_dir / "template_construction_organization_20261023.docx"
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual([canonical(table._tbl) for table in Document(reference).tables], expected_tables)
        self.assertEqual([canonical(table._tbl) for table in Document(output).tables], expected_tables)
        self.assertEqual(manifest["source_sha256"], hashlib.sha256(original).hexdigest().upper())
        disk_manifest = json.loads((self.output_dir / "latest-template-manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest, disk_manifest)
        for path in (reference, output):
            self.assertEqual(manifest["artifacts"][path.name]["sha256"], hashlib.sha256(path.read_bytes()).hexdigest().upper())
            with ZipFile(path) as package:
                self.assertIsNone(package.testzip())
                metadata = package.read("docProps/core.xml")
                self.assertNotIn(b"private-test-author", metadata)
                self.assertNotIn(b"private-test-editor", metadata)
        with ZipFile(self.source) as source_package, ZipFile(reference) as reference_package:
            self.assertEqual(source_package.read("word/document.xml"), reference_package.read("word/document.xml"))

    def test_new_body_properties_and_page_geometry_without_normal_style_cleanup(self):
        source_doc = Document(self.source)
        original_normal = canonical(source_doc.styles["Normal"]._element)
        original_section = source_doc.sections[0]
        self.prepare_reviewed_fixture()
        document = Document(self.output_dir / "template_construction_organization_20261023.docx")
        body = document.styles["BiddingBody"]
        self.assertEqual(body.font.name, "宋体")
        self.assertEqual(body._element.rPr.rFonts.get(qn("w:eastAsia")), "宋体")
        self.assertEqual(body.font.size.pt, 12)
        self.assertEqual(body.paragraph_format.line_spacing, 1.5)
        self.assertEqual(body.paragraph_format.first_line_indent.pt, 24)
        self.assertEqual(canonical(document.styles["Normal"]._element), original_normal)
        for property_name in (
            "page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin",
            "header_distance", "footer_distance", "orientation",
        ):
            self.assertEqual(getattr(document.sections[0], property_name), getattr(original_section, property_name))
        self.assertEqual(sum(p.style.name == "BiddingHeading1" for p in document.paragraphs), 15)
        self.assertEqual(sum(p.style.name == "BiddingAppendix" for p in document.paragraphs), 6)
        self.assertGreaterEqual(sum(p.style.name == "BiddingBody" for p in document.paragraphs), 15)
        instructions = document._element.xpath(".//w:instrText")
        self.assertTrue(any("TOC" in (node.text or "") for node in instructions))
        footer_instructions = document.sections[0].footer._element.xpath(".//w:instrText")
        self.assertTrue(any("PAGE" in (node.text or "") for node in footer_instructions))


if __name__ == "__main__":
    unittest.main()
