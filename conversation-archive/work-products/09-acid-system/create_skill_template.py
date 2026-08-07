from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


SOURCE = Path(r"C:\Users\Administrator\Desktop\Metallurgical_BOQ_AgentSkill_Fixed\agents\_work_qc_canteen_gas_to_electric\output\质检中心食堂气改电施工方案_修订版.docx")
OUTPUT = Path(r"C:\Users\Administrator\Desktop\Metallurgical_BOQ_AgentSkill_Fixed\agents\industrial-construction-scheme\assets\条材厂施工方案版式模板.docx")
FONT = "方正仿宋_GB2312"
HEADINGS = [
    "一、编制依据及工程概况",
    "二、施工组织体系",
    "三、施工预设目标及组织管理体系",
    "四、主要施工工序及方法",
    "五、技术质量标准",
    "六、质量保证措施",
    "七、安全环保措施",
    "八、资源需求计划",
    "九、应急预案",
]


def set_run_font(run, size=None, bold=None):
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), FONT)


def replace_preserving_runs(paragraph, text):
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def table_borders(table):
    table_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")
        borders.append(element)
    table_pr.append(borders)


document = Document(SOURCE)
body_start = document.paragraphs[31]._p
body = document._element.body
current = body_start
while current is not None and current.tag != qn("w:sectPr"):
    following = current.getnext()
    body.remove(current)
    current = following

replace_preserving_runs(document.paragraphs[2], "工程名称：【工程名称】")
replace_preserving_runs(document.paragraphs[3], "工程编号：以建设单位下达编号为准")
replace_preserving_runs(document.paragraphs[13], "武汉兴达建筑安装工程有限责任公司")
replace_preserving_runs(document.paragraphs[14], "【编制日期】")

portrait = document.sections[0]
portrait.orientation = WD_ORIENT.PORTRAIT
portrait.page_width = Mm(210)
portrait.page_height = Mm(297)
portrait.left_margin = Inches(1.25)
portrait.right_margin = Inches(1.25)
portrait.top_margin = Inches(1.0)
portrait.bottom_margin = Inches(1.0)

for heading in HEADINGS:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(heading)
    set_run_font(run, 12, True)

landscape = document.add_section(WD_SECTION.NEW_PAGE)
landscape.orientation = WD_ORIENT.LANDSCAPE
landscape.page_width = Mm(297)
landscape.page_height = Mm(210)
landscape.left_margin = Inches(1.0)
landscape.right_margin = Inches(1.0)
landscape.top_margin = Inches(0.75)
landscape.bottom_margin = Inches(0.75)

paragraph = document.add_paragraph()
paragraph.paragraph_format.keep_with_next = True
run = paragraph.add_run("十、施工进度计划")
set_run_font(run, 12, True)

table = document.add_table(rows=3, cols=18)
table.autofit = False
table_borders(table)
headers = ["序号", "工作内容", "开始", "完成", "天数"] + [f"时段{i}" for i in range(1, 14)]
for index, text in enumerate(headers):
    cell = table.cell(0, index)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].add_run(text)
    set_run_font(run, 7, True)
for row_index in (1, 2):
    for col_index in range(18):
        cell = table.cell(row_index, col_index)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run("【填入】" if col_index < 5 else "")
        set_run_font(run, 7, False)

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        set_run_font(run)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)
for section in document.sections:
    for story in (
        section.header,
        section.first_page_header,
        section.even_page_header,
        section.footer,
        section.first_page_footer,
        section.even_page_footer,
    ):
        for paragraph in story.paragraphs:
            for run in paragraph.runs:
                set_run_font(run)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
print(f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} sections={len(document.sections)}")
