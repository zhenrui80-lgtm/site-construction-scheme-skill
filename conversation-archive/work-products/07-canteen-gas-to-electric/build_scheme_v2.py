from __future__ import annotations

from pathlib import Path
import shutil
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "template-reference.docx"
OUTPUT = BASE / "output" / "质检中心食堂气改电施工方案_修订版.docx"
ASSETS = BASE / "assets"
SCHEDULE_RENDER = BASE / "spreadsheet" / "renders" / "文档版横道图.png"

PROJECT = "（实事工程）质检中心食堂气改电"
PROJECT_NO = "工程编号：以建设单位下达编号为准"
COMPANY = "武汉兴达建筑安装工程有限责任公司"
START = "2026年8月10日"
FINISH = "2026年12月31日"
COMPILE_DATE = "2026年7月31日"

BODY_FONT = "仿宋_GB2312"
COVER_FONT = "方正仿宋_GB2312"


def set_run(run, size=12, bold=False, font=BODY_FONT, color=(0, 0, 0)):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, size=6):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), str(size))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    tbl_pr.append(borders)


def set_table_widths(table, widths_cm):
    grid_columns = table._tbl.tblGrid.gridCol_lst
    for grid_column, width_cm in zip(grid_columns, widths_cm):
        grid_column.set(qn("w:w"), str(round(width_cm / 2.54 * 1440)))
    for row in table.rows:
        for cell, width_cm in zip(row.cells, widths_cm):
            width_twips = str(round(width_cm / 2.54 * 1440))
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), width_twips)
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Cm(width_cm)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run(run, 9)


def configure_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        section.top_margin = Cm(3.175)
        section.bottom_margin = Cm(3.175)
    else:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3.175)
        section.right_margin = Cm(3.175)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(1.75)


def set_header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    for paragraph in list(header.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    header.add_paragraph()
    footer = section.footer
    for paragraph in list(footer.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)
    footer.add_paragraph()
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(paragraph)


def remove_body(document):
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def body(document, text, indent=True, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if indent:
        paragraph.paragraph_format.first_line_indent = Pt(24)
    run = paragraph.add_run(text)
    set_run(run, size, bold)
    return paragraph


def heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(6 if level == 1 else 2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run(run, 12, True)
    return paragraph


def list_item(document, prefix, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(-24)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{prefix}{text}")
    set_run(run, 12)
    return paragraph


def add_table(document, headers, rows, widths, font_size=9.5, aligns=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_table_widths(table, widths)
    repeat_header(table.rows[0])
    prevent_split(table.rows[0])
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(value))
        set_run(run, font_size, True)
    for values in rows:
        row = table.add_row()
        prevent_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Cm(widths[index])
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                aligns[index]
                if aligns
                else WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run(run, font_size)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.5
    return table


def add_picture(document, path, caption, width_cm, source=None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    caption_p = document.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(0)
    caption_p.paragraph_format.space_after = Pt(0)
    run = caption_p.add_run(caption)
    set_run(run, 9, False)
    if source:
        source_p = document.add_paragraph()
        source_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        source_p.paragraph_format.space_before = Pt(0)
        source_p.paragraph_format.space_after = Pt(2)
        run = source_p.add_run(source)
        set_run(run, 8, False, color=(90, 90, 90))


def add_cover(document):
    for _ in range(2):
        document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("工程名称：")
    set_run(run, 14, False, COVER_FONT)
    run = paragraph.add_run(PROJECT)
    set_run(run, 14, False, COVER_FONT)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(PROJECT_NO)
    set_run(run, 14, False, COVER_FONT)
    paragraph.paragraph_format.space_after = Pt(54)

    for character in "施工方案":
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(9)
        run = paragraph.add_run(character)
        set_run(run, 26, False, COVER_FONT)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(20)
    for label in ("编制：________________", "审核：________________", "审批：________________"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(label)
        set_run(run, 15, False, COVER_FONT)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(26)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(COMPANY)
    set_run(run, 14, False, COVER_FONT)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(COMPILE_DATE)
    set_run(run, 14, False, COVER_FONT)


def add_approval(document):
    document.add_page_break()
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, 6)
    set_table_widths(table, [1.45, 13.1])
    row = table.rows[0]
    row.height = Cm(22.4)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    left, right = row.cells
    left.width = Cm(1.45)
    right.width = Cm(13.1)
    left.text = ""
    right.text = ""
    for cell in (left, right):
        set_cell_margins(cell, 120, 100, 120, 100)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = left.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run("产\n权\n厂\n审\n批\n意\n见")
    set_run(run, 12)
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    paragraph = right.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("签字：                    （盖章）        年   月   日")
    set_run(run, 12)


def add_toc(document):
    document.add_page_break()
    for _ in range(2):
        document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(30)
    run = paragraph.add_run("目  录")
    set_run(run, 20)
    items = [
        "一、编制依据及工程概况",
        "二、施工组织体系",
        "三、施工预设目标及组织管理体系",
        "四、主要施工工序及方法",
        "五、技术质量标准",
        "六、质量保证措施",
        "七、安全环保措施",
        "八、资源需求计划",
        "九、应急预案",
        "十、施工进度计划",
    ]
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(28)
        paragraph.paragraph_format.line_spacing = 2.0
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(item)
        set_run(run, 14)


def add_standards_evidence(document):
    heading(document, "规范查新网络核验截图", 1)
    body(
        document,
        "下图仅保留与本工程直接相关的官方公告关键区域，作为2026年7月31日规范版本核验的过程证据。截图不替代正式标准文本，施工和验收仍应使用建设单位确认的正版现行文本。",
    )
    add_picture(
        document,
        ASSETS / "规范查新_GB55024-2022_正文.png",
        "图1  《建筑电气与智能化通用规范》（GB 55024—2022）发布公告关键内容",
        12.2,
        "来源：住房和城乡建设部标准公告页面，核验日期：2026年7月31日",
    )
    body(
        document,
        "全国标准信息公共服务平台核验显示，《工业企业煤气安全规范》（GB 6222—2025）为强制性现行国家标准，自2026年5月1日起实施并代替GB 6222—2005。本工程只有在产权单位确认介质属于工业煤气时才按该规范控制；天然气或城镇燃气系统按GB 55009—2021及产权单位规程执行。",
    )


def add_main_content(document):
    heading(document, "一、编制依据及工程概况", 1)
    heading(document, "（一）编制依据", 2)
    body(
        document,
        "本施工方案依据《（实事工程）质检中心食堂气改电工程量清单》、建设单位现场管理制度、施工区域实际条件及甲供厨房设备接口资料编制。工程量清单共49项，涉及燃气停用与拆除、低压配电、封闭母线、桥架及钢管、电缆与接地、厨房设备就位接电、混凝土和砌体、屋面防水、装饰恢复、排水沟以及措施项目等内容。",
    )
    body(
        document,
        "方案编制日期为2026年7月31日，计划于2026年8月10日开工、2026年12月31日完工，总工期144日历天。施工期间如设计图纸、设备参数、燃气介质、停送电条件或食堂运营窗口发生变化，应履行图纸会审、技术核定和方案交底程序后实施，不得以本方案替代专项作业票和能源介质操作指令。",
    )
    body(
        document,
        "国家及行业标准按开工日期已经实施的现行版本执行。标准条文与强制性工程建设规范不一致时，以强制性工程建设规范为准；建设单位企业标准高于国家、行业标准时，按较严要求执行。涉及燃气系统时，必须先由产权单位书面确认介质属性，再确定采用天然气或工业煤气的相应安全规范。",
    )
    standards = [
        "《建筑电气与智能化通用规范》（GB 55024—2022）；",
        "《建筑电气工程施工质量验收规范》（GB 50303—2015）；",
        "《电气装置安装工程 电缆线路施工及验收标准》（GB 50168—2018）；",
        "《电气装置安装工程 接地装置施工及验收规范》（GB 50169—2016）；",
        "《电气装置安装工程 低压电器施工及验收规范》（GB 50254—2014）；",
        "《建筑与市政工程施工现场临时用电安全技术标准》（JGJ/T 46—2024）；",
        "《燃气工程项目规范》（GB 55009—2021）；",
        "《工业企业煤气安全规范》（GB 6222—2025）；",
        "《危险化学品企业特殊作业安全规范》（GB 30871—2022，适用范围满足时执行）；",
        "《建筑防火通用规范》（GB 55037—2022）；",
        "《混凝土结构工程施工质量验收规范》（GB 50204—2015）；",
        "《砌体结构工程施工质量验收规范》（GB 50203—2011）；",
        "《屋面工程质量验收规范》（GB 50207—2012）；",
        "《建筑地面工程施工质量验收规范》（GB 50209—2010）；",
        "《建筑装饰装修工程质量验收标准》（GB 50210—2018）；",
        "《建筑给水排水及采暖工程施工质量验收规范》（GB 50242—2002）；",
        "《建筑工程施工质量验收统一标准》（GB 50300—2013）；",
        "《建筑施工高处作业安全技术规范》（JGJ 80—2016）；",
        "《坠落防护 安全带》（GB 6095—2021）；",
        "《生产经营单位生产安全事故应急预案编制导则》（GB/T 29639—2020）。",
    ]
    for index, standard in enumerate(standards, 1):
        list_item(document, f"{index}、", standard)
    body(
        document,
        "版本控制说明：GB 6222—2005已于2026年5月1日废止，工业企业煤气相关作业应执行GB 6222—2025；施工现场临时用电采用JGJ/T 46—2024，不再引用JGJ 46—2005。上述标准在开工前由技术负责人再次核验，标准清单、电子或纸质受控文本和查新记录纳入项目资料。",
    )
    add_standards_evidence(document)

    heading(document, "（二）工程概况", 2)
    heading(document, "1、工程范围", 3)
    body(
        document,
        "本工程位于武汉钢铁有限公司质检中心食堂区域，施工目标是使原有燃气用能系统有序退出，完成厨房主要设备电气化改造，并同步恢复受施工影响的建筑、屋面、地面、排水和装饰系统。施工区域与食堂运营区相邻，存在人员流动、食品卫生、能源介质和多专业交叉作业等特点，必须实行硬隔离和分区移交。",
    )
    body(
        document,
        "燃气专业主要包括原DN120燃气管道停气、阀门关闭、可靠隔断或盲板、氮气置换、检测确认、约3.6t管道拆除和封口标识；清单暂列φ133×4.5管道约252m、理论容积约3.04m³，按不少于5倍容积估算置换量约15.2m³，氮气暂列40m³。实际置换方案和氮气用量必须根据介质、系统分段、残余支管和检测结果复核。",
    )
    body(
        document,
        "电气专业主要包括2台800×800×2200低压GGD柜、1台JXF3户外不锈钢动力箱、20m封闭铜母线、约730m动力电缆、镀锌桥架、镀锌钢管、接地干线、等电位端子箱以及厨房设备末端配管接线。甲供设备到货后应复核铭牌功率、电压、相数、启动方式、保护要求和接口位置，未经书面确认不得仅按清单估算值制作终端。",
    )
    body(
        document,
        "土建和装饰专业包括墙体及地坪拆除、墙洞开设、C30设备基础和结构、室内外电缆沟、排水沟、沟盖板、屋面防水及保温、墙地面和吊顶恢复、门窗和雨水系统修复。所有拆改范围应先探查暗埋管线并确认结构安全，恢复做法应与原建筑的材料、颜色、标高和排水坡向协调。",
    )
    quantities = [
        ["燃气停用与拆除", "DN120燃气管道拆除约3.6t；φ133×4.5管道系统置换；氮气暂列40m³", "产权单位操作、可靠隔断、连续检测"],
        ["低压配电设备", "GGD柜2台；JXF3户外动力箱1台；TD28等电位端子箱1台", "参数、短路容量及回路需深化确认"],
        ["母线及电缆", "封闭铜母线20m；动力电缆约730m", "敷设前后绝缘测试、终端挂牌"],
        ["桥架与配管", "200×100镀锌桥架约1.364t；镀锌钢管约5.262t；挠性管35m", "支架、防腐、跨接和防火封堵"],
        ["接地系统", "镀锌接地材料约0.038t；BVR 1×6黄绿线100m", "形成可靠PE及等电位网络"],
        ["混凝土工程", "C30设备基础4.02m³；C30结构10.56m³", "轴线、标高、预埋件和养护"],
        ["拆除与沟槽", "砌体拆除15.96m²；地坪开挖14.4m²；排水沟40m", "探管、临边防护、湿法降尘"],
        ["屋面及装饰", "屋面防水52.92m²；墙面、地面、吊顶、门窗等恢复", "防水节点和成品保护"],
        ["机械措施", "12m登高车5台班；25t汽车吊2台班；10t平板车2台班", "专项交底、持证指挥和警戒"],
    ]
    add_table(document, ["专业/分部", "主要工程量", "施工控制重点"], quantities, [3.0, 6.5, 5.0], 9.2)

    heading(document, "2、施工特点和管理难点", 3)
    body(
        document,
        "本工程的第一项难点是燃气系统边界不明和介质风险。施工人员无权自行操作产权单位阀门，停气、隔断、放散、置换、取样和恢复管理必须由产权单位牵头。施工单位负责在交接边界外组织拆除，但只有在书面交接、检测合格和现场监护条件同时满足后才能开始作业。",
    )
    body(
        document,
        "第二项难点是食堂不停产或分区运行条件下的施工组织。施工区必须采用封闭围挡与运营区隔离，人员、材料、废弃物分别设置通道；粉尘、异味、焊烟、噪声和临时停电不得影响食品加工。每日作业结束后进行清场、吸尘、消毒和通风，经使用单位检查后方可撤离。",
    )
    body(
        document,
        "第三项难点是甲供设备参数与配电深化之间的联动。GGD柜、封闭母线、电缆截面、桥架容量和保护整定均与设备功率有关。项目部建立设备参数台账，执行“铭牌核对—接口确认—回路深化—材料下单—到货复核—终端接线—保护试验”的闭环。",
    )
    body(
        document,
        "第四项难点是土建、电气、设备和装饰多专业交叉。项目采用样板先行和区域移交制度，先确认沟槽、基础、墙洞和支架的综合定位，再施工桥架、钢管和电缆，最后完成设备接电和装饰收口。未经交接的区域不得平行堆放材料或安排相互干扰的作业。",
    )

    heading(document, "二、施工组织体系", 1)
    heading(document, "（一）项目组织设置", 2)
    body(
        document,
        "项目实行项目经理负责制，设置技术负责人、施工员、安全员、质量员、材料员、资料员和专业班组长。项目经理负责工期、质量、安全、协调和成本的总体控制；技术负责人负责方案、深化设计、技术交底、能源边界确认和试验策划；安全员负责危险作业许可、现场巡查、气体检测和应急资源检查。",
    )
    body(
        document,
        "项目部与建设单位、产权单位、食堂使用单位、监理单位、设计单位和甲供设备厂家建立联合协调机制。燃气停用、停送电、设备接口、食堂运营窗口、吊装通道和屋面作业均形成书面协调记录；涉及生产能源的指令由产权单位发布，项目部不得越权操作。",
    )
    roles = [
        ["项目经理", "全面履约；审批施工安排；协调能源、运营和外部单位；组织验收与移交"],
        ["技术负责人", "技术核定、方案交底、深化设计、危险作业技术条件确认、试验和竣工资料审核"],
        ["施工员", "日计划、作业面移交、工序协调、班组交底、实物量和施工记录"],
        ["安全员", "作业票、隔离警戒、气体检测、临电、高处、吊装、消防和应急检查"],
        ["质量员", "材料验收、样板、隐蔽、停检点、试验记录、分项验收和问题闭环"],
        ["材料员", "报审、采购、甲供接收、标识、保管、防潮防火和余料退库"],
        ["资料员", "图纸、变更、交底、检验、影像、试验、会议、竣工图和移交资料"],
        ["班组长", "班前会、人员资质、工器具点检、工序自检、清场和异常报告"],
    ]
    add_table(document, ["岗位", "主要职责"], roles, [3.2, 11.3], 9.8)
    body(
        document,
        "每日开工前由施工员组织班前会，说明当日作业范围、能源状态、交叉作业、风险控制和撤离路线；每日收工后由安全员、质量员和使用单位代表进行联合检查。每周召开进度协调会，对材料到货、停气停电窗口、土建条件、甲供设备和验收资料进行滚动更新。",
    )
    body(
        document,
        "项目实行作业面移交制度。土建向电气移交时确认基础强度、沟槽、预埋件、墙洞和标高；电气向设备安装移交时确认回路、接地、绝缘和接口；装饰恢复前确认隐蔽工程和试验记录。移交双方在检查表上签字，遗留问题标明责任和完成期限。",
    )
    body(
        document,
        "技术文件实行受控发放。施工现场只使用经批准的最新图纸、清单、技术核定和设备资料，旧版文件回收或明显标识作废。班组收到变更后由技术负责人进行二次交底，施工员检查现场是否存在已按旧版实施的部位，并及时组织复核处理。",
    )
    body(
        document,
        "项目建立接口台账，对燃气隔断点、电源接入点、盘柜回路、母线连接、设备端子、排水接口和建筑收口逐项编号。每个接口明确上游、下游、责任单位、检查条件和验收记录，避免专业之间出现“已完成但无人接收”或“边界遗漏”的问题。",
    )

    heading(document, "三、施工预设目标及组织管理体系", 1)
    heading(document, "（一）管理目标", 2)
    body(
        document,
        f"工期目标：{START}开工，{FINISH}完成，计划工期144日历天。项目以燃气停用窗口、配电设备和甲供设备到货、土建移交及送电条件为主要节点，采用月度总控、周计划和三日滚动计划逐级分解，不以压缩置换、试验、养护和验收时间换取表面进度。",
    )
    body(
        document,
        "质量目标：分项工程一次验收合格，燃气保留端口安全可靠，配电、接地和保护系统功能完整，厨房设备单机及负荷试运行正常，屋面防水、装饰和排水恢复无渗漏、空鼓、倒坡等质量缺陷。竣工图、试验报告、隐蔽记录和设备资料与实体同步完成。",
    )
    body(
        document,
        "安全目标：杜绝燃气泄漏、中毒窒息、火灾爆炸、触电、高处坠落、起重伤害和食堂污染事故，控制未遂事件和违章行为。重大风险作业全部实行书面许可、专人监护、过程检测和项目负责人带班，任何人不得在安全条件不具备时强令施工。",
    )
    body(
        document,
        "文明施工目标：施工区和食堂运营区硬隔离，人员、材料、废弃物通道分开，材料按区域和使用顺序定置摆放；每天工完场清，围挡、消防、临电和通道保持完好，不占用疏散通道，不影响既有交通和必要的生产组织。",
    )
    body(
        document,
        "环境和卫生目标：控制粉尘、噪声、焊烟、废水、废气和固体废弃物，建筑垃圾和废旧材料分类清运；施工产生的异味和异物不得进入食品加工、储存和就餐区域。每次作业面移交前完成吸尘、清洁、消毒和通风，经使用单位确认后恢复运营。",
    )
    heading(document, "（二）施工准备和开工条件", 2)
    body(
        document,
        "开工前完成图纸会审、清单复核和现场测绘，核对燃气介质、阀门和盲板位置、配电电源和短路容量、设备功率及接口、沟槽路径、结构墙洞、消防通道和食堂运营边界。测绘结果形成综合平面和标高控制记录，冲突点经确认后再放线。",
    )
    body(
        document,
        "材料和设备实行报审制度。盘柜、母线、电缆、桥架、钢管、接地材料、防水卷材、混凝土和装饰材料应具有合格证、检验报告及必要的型式试验资料。盘柜和甲供设备到货后联合开箱，核对型号、数量、附件、外观、端子和随机文件。",
    )
    body(
        document,
        "施工围挡、人员通道、材料通道、临时配电箱、消防器材、吸尘设备、气体检测仪、应急照明和应急集合点应在开工前布置完成。燃气、临电、高处、吊装、动火、开挖等作业的人员资质和工器具状态须经检查，检测仪器在校准有效期内。",
    )
    heading(document, "（三）施工分区与总体部署", 2)
    body(
        document,
        "现场划分为燃气停用拆除区、配电柜及母线安装区、桥架配管和电缆敷设区、厨房设备安装区、土建及装饰恢复区。各区设置责任牌、边界标识和材料定置点；燃气作业区与电气火源作业区在时间和空间上隔离，禁止未经批准的交叉作业。",
    )
    body(
        document,
        "总体施工顺序为：深化核对和材料报审→施工隔离和临时设施→燃气停用、隔断、置换和检测→原管道拆除→沟槽、墙洞和基础土建→桥架、钢管、接地和盘柜母线安装→电缆敷设及终端→甲供设备就位接电→建筑装饰恢复→电气试验和设备调试→负荷试运行→清洁消毒、验收和移交。",
    )
    body(
        document,
        "第一阶段重点完成接口确认和前置条件，不安排不可逆的大面积拆除。燃气介质、配电电源、设备功率和结构开孔确认后，才允许形成最终材料清单和加工尺寸。盘柜、封闭母线和设备基础之间至少进行一次联合测量，避免设备到场后返工。",
    )
    body(
        document,
        "第二阶段以燃气退出和土建作业面形成作为关键。燃气拆除完成并封口验收后，释放原管线路径；土建分区完成沟槽、基础和墙洞，达到移交条件的区域立即转入桥架、钢管和接地施工。未达到养护强度或防水隐蔽未验收的区域禁止抢先安装设备。",
    )
    body(
        document,
        "第三阶段以电气安装、设备接电和恢复收口为主。桥架、盘柜、母线和电缆按系统逐回路完成，测试合格后向设备班组移交。装饰恢复与调试错时进行，试运行区域保持清洁和可进入，最终按“单机—分系统—全系统—负荷”的顺序验收。",
    )
    add_picture(
        document,
        ASSETS / "总体施工示意图.png",
        "图3  质检中心食堂气改电总体施工关系示意图",
        13.4,
        "根据工程量清单和现场条件绘制，实际接口以深化设计及现场交底为准",
    )

    heading(document, "四、主要施工工序及方法", 1)
    heading(document, "（一）原燃气系统停用、置换与拆除", 2)
    heading(document, "1、介质确认与作业许可", 3)
    body(
        document,
        "燃气作业前由建设单位组织产权管理、使用、设计和施工单位确认介质种类、运行压力、供气方向、阀门编号、影响用户和停气范围。若介质为天然气或城镇燃气，执行GB 55009—2021及产权单位规程；若确认为工业煤气，执行GB 6222—2025。介质未确认、管线边界不清或停气影响未经批准时不得作业。",
    )
    body(
        document,
        "停气、盲板抽堵、置换、检测、动火、高处和吊装作业按建设单位制度办理作业票。作业票必须写明作业地点、介质、隔离点、检测点、检测频次、监护人、有效时限和异常处置方式。班组人员在票证有效范围内作业，条件变化立即停止并重新确认。",
    )
    heading(document, "2、停气和可靠隔断", 3)
    body(
        document,
        "阀门关闭由产权单位授权人员操作，施工单位只负责配合和确认。关闭前核对阀门编号、流向和下游用户，关闭后挂牌上锁并记录状态。仅关闭阀门不能作为拆除的唯一隔离措施，应采用盲板、断开加封堵或产权单位认可的可靠物理隔断。",
    )
    body(
        document,
        "盲板的材质、厚度、直径和耐压能力按系统参数确定，设置在便于检查和拆装的位置。盲板两侧需要放散或检测时按专项措施布置。完成隔断后由产权单位、监护人和施工负责人共同签字，隔离点未经许可不得移动。",
    )
    heading(document, "3、氮气置换和检测", 3)
    body(
        document,
        "置换前划定警戒区，控制火源、电气火花和无关人员。氮气瓶组、减压装置、汇流排、软管和接头经检查无泄漏，置换气体从有利于气体排出的方向缓慢送入，放散口引至安全地点并处于下风侧。严禁向室内、人员通道或可能积聚的低洼处直接放散。",
    )
    body(
        document,
        "清单管道按外径133mm、壁厚4.5mm、长度约252m估算，内径约124mm，理论容积约3.04m³。按5倍容积计算仅为策划基准，实际置换量应考虑支管、设备腔体和死角，并以连续检测结果为准。氮气具有窒息风险，室内作业必须保持强制通风并同时监测氧含量。",
    )
    body(
        document,
        "检测从放散点、末端、低点和可能积聚部位取样，检测仪须适配介质且在校准期内。达到产权单位和适用标准规定的允许拆除条件后，至少复测确认一次并记录检测时间、位置、仪器编号和数值。停止作业后重新开工、隔离状态变化或超过作业票时限时重新检测。",
    )
    add_picture(
        document,
        ASSETS / "燃气停用置换示意图.png",
        "图4  原燃气管道停气、隔断与氮气置换工序示意图",
        13.4,
        "示意图不替代产权单位操作票、盲板图和专项作业许可",
    )
    heading(document, "4、管道拆除与封口", 3)
    body(
        document,
        "拆除前再次确认隔断、通风、检测和监护条件。先拆除保温、支架附件和非承重连接，再按“先支撑、后切断；先小段、后大段；先高处、后低处”的顺序分段拆除。管段切割前设置可靠支撑和溜绳，禁止自由坠落或利用原管道承受拆除荷载。",
    )
    body(
        document,
        "介质确认前不得使用产生火花或高温的切割方式。允许动火时仍须办理动火票、连续检测、设置消防器材和专人监火。拆下管段及时封堵、标识、分类和转运，原系统保留端口采用盲板或焊接封头永久封闭，并标注介质、状态、日期和责任单位。",
    )
    body(
        document,
        "支架拆除不得损伤主体结构和保留设施，膨胀螺栓、焊疤和孔洞按原基层做防腐、防水或饰面修复。拆除物按产权单位要求移交或清运，建立称量、交接和影像记录，严禁将残余管段堆放在疏散通道和食堂运营区域。",
    )

    heading(document, "（二）低压配电、母线、桥架和电缆工程", 2)
    heading(document, "1、施工测量与综合排布", 3)
    body(
        document,
        "电气施工前复核电源接入点、柜体基础、封闭母线路径、桥架标高、设备进线方向和检修空间，结合风管、水管、燃气保留设施、吊顶和结构梁柱进行综合排布。桥架和钢管尽量沿结构构件敷设，不占用消防通道和设备操作面。",
    )
    body(
        document,
        "放线时标出支架、弯通、伸缩、穿墙和防火封堵位置。需要开孔时先使用探测设备检查暗埋管线，结构墙洞和楼板洞口经设计或建设单位确认。支架生根点不得落在薄弱构件、装饰面或防水节点上。",
    )
    heading(document, "2、盘柜和动力箱安装", 3)
    body(
        document,
        "基础槽钢制作前复核柜体排列、进出线方式和母线接口，槽钢调直、除锈、防腐后安装，水平度、直线度和标高满足验收要求，并与接地干线可靠连接。盘柜搬运采用专用滚杠或吊点，严禁从门框、母排和仪表部位受力。",
    )
    body(
        document,
        "GGD柜和动力箱就位后调整垂直度、柜间缝隙和排列面，采用螺栓固定；柜门、抽屉、机械联锁和接地软铜带动作可靠。柜内母排相序、绝缘间距、紧固扭矩和防松标记逐项检查，临时敞开的孔洞安装封板，备用回路保留清晰标识。",
    )
    body(
        document,
        "盘柜二次回路按图核线，线号、端子号和回路名称一致，控制线与动力线分开敷设。设备厂家有特殊整定要求时，由设计、建设单位和厂家共同确认；保护定值未批准、回路绝缘未合格、联锁未验证时不得送电。",
    )
    heading(document, "3、封闭母线安装", 3)
    body(
        document,
        "封闭母线到货后检查型号、额定电流、防护等级、导体规格、连接器、绝缘件和外壳完整性。安装前根据实测尺寸复核直线段、弯头、始端箱和伸缩节，支架间距及吊架形式按厂家文件和设计确定，不得由母线连接端承受额外机械荷载。",
    )
    body(
        document,
        "母线连接面保持清洁、干燥，连接螺栓使用经校验的扭矩扳手按规定顺序紧固并作标记。外壳跨接和保护接地连续可靠，穿越防火分区处完成防火封堵。安装完成后进行绝缘电阻、相序、连续性和外观检查，合格后再与盘柜和设备连接。",
    )
    heading(document, "4、桥架、支架和钢管安装", 3)
    body(
        document,
        "镀锌桥架支架按荷载和现场结构设置，转弯、三通、进出盘柜和跨越伸缩缝处增设支撑。桥架连接板齐全、螺栓由内向外安装且无尖锐端头，盖板完整；直线段超过规定长度或跨越建筑伸缩缝时设置补偿措施。",
    )
    body(
        document,
        "桥架全长形成可靠电气通路，非镀锌连接处设置跨接线，桥架与盘柜、动力箱和设备外壳接地系统可靠连接。电缆桥架不得作为燃气管道、工艺管道或脚手架的支撑，也不得在桥架上进行与本体无关的焊接。",
    )
    body(
        document,
        "镀锌钢管采用螺纹或专用连接件连接，管口平整、无毛刺，弯曲半径满足电缆穿管要求。沿立柱引下时设置独立卡具和防机械损伤措施，距地面及设备操作区域的管段应加密固定。金属管只作为辅助保护导体时仍应验证连接连续性，不得替代设计要求的PE导体。",
    )
    heading(document, "5、电缆敷设、终端和防火封堵", 3)
    body(
        document,
        "电缆到货核对型号、截面、电压等级和长度，盘体无损伤，合格证和检验资料齐全。敷设前测量绝缘电阻并编制电缆表，检查桥架、钢管、穿墙套管和牵引路径，安排人员、滑轮、放线架和通讯方式。",
    )
    body(
        document,
        "电缆从盘上端引出，控制牵引力和侧压力，弯曲半径不小于产品和规范要求。不同电压等级和用途的电缆按设计分层或分隔敷设，转弯、分支、进出设备和直线适当间距设置标牌。电缆在桥架内排列整齐，不交叉堆压，不用金属丝直接绑扎绝缘层。",
    )
    body(
        document,
        "终端制作由合格电工实施，切剥绝缘不得伤及导体，铜鼻子规格与导体匹配，压接模具和压接次数符合产品要求。端子连接使用扭矩工具，完成后设置防松标记。穿墙、穿楼板、进盘柜和防火分区处采用与构造匹配的防火材料封堵，表面平整并形成验收记录。",
    )
    heading(document, "6、接地和等电位", 3)
    body(
        document,
        "配电柜、动力箱、封闭母线外壳、桥架、钢管、厨房设备金属外壳及其他可导电部分按设计接入PE系统。接地干线的规格、路径、搭接和防腐符合要求，连接点设置防松和标识，隐蔽前检查焊缝或机械连接质量。",
    )
    body(
        document,
        "厨房设备区设置等电位联结，TD28端子箱位置便于检查和维护。柔性连接处使用黄绿双色铜导体，门、盖板和可拆卸部件的接地跨接不妨碍正常开启。送电前测量接地电阻并检查保护导体连续性，测试结果与回路、设备编号对应。",
    )

    heading(document, "（三）甲供厨房设备就位、接线和调试", 2)
    body(
        document,
        "甲供设备到场由建设单位、使用单位、厂家和施工单位联合开箱，核对设备名称、型号、数量、外观、附件、随机资料和铭牌参数。设备基础、地脚、排水、电源和检修空间未验收前不得强行就位。",
    )
    body(
        document,
        "设备搬运采用适合重量和通道的搬运工具，吊点使用厂家指定位置，门板、操作面板、加热元件和进线端不得受力。就位后调整水平、垂直和相互间距，固定可靠，设备与建筑之间便于清洁，连接口不承受附加应力。",
    )
    body(
        document,
        "接线前确认上级开关断开并挂牌上锁，核对电压、相数、额定功率、保护导体和端子说明。电缆端子压接可靠，PE导体先接后拆，软管和电缆入口采取防水、防油和防机械损伤措施。控制面板、急停、门控和温控保护按厂家要求逐项检查。",
    )
    body(
        document,
        "单机调试先空载后带负荷，记录电压、电流、相序、温升、振动、噪声、控制和保护动作。加热类设备检查绝缘、接地、超温和干烧保护；旋转设备检查方向和防护罩；涉水设备检查漏水和防护等级。使用单位和厂家参加功能确认并签字。",
    )

    heading(document, "（四）土建、给排水及装饰恢复", 2)
    heading(document, "1、拆除、开孔和沟槽施工", 3)
    body(
        document,
        "墙体、地坪、屋面和装饰拆除前设置围挡、接尘和成品保护，复核暗埋电线、给排水和结构构件。拆除采用小型工具分层分块进行，湿法降尘，禁止大面积冲击造成保留结构开裂。洞口形成后立即设置临边防护或临时封闭。",
    )
    body(
        document,
        "室内地坪开挖和室外电缆沟施工前放线探测，人工试探确认地下设施。沟槽两侧禁止集中堆载，设置挡水和排水措施，夜间设置警示。验槽后及时施工垫层、沟体和盖板，盖板承载、平整度和开启维护条件符合使用要求。",
    )
    heading(document, "2、混凝土基础和结构", 3)
    body(
        document,
        "C30设备基础和结构施工前复核轴线、标高、尺寸、预埋件、地脚孔和设备接口。模板拼缝严密、支撑稳定，钢筋品种、规格、间距和保护层经隐蔽验收。混凝土分层浇筑、振捣密实，避免碰动预埋件和预留孔。",
    )
    body(
        document,
        "浇筑完成后及时养护，拆模和设备安装时混凝土强度应满足设计及施工要求。基础表面平整，标高和中心线复测记录齐全。二次灌浆前清理并湿润基层，灌浆料按产品要求搅拌、灌注和养护。",
    )
    heading(document, "3、屋面防水和保温恢复", 3)
    body(
        document,
        "屋面恢复采用3+3厚SBS防水卷材、70mm挤塑聚苯板保温及20mm水泥砂浆保护层。施工前检查基层强度、含水率、平整度和排水坡向，管根、落水口、阴阳角和洞口先施工附加层，卷材铺贴方向、搭接宽度和热熔程度符合工艺要求。",
    )
    body(
        document,
        "防水层完成后组织外观和蓄水或淋水检查，合格后及时施工保温和保护层。穿屋面支架及管线节点采用与原防水体系相容的密封材料，禁止在完成的防水层上直接焊接、切割或拖拉尖锐构件。",
    )
    heading(document, "4、砌体、地面、墙面和吊顶恢复", 3)
    body(
        document,
        "砌体恢复前清理基层并处理拉结，砌块规格、砂浆强度和灰缝满足要求。墙面抹灰分层施工，基层湿润、冲筋找平，防止空鼓和开裂。内外墙饰面颜色和纹理与周边协调，门窗洞口收边顺直，密封完整。",
    )
    body(
        document,
        "地砖及800×800防静电不发火地面施工前检查基层强度、含水率、平整度和接地网络。材料性能满足设计要求，排版、标高和坡向经样板确认，面层粘结牢固、缝格顺直。吊顶龙骨和面板恢复前检查上部电气、消防和管线隐蔽工程，检修口位置与设备对应。",
    )
    heading(document, "5、排水沟和管道", 3)
    body(
        document,
        "新建300×300排水沟按放线施工，沟底和沟壁采用清单要求的混凝土做法，控制纵坡、标高和转角。沟内表面平顺、无倒坡和积水，盖板排列整齐、启闭方便。PVC-U DN110排水管使用配套胶粘剂和管件，承插深度、坡度、支吊架和伸缩措施符合要求。",
    )
    body(
        document,
        "DN120不锈钢管回装前确认材质和接口形式，切口平整，焊接或机械连接满足设计。管道完成后进行通水和严密性检查，清理沟内杂物并恢复周边地面。厨房排水系统交付前进行冲洗和卫生清洁。",
    )

    heading(document, "五、技术质量标准", 1)
    heading(document, "（一）材料设备质量控制", 2)
    body(
        document,
        "所有材料设备实行进场报验和标识管理。对盘柜、母线、电缆、桥架、钢管、接地材料、防水卷材、保温板、混凝土、砌块、地砖和管材核对合格证、检验报告、规格型号和外观。对防火封堵材料、检测仪器和需要复验的材料按规定取样或核验有效证明。",
    )
    body(
        document,
        "材料存放按防潮、防晒、防火和防污染要求分区。电缆盘竖直存放并防止滚动，盘柜和母线覆盖防潮，防水材料远离火源，装饰材料不得进入食品加工区。甲供设备发现缺件、损伤或参数不符时形成开箱记录，不得擅自修复或代换。",
    )
    heading(document, "（二）工序检验和停检点", 2)
    inspections = [
        ["燃气隔断", "阀门、盲板、封堵、挂牌和交接边界", "产权单位书面交接及影像记录"],
        ["置换检测", "可燃/有毒气体、氧含量、检测点和频次", "检测记录及允许作业签字"],
        ["沟槽及基础", "槽底、轴线、标高、钢筋、模板、预埋件", "隐蔽验收和混凝土记录"],
        ["盘柜母线", "基础、垂直度、相序、绝缘间距、扭矩、接地", "安装检查和扭矩记录"],
        ["桥架配管", "路径、支架、连接、跨接、防腐和封堵", "隐蔽验收及影像记录"],
        ["电缆终端", "型号、绝缘、排列、标识、压接、扭矩", "绝缘和敷设记录"],
        ["接地等电位", "干线、跨接、设备PE和端子箱", "连续性及接地电阻记录"],
        ["屋面防水", "基层、附加层、搭接、收头、蓄水/淋水", "隐蔽和防水检查记录"],
        ["设备调试", "控制、保护、空载、负荷、温升和异常", "单机及联动试运行记录"],
    ]
    add_table(document, ["停检点", "主要检查内容", "验收输出"], inspections, [3.0, 7.0, 4.5], 9.2)
    heading(document, "（三）电气试验和送电条件", 2)
    body(
        document,
        "送电前完成电缆和设备绝缘电阻、保护导体连续性、接地电阻、相序、控制回路、联锁、断路器和漏电保护动作等试验。试验仪器在有效期内，测试回路与图纸和设备编号一致。发现绝缘下降、接地不连续、相序错误或保护拒动时不得送电。",
    )
    body(
        document,
        "送电实行书面申请和统一指挥，清点人员和工器具，确认柜门、盖板、封堵、警示和消防条件。先空载检查母线和回路，再按区域逐级送电；每级送电后测量电压和相序，观察异响、异味和温升。负荷试运行期间记录主要设备电流和保护状态。",
    )
    heading(document, "（四）土建装饰验收标准", 2)
    body(
        document,
        "混凝土、砌体、屋面、地面、装饰和排水工程按相应分项验收。重点检查设备基础尺寸和强度、防水节点和渗漏、地面平整及坡向、墙面空鼓开裂、吊顶检修条件、沟槽盖板稳定性和排水通畅。隐蔽部位未经检查不得覆盖。",
    )

    heading(document, "六、质量保证措施", 1)
    heading(document, "（一）质量管理制度", 2)
    body(
        document,
        "项目执行班组自检、专业互检和项目专检的三级检查制度。施工员负责工序组织，质量员负责停检点和资料核验，技术负责人负责技术复核和重大质量问题处理。每道工序明确输入条件、过程记录和验收输出，未验收不得转序。",
    )
    body(
        document,
        "执行样板先行制度。桥架支架、钢管引下、电缆终端、接地跨接、防火封堵、墙地面收口和防水节点先做样板，经建设、监理和使用单位确认后推广。样板作为班组交底和实物验收的统一标准。",
    )
    body(
        document,
        "对清单、图纸和现场不一致的内容实行技术核定。任何材料代用、线路改向、设备接口调整和恢复做法变化均须书面确认。施工人员不得以现场方便为由改变保护接地、开关容量、母线连接、隔断方式和结构构造。",
    )
    heading(document, "（二）过程记录和资料管理", 2)
    body(
        document,
        "技术交底、作业票、材料报验、开箱、测量、隐蔽、检测、试验、调试、影像和交接记录与施工同步。影像资料显示日期、位置、工序和关键尺寸；测试记录写明仪器编号、有效期、环境条件和人员。资料不完整的工序不得视为完成。",
    )
    body(
        document,
        "竣工图根据批准的深化图和现场实测编制，准确反映电缆路径、回路编号、盘柜和设备接口、接地系统、燃气保留端口和隐蔽管线。设备说明书、合格证、保修资料和操作维护要点按系统整理并移交。",
    )
    heading(document, "（三）成品保护和工程移交", 2)
    body(
        document,
        "盘柜和设备安装后采取防尘、防潮和防碰撞保护，柜内不得存放工具。桥架盖板、电缆标识、接地跨接和防火封堵完成后设置保护标识。屋面防水、地面、墙面和吊顶完成后限制交叉作业，确需进入时铺设保护层并办理作业面移交。",
    )
    body(
        document,
        "工程完工后组织燃气遗留接口、电气系统、厨房设备、屋面防水、装饰和排水全专业联合检查。问题形成清单，明确责任人和完成日期，复验合格后办理移交。保修期内对施工质量问题及时响应，设备本体和甲供范围按合同界面处理。",
    )
    heading(document, "（四）质量问题预防和纠偏", 2)
    body(
        document,
        "项目部每周汇总实测和验收问题，区分设计接口、材料设备、施工工艺和成品保护原因。重复出现的问题由技术负责人组织专题分析，形成图示交底、样板或检查清单；同类区域在整改样板验收前暂停继续施工，防止问题批量扩散。",
    )
    body(
        document,
        "对盘柜母线扭矩、电缆绝缘、保护导体连续性、接地电阻、设备基础尺寸、防水搭接和排水坡度等关键指标实行实测实量。数据直接记录原始值，不只填写“合格”；超差项标明位置、整改方法和复验结果，原始记录不得涂改或事后集中补填。",
    )
    body(
        document,
        "质量问题整改遵循先查明原因、再确定措施、实施后复验的程序。涉及结构、保护整定、母线连接、燃气封堵和防水体系的整改必须由技术负责人审核，必要时征求设计、设备厂家或产权单位意见。未经复验关闭的问题不得进入竣工移交。",
    )

    heading(document, "七、安全环保措施", 1)
    heading(document, "（一）通用安全管理", 2)
    body(
        document,
        "所有人员进场接受厂级和项目级安全教育，特种作业人员持有效证件。每日班前会确认能源状态、作业票、监护人、防护用品、工器具和撤离路线。施工区域实行封闭管理，无关人员不得进入，作业条件变化时立即停工复核。",
    )
    body(
        document,
        "燃气、动火、盲板抽堵、临时用电、停送电、高处、吊装和开挖作业执行建设单位许可制度。作业票只在规定区域、时段和条件内有效，不得涂改、转借或跨区域使用。监护人不得同时从事影响监护的操作。",
    )
    heading(document, "（二）专项安全措施", 2)
    body(
        document,
        "燃气作业区严禁火种和非防爆电器，保持通风并按介质配置可燃或有毒气体检测仪。氮气置换同时监测氧含量，禁止人员进入可能积聚氮气的封闭或低洼空间。发生报警、异味或人员不适立即停止、撤离并报告。",
    )
    body(
        document,
        "停送电执行工作票和上锁挂牌程序，核对回路、切断电源、验电、必要时放电和接地，再开始作业。多人作业由负责人统一保管或控制锁具，任何人不得擅自拆除他人挂牌。临时用电采用分级配电和漏电保护，电缆避开积水、热源和机械碾压。",
    )
    body(
        document,
        "约5m高处的桥架、配管和拆除优先使用合格登高平台或脚手架。作业平台、护栏、踢脚板和上下通道经验收，人员正确系挂安全带，工具设置防坠绳。大风、雷雨、照明不足和脚下湿滑时停止露天或临边高处作业。",
    )
    body(
        document,
        "吊装前核实设备重量、重心、吊点、通道和地面承载，编制作业交底。起重机、吊索具和指挥人员符合要求，设置警戒区，执行试吊确认。设备下方和回转半径内严禁站人，临时停放使用垫木和防倾覆措施。",
    )
    body(
        document,
        "动火前清除可燃物，不能清除的采用防火毯隔离，设置灭火器、监火人和火花接收措施。燃气区域必须在可靠隔断、置换检测和产权单位许可后动火。作业结束后清理焊渣并持续检查余火，确认无复燃可能后撤离。",
    )
    body(
        document,
        "拆除、沟槽和开孔区域设置硬质围挡、临边防护和警示，洞口不能及时封闭时安装承载可靠的盖板并固定。交叉作业实行上下隔离，禁止在人员通道上方无防护施工。拆除物和工具不得抛掷。",
    )
    heading(document, "（三）食堂卫生、环境保护和季节施工", 2)
    body(
        document,
        "施工区与食品加工、储存和就餐区域采用连续封闭围挡，施工人员不得进入运营区。粉尘作业采用湿法、集尘和负压隔离，焊烟和异味及时排出室外。恢复营业前清除异物，完成吸尘、擦洗、消毒和通风，经使用单位确认。",
    )
    body(
        document,
        "废旧管道、电缆、金属、包装和建筑垃圾分类堆放并及时清运。废油、胶粘剂和污染吸附物按要求收集，不得进入排水系统。噪声较大的切割、钻孔安排在允许时段，运输车辆覆盖并清洁轮胎，保持厂区道路和消防通道畅通。",
    )
    body(
        document,
        "施工跨越武汉夏季、秋季和初冬。8月至9月重点防高温、暴雨和雷电，调整作息、配置饮水和防暑物资，雷雨时停止露天高处和吊装。11月至12月检查混凝土、砂浆、防水和涂料的施工温度，低温时采取保温养护；送电前检查设备是否受潮凝露。",
    )
    heading(document, "（四）安全检查和违章治理", 2)
    body(
        document,
        "项目部实行班组日检、项目周检和专项作业旁站检查。日检重点为围挡、通道、临电、消防、登高设施和工完清场；周检核对作业票、人员资质、检测记录、吊索具和整改闭环；燃气、动火、停送电和吊装由项目负责人或技术负责人现场确认开工条件。",
    )
    body(
        document,
        "检查发现一般隐患时明确责任人和完成时限，整改后由原检查人复查；发现能源隔离失效、气体报警、无证作业、保护装置缺失、人员进入吊物下方或高处无防护等重大隐患时立即停工撤人。未经复查批准，不得以口头承诺代替整改。",
    )
    body(
        document,
        "对重复违章和措施执行不到位的班组进行再教育和重新交底，必要时调整人员。隐患记录写明部位、事实、风险、整改和复验，不使用笼统表述。项目每周分析违章类型和高发区域，将结果用于下周计划、资源配置和旁站安排。",
    )

    heading(document, "危险因素分析及对策措施", 1)
    heading(document, "（一）辨识和评价方法", 2)
    body(
        document,
        "项目部按作业活动、能源介质、设备设施、施工环境和人员行为进行危险源辨识，综合考虑事故发生可能性、暴露频次和后果严重程度确定控制等级。燃气隔断置换、动火、停送电、高处、吊装和交叉作业按重大风险管理，实施项目负责人带班、作业许可、旁站监护和过程检测。",
    )
    body(
        document,
        "危险源清单在图纸会审、现场条件变化、甲供设备到货、施工转序和发生异常后动态更新。班组交底必须说明风险、后果、控制措施和应急动作，不能只签字不讲解。发现措施无法落实时，任何人员有权停止作业并报告。",
    )
    hazards = [
        ["1", "燃气介质或边界识别错误", "中毒、窒息、燃烧爆炸", "重大", "产权单位书面确认介质、压力、流向、阀门和影响用户；未确认不得作业。"],
        ["2", "停气或可靠隔断失效", "燃气泄漏、火灾爆炸", "重大", "授权人员操作；挂牌上锁；盲板或断开封堵；双人复核并书面交接。"],
        ["3", "氮气置换和放散", "缺氧窒息、人员不适", "重大", "设置警戒、强制通风、监测氧含量；放散口引至室外安全地点。"],
        ["4", "检测点或检测频次不足", "误判安全条件", "重大", "使用适配且校准有效的仪器；末端、低点和死角多点检测；复工复测。"],
        ["5", "燃气管道切割拆除", "火花、坠落、物体打击", "重大", "先支撑后切断、分段拆除；介质合格后选择机具；设置监护和消防。"],
        ["6", "停送电和回路识别错误", "触电、电弧、设备损坏", "重大", "工作票、上锁挂牌、验电、核相；统一指挥送电，保护未试验不得投运。"],
        ["7", "施工临时用电", "触电、短路、火灾", "较大", "分级配电、漏电保护、一机一闸；电缆防水防压；持证电工每日检查。"],
        ["8", "约5m高处桥架配管", "高处坠落、工具坠落", "重大", "合格平台或脚手架；验收后使用；安全带高挂低用；工具防坠和下方隔离。"],
        ["9", "盘柜母线和设备吊装", "起重伤害、倾覆挤压", "重大", "核重心和吊点；持证指挥；试吊；警戒；严禁人员进入吊物下方。"],
        ["10", "沟槽开挖和洞口", "坍塌、跌落、破坏管线", "较大", "探管试挖、临边防护、控制堆载、排水；洞口盖板固定并标识。"],
        ["11", "墙体和楼板开孔", "结构损伤、飞溅、粉尘", "较大", "结构确认、探测暗埋设施、湿法作业、护目和集尘；洞口及时封闭。"],
        ["12", "动火及防水热作", "火灾、灼伤、烟气", "重大", "办理动火票、清除可燃物、监火和灭火器；作业后复查余火。"],
        ["13", "多专业交叉作业", "物体打击、相互干扰", "较大", "作业面移交、上下隔离、错时施工、统一指挥；禁止无防护垂直交叉。"],
        ["14", "防水和胶粘材料", "火灾、中毒、皮肤伤害", "较大", "通风、远离火源、控制用量、佩戴防护用品；废料密闭回收。"],
        ["15", "食堂相邻区域施工", "食品污染、异物混入", "重大", "硬隔离、独立通道、吸尘消毒、工完清场；使用单位检查后恢复运营。"],
        ["16", "高温、暴雨、雷电及低温", "中暑、触电、滑跌、质量缺陷", "较大", "天气预警、调整作息、排水防雷、防潮检查和低温养护；恶劣天气停工。"],
    ]
    add_table(
        document,
        ["序号", "危险源辨识", "可能后果", "等级", "主要控制措施"],
        hazards,
        [0.9, 3.0, 2.7, 1.1, 6.8],
        8.5,
    )
    heading(document, "（二）重大风险专项控制", 2)
    body(
        document,
        "燃气泄漏和置换风险实行能源产权单位与施工项目部双重控制。能源状态由产权单位确认，施工边界由项目负责人确认，气体检测由经培训人员实施，监护人记录全过程。任何检测报警、异味、流量异常、阀门状态不明或隔断松动都视为条件失效，立即撤离和重新交接。",
    )
    body(
        document,
        "触电和电弧风险以能量隔离为核心。停电后仍须核对回路和验电，严禁仅凭指示灯、开关位置或他人通知判断无电。送电前执行人员清点、工器具清点、封板检查和保护试验四项确认，送电过程设置警戒并由一人统一下令。",
    )
    body(
        document,
        "高处和起重风险以可靠作业平台、人员防坠和吊装隔离为核心。登高车支腿、地面承载、护栏和应急下降功能开工前检查；吊装设备按实物重量选择吊具，设置溜绳控制摆动。高处和吊装作业禁止与下方通行或其他班组同时进行。",
    )
    body(
        document,
        "食品卫生风险以空间隔离、污染控制和恢复确认三道防线控制。围挡连续到顶或采取等效封闭，粉尘作业设置局部负压和吸尘，施工工具材料不得跨区混用。每日收工和阶段移交均由使用单位检查，发现粉尘、异物或异味时不得恢复经营。",
    )
    body(
        document,
        "动火和火灾风险以清除可燃物、隔离火花、连续监护和作业后复查为核心。燃气管线附近动火必须与停气、盲板和检测条件联动，任何一项失效即终止动火。防水热熔、焊接和切割不能同时布置在同一垂直作业面，灭火器材放在人员可迅速取得的位置。",
    )
    body(
        document,
        "沟槽、开孔和结构风险以先探测、后拆改为核心。开挖前使用图纸、探测和人工试挖交叉确认，发现不明管线立即保护并报告。墙体和楼板开孔按批准位置施工，出现异常裂缝、振动、渗水或构件变形时停止作业，由技术人员复核。",
    )
    body(
        document,
        "多专业交叉风险以作业面移交和空间隔离为核心。项目日计划明确每个区域的主导班组，未经协调不得自行进入他人作业面。高处作业、吊装和动火区域设置排他性警戒，电气试验和送电期间停止相关回路及设备周边的其他施工。",
    )
    body(
        document,
        "季节和环境风险实行预警控制。项目每天关注气象信息，高温、雷电、暴雨和大风预警纳入班前会。雨后复查登高设备、脚手基础、沟槽和临电，低温时复核材料施工条件；食品加工区相邻作业还应关注通风方向，避免焊烟、粉尘和异味倒灌。",
    )

    heading(document, "八、资源需求计划", 1)
    heading(document, "（一）劳动力计划", 2)
    body(
        document,
        "劳动力按施工阶段动态配置，燃气作业、电气安装、起重、登高和试验人员必须具备相应资格。高峰期安排土建、电气和设备班组平行作业，但通过分区和日计划避免相互干扰。人员变更时重新进行入场教育和专项交底。",
    )
    labor = [
        ["项目管理", "项目经理、技术负责人、施工员、安全员、质量员、材料员、资料员", "7人", "全工期"],
        ["燃气停用拆除", "管道作业、检测、监护及产权单位接口人员", "约6人", "8月下旬至9月上旬"],
        ["电气安装", "盘柜母线、桥架配管、电缆接地班组", "8～10人", "9月至11月"],
        ["土建装饰", "拆除、混凝土、砌筑、防水、地面和装饰班组", "10～12人", "8月至11月"],
        ["厨房设备", "设备安装和电气末端接线人员", "约6人", "10月下旬至11月"],
        ["检测调试", "电气试验、设备调试及使用单位配合人员", "4～6人", "11月下旬至12月"],
    ]
    add_table(document, ["工种/班组", "人员构成", "计划投入", "主要阶段"], labor, [3.0, 6.2, 2.0, 3.3], 9.2)
    heading(document, "（二）主要机械和工器具", 2)
    equipment = [
        ["25t汽车吊", "2台班", "盘柜、母线和设备吊装"],
        ["10t平板车", "2台班", "设备和材料运输"],
        ["12m登高作业平台", "5台班", "高处桥架、配管和拆除"],
        ["氮气瓶组及汇流排", "按专项配置", "燃气置换，数量按计算和检测调整"],
        ["气体检测仪和氧检测仪", "不少于2套", "适配介质、校准有效，配置备用电池"],
        ["液压弯管器、压接钳、牵引滑轮", "按两作业面配置", "钢管、电缆及终端施工"],
        ["绝缘、接地、相序和扭矩仪表", "各不少于1套", "电气试验和连接质量检查"],
        ["灭火器、防火毯、吸尘设备", "按作业面配置", "动火消防、粉尘和食堂卫生控制"],
    ]
    add_table(document, ["机械/器材", "数量", "主要用途"], equipment, [4.2, 3.0, 7.3], 9.5)
    heading(document, "（三）材料供应和保管", 2)
    body(
        document,
        "项目编制材料设备到货计划，优先保证盘柜、封闭母线、电缆、桥架和甲供厨房设备等长周期物资。每批材料到货后先报验再领用，现场只保留满足近期施工的数量。剩余材料及时退库，防止在运营通道和屋面集中堆载。",
    )

    heading(document, "九、应急预案", 1)
    heading(document, "（一）编制目的、适用范围和原则", 2)
    body(
        document,
        "本预案用于本工程施工期间发生或可能发生的燃气泄漏、中毒窒息、火灾爆炸、触电、高处坠落、起重和机械伤害、坍塌及物体打击、极端天气和食品卫生污染等事件。目的在于迅速控制危险源、保护人员生命、减少财产和环境损失，并与建设单位厂区应急体系有效衔接。",
    )
    body(
        document,
        "应急处置坚持“人员生命优先、先控制后处置、统一指挥、分级响应、专业救援、保护现场”的原则。现场人员发现险情有权立即停工和报警，不得因赶工、保护设备或未经授权抢险而冒险进入危险区域。",
    )
    heading(document, "（二）应急组织和职责", 2)
    body(
        document,
        "项目经理担任现场应急总指挥，负责启动预案、资源调配、对外报告和与建设单位应急体系衔接；技术负责人负责判断能源和设备状态、提出隔离及技术处置措施；安全员负责报警、警戒、人员清点、疏散和现场记录；施工员组织停工和撤离；班组长清点本班人员并报告。",
    )
    body(
        document,
        "应急小组下设抢险隔离组、疏散警戒组、医疗救护组、通讯联络组和后勤保障组。未经专业训练的人员不得进入燃气、缺氧、带电、火场或不稳定结构区域抢险。专业消防、医疗、能源和救援力量到场后，现场指挥权按建设单位应急制度移交。",
    )
    emergency_contacts = [
        ["建设单位总调度/能源管理", "开工前填写", "燃气、电力、生产及厂区应急协调"],
        ["项目经理", "开工前填写", "现场总指挥、对外报告"],
        ["安全员", "开工前填写", "警戒、疏散、清点和记录"],
        ["技术负责人", "开工前填写", "能源状态和技术处置"],
        ["消防报警", "119或厂区消防电话", "火灾、爆炸报警"],
        ["医疗急救", "120或厂区急救电话", "人员伤害救护"],
    ]
    add_table(document, ["联系人/单位", "电话", "职责"], emergency_contacts, [4.0, 4.0, 6.5], 9.5)
    heading(document, "（三）监测预警、信息报告和响应程序", 2)
    body(
        document,
        "燃气检测仪、氧含量检测仪、临时用电巡检、天气预警、火灾报警和人员观察构成现场监测体系。出现气体报警、异味、阀门或隔断异常、人员头晕、设备冒烟、漏电保护动作、高处设施不稳、沟槽积水或雷电预警时，立即停止相关作业并报告安全员。",
    )
    body(
        document,
        "事故报告内容包括时间、地点、作业活动、人员伤害、能源状态、现场危险、已采取措施和所需支援。首报可简明但不得迟报、漏报或瞒报，后续随情况更新。报警时安排人员在厂区入口和路口引导消防、医疗或专业救援车辆。",
    )
    body(
        document,
        "响应基本程序为：发现异常并呼救→立即停工和切断可安全切断的能源→人员向上风侧或指定路线撤离→设置警戒并清点人员→报警和逐级报告→专业人员控制危险源→救治伤员→保护现场→调查整改→验收批准后复工。",
    )
    body(
        document,
        "疏散路线在开工布置图和现场围挡上标明，至少保持一条不经过燃气放散点、吊装回转区和临时配电设施的通道。人员撤离后到指定集合点按班组清点，班组长报告应到、实到和失联人数。未经总指挥许可，任何人不得返回警戒区寻找物品。",
    )
    body(
        document,
        "警戒范围根据介质、风向、火势、带电范围、吊物状态和结构稳定性动态调整。警戒人员站在安全位置，阻止无关人员和车辆进入，同时为消防、医疗和能源抢险留出通道。涉及食堂运营时立即通知使用单位停止受影响区域的加工和人员进入。",
    )
    body(
        document,
        "现场急救坚持先脱离危险、再判断伤情、及时呼救和连续观察。伤员没有继续暴露风险前不得急于搬运；怀疑脊柱、骨盆、内脏或严重挤压伤时由专业急救人员指导转运。项目安排人员携带伤员身份、事故介质和初步处置信息陪同就医。",
    )
    heading(document, "（四）燃气泄漏、中毒和缺氧窒息处置", 2)
    body(
        document,
        "发现燃气报警、异味、异常声响或人员头晕恶心时，立即停止作业，禁止启闭可能产生火花的电器，禁止使用非防爆通讯设备靠近泄漏点。人员沿上风侧迅速撤离，扩大警戒并报告产权单位，未经授权不得操作阀门或进入查漏。",
    )
    body(
        document,
        "产权单位组织切断气源、加强通风和连续检测。若人员倒地，救援者必须佩戴与介质相适应的呼吸防护并采用安全绳或专业救援方式，严禁无防护盲目施救。伤员移至空气新鲜处，保持呼吸道通畅，必要时实施心肺复苏并立即送医。",
    )
    body(
        document,
        "氮气置换造成缺氧时，同样执行先通风、检测、专业防护后救援。普通口罩、防尘面具不能用于缺氧或有毒气体环境。现场检测达到安全条件、泄漏点完成修复且经产权单位批准后，方可解除警戒。",
    )
    heading(document, "（五）火灾和爆炸处置", 2)
    body(
        document,
        "发生初起火灾时立即报警并停止相关能源作业。在确保人员安全和具备适用灭火器材的前提下进行初期扑救；燃气火灾不得在气源未能可靠切断时盲目灭火，以防形成未燃烧气体积聚。火势扩大、容器受热或存在爆炸危险时立即撤离。",
    )
    body(
        document,
        "电气设备未确认断电前使用适用于带电设备的灭火器，严禁直接用水。防水材料、胶粘剂或包装物着火时，移走周边可燃物并阻止火势向食堂运营区和电缆通道蔓延。消防人员到场后提供燃气、电源、材料和人员清点信息。",
    )
    body(
        document,
        "火灾扑灭后继续监测复燃、燃气和电气风险，不得擅自恢复供气供电。事故区域由有关单位勘查，受热盘柜、电缆、支架、防水层和结构经检查评估，隐患整改并验收后方可恢复。",
    )
    heading(document, "（六）触电和电弧伤害处置", 2)
    body(
        document,
        "发现触电时首先切断电源或使用绝缘工具使伤员脱离电源，施救者不得直接接触带电人员。高处触电应同时防止伤员坠落。电源切断后检查意识、呼吸和脉搏，立即拨打急救电话，呼吸心跳停止时持续进行心肺复苏直至专业人员接替。",
    )
    body(
        document,
        "电弧灼伤人员转移到安全地点，覆盖清洁敷料，不随意剥离粘连衣物或涂抹药物。事故回路保持隔离，未经调查不得重新送电。项目部检查回路识别、上锁挂牌、绝缘工具和保护装置，查明原因并形成整改。",
    )
    heading(document, "（七）高处坠落、物体打击和坍塌处置", 2)
    body(
        document,
        "发生高处坠落或物体打击后立即停止上方作业，封锁区域，防止工具、构件或平台再次坠落。伤员疑似颈椎、脊柱、骨盆或内脏损伤时不得随意搬动，由急救人员指导固定和转运；出现大出血时采用加压包扎等方法止血。",
    )
    body(
        document,
        "沟槽、围挡、脚手或临时支撑发生坍塌征兆时迅速撤离，切断附近水电和机械动力，禁止未经评估进入。救援前确认结构稳定、二次坍塌和地下管线风险，采用专业支撑和挖掘方式，严禁大型机械直接盲目开挖。",
    )
    heading(document, "（八）起重和机械伤害处置", 2)
    body(
        document,
        "吊装中出现设备倾斜、吊索异常、地基沉陷或人员闯入时，指挥人员立即发出停止信号，操作人员将吊物在安全条件下缓慢落地。吊物无法落地时扩大警戒，保持设备稳定并联系专业人员处理，任何人不得进入吊物下方。",
    )
    body(
        document,
        "发生夹挤、卷入或切割伤害时先停机并切断能源，不得反向启动设备强行解救。根据伤情止血、包扎和固定，严重挤压伤由专业急救人员处理。事故设备、吊具和操作记录保留，检查合格前不得复用。",
    )
    heading(document, "（九）极端天气、中暑和食堂污染处置", 2)
    body(
        document,
        "收到暴雨、雷电或大风预警时停止露天高处、吊装和屋面作业，固定材料、切断室外临电、疏通排水并撤离到安全区域。沟槽进水或脚手基础受冲刷后必须重新检查验收。设备受潮后干燥并复测绝缘，禁止直接送电。",
    )
    body(
        document,
        "发现中暑人员立即移至阴凉通风处，解开衣物、补充适量水分并降温；出现意识障碍、抽搐或高热时立即送医。高温作业调整时段和轮换，燃气监护、盘柜内作业和屋面施工避免长时间连续暴露。",
    )
    body(
        document,
        "发现粉尘、金属屑、油污、异味或施工废水进入食堂运营区时，立即停止相关作业并封闭污染范围，通知使用单位。污染物按食品卫生要求清除，对受影响区域清洗、消毒和通风，必要时由专业机构检测；未经使用单位确认不得恢复运营。",
    )
    heading(document, "（十）应急物资、培训演练和恢复", 2)
    body(
        document,
        "现场配置灭火器、防火毯、急救箱、担架、警戒带、应急照明、绝缘工具、吸附材料、气体和氧含量检测仪等物资。安全员每周检查数量、位置、有效期和完好状态，任何人不得挪作施工用途。燃气专业防护和救援装备由产权单位或专业救援力量配置。",
    )
    body(
        document,
        "开工后组织一次综合应急演练，燃气停用前组织泄漏和人员不适专项桌面或现场演练，送电前组织触电和停电应急演练。演练检验报警、撤离、人员清点、现场隔离和信息报告，发现问题立即修订预案和交底。",
    )
    body(
        document,
        "事故或险情处置后由建设单位和项目部组织调查，确认危险源已经消除、设备设施检查合格、作业票和技术措施重新批准、人员完成再教育后方可复工。涉及人身伤害、火灾、燃气或生产影响的事件按规定报告，不得擅自处理或隐瞒。",
    )
    body(
        document,
        "应急演练和实际响应均应记录时间线，包括异常发现、停工、报警、撤离、人员清点、专业力量到场和危险解除等节点。项目部对照预案检查通讯是否畅通、路线是否可用、器材是否适配以及各岗位是否清楚职责，形成书面评估和改进清单。",
    )
    body(
        document,
        "应急物资消耗或使用后立即补充，气体检测仪、绝缘工具和急救用品按规定维护校验。施工阶段或作业面变化时重新检查应急通道、集合点和报警电话，防止围挡调整、材料堆放或临时设施占用原有疏散条件。",
    )
    body(
        document,
        "复工前除完成设备和环境检查外，还应向全体相关人员通报事件原因、整改措施和新的作业边界。涉及产权单位能源系统、食堂运营或公共安全的，由建设单位组织联合验收；项目部不得自行宣布恢复供气、送电或营业。",
    )

    heading(document, "十、施工进度计划", 1)
    body(
        document,
        f"计划工期为{START}至{FINISH}，共144日历天。进度控制以燃气停用窗口、盘柜母线和电缆到货、甲供厨房设备接口、土建作业面和送电试运行条件为关键约束，实行周计划、三日滚动计划和日协调。",
    )
    body(
        document,
        "关键线路为深化和设备策划→燃气停用置换→原管道拆除→基础和沟槽→桥架盘柜母线→电缆敷设→厨房设备接电→电气试验→系统联调→联合验收。建筑装饰恢复与电气安装在分区移交后穿插进行，但防水、隐蔽和试验条件不得因赶工压缩。",
    )
    body(
        document,
        "8月10日至8月23日完成进场、会审、综合测绘、设备参数核对、材料报审和施工隔离；8月24日至9月10日利用批准的停气窗口完成阀门关闭、可靠隔断、氮气置换、检测和原燃气管道拆除。停气日期如受生产组织调整，后续关键工作同步滚动修正。",
    )
    body(
        document,
        "8月下旬至9月底形成沟槽、基础、墙洞和结构作业面，并开始建筑恢复；9月中旬至10月下旬完成桥架、钢管、接地、盘柜和封闭母线安装；10月至11月上旬完成动力电缆敷设、终端制作和防火封堵。各区域按移交条件穿插，不在同一作业面无序叠加。",
    )
    body(
        document,
        "10月下旬至11月20日完成甲供厨房设备就位、固定和电气末端连接，11月底前完成屋面防水、地面、墙面、吊顶、门窗和排水恢复。11月21日至12月22日组织绝缘、接地、相序、保护动作、单机和系统联调，并结合食堂条件安排负荷试运行。",
    )
    body(
        document,
        "12月23日至12月31日集中完成缺陷整改、清洁消毒、竣工图和资料整理、联合验收及工程移交。项目部每周核对关键线路和剩余工作量，出现材料延误、作业窗口变化或验收问题时优先调整非关键工作的作业顺序，保证安全和质量条件不被压缩。",
    )


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, OUTPUT)
    document = Document(OUTPUT)
    remove_body(document)
    first = document.sections[0]
    configure_section(first, landscape=False)
    set_header_footer(first)

    add_cover(document)
    add_approval(document)
    add_toc(document)
    document.add_page_break()
    add_main_content(document)

    landscape = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape, landscape=True)
    set_header_footer(landscape)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(SCHEDULE_RENDER), width=Cm(19.2))

    document.core_properties.title = PROJECT + "施工方案（修订版）"
    document.core_properties.subject = "食堂气改电、燃气停用拆除、低压配电及土建恢复施工"
    document.core_properties.author = COMPANY
    document.core_properties.comments = "严格参照条材厂施工方案模板版式编制"
    document.save(OUTPUT)
    with zipfile.ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
        assert "word/styles.xml" in archive.namelist()
    print(OUTPUT)


if __name__ == "__main__":
    build()
