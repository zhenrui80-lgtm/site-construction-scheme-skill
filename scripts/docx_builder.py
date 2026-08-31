#!/usr/bin/env python3
"""Stable DOCX helpers for project-specific construction-scheme builders.

Project builders should import this module from the skill. They must not import
another project's build script because doing so leaks constants and formatting.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BODY_FONT = "方正仿宋_GB2312"
ACCENT = "1F4E79"
GRAY = "666666"
ALT_FILL = "F6F9FC"


def set_run_font(run, size=None, bold=None, color="000000", font_name=BODY_FONT):
    """Set every Word font slot without changing unspecified point size/bold."""
    run.font.name = font_name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    return run


def replace_paragraph_text_preserve_runs(paragraph, text):
    """Replace text in retained template paragraphs without flattening pPr/rPr."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return paragraph
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""
    return paragraph


def replace_exact_paragraph(document, old_text, new_text, occurrence=1):
    hits = 0
    for paragraph in document.paragraphs:
        if paragraph.text == old_text:
            hits += 1
            if hits == occurrence:
                replace_paragraph_text_preserve_runs(paragraph, new_text)
                return paragraph
    raise ValueError(f"Paragraph not found at occurrence {occurrence}: {old_text}")


def body_blocks(document):
    return [
        child
        for child in document.element.body
        if child.tag in {qn("w:p"), qn("w:tbl")}
    ]


def find_body_boundary(document, marker, occurrence=2):
    hits = 0
    for index, block in enumerate(body_blocks(document)):
        text = "".join(block.xpath(".//w:t/text()"))
        if marker in text:
            hits += 1
            if hits == occurrence:
                return index
    raise ValueError(f"Body marker not found at occurrence {occurrence}: {marker}")


def remove_existing_body(document, marker="一、编制依据及工程概况", occurrence=2):
    """Remove template body while leaving the frozen front-matter blocks intact."""
    blocks = body_blocks(document)
    boundary = find_body_boundary(document, marker, occurrence)
    for block in blocks[boundary:]:
        document.element.body.remove(block)
    return boundary


def ensure_style(
    document,
    name,
    size,
    bold=False,
    color="000000",
    alignment=None,
    first_indent=False,
    font_name=BODY_FONT,
):
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font_name)
    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if alignment is not None:
        fmt.alignment = alignment
    if first_indent:
        fmt.first_line_indent = Pt(size * 2)
    return style


def configure_body_styles(document, prefix="SCV2", font_name=BODY_FONT, accent=ACCENT):
    """Create body-only styles; retained template styles are never normalized."""
    ensure_style(document, f"{prefix}_Chapter", 16, True, accent, WD_ALIGN_PARAGRAPH.LEFT, font_name=font_name)
    ensure_style(document, f"{prefix}_Section", 14, True, "000000", WD_ALIGN_PARAGRAPH.LEFT, font_name=font_name)
    ensure_style(document, f"{prefix}_Subsection", 12, True, "000000", WD_ALIGN_PARAGRAPH.LEFT, font_name=font_name)
    ensure_style(document, f"{prefix}_Minor", 11.5, True, accent, WD_ALIGN_PARAGRAPH.LEFT, font_name=font_name)
    ensure_style(document, f"{prefix}_Body", 12, False, "000000", WD_ALIGN_PARAGRAPH.JUSTIFY, True, font_name)
    ensure_style(document, f"{prefix}_NoIndent", 12, False, "000000", WD_ALIGN_PARAGRAPH.JUSTIFY, False, font_name)
    ensure_style(document, f"{prefix}_Caption", 10.5, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, False, font_name)
    ensure_style(document, f"{prefix}_Note", 10.5, False, GRAY, WD_ALIGN_PARAGRAPH.JUSTIFY, False, font_name)
    return prefix


def add_heading(document, text, level=1, page_break=False, prefix="SCV2", font_name=BODY_FONT, accent=ACCENT):
    style = {
        1: f"{prefix}_Chapter",
        2: f"{prefix}_Section",
        3: f"{prefix}_Subsection",
        4: f"{prefix}_Minor",
    }[level]
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break
    sizes = {1: 16, 2: 14, 3: 12, 4: 11.5}
    colors = {1: accent, 2: "000000", 3: "000000", 4: accent}
    set_run_font(paragraph.add_run(text), sizes[level], True, colors[level], font_name)
    return paragraph


def add_body(document, text, bold_prefix=None, prefix="SCV2", font_name=BODY_FONT, accent=ACCENT):
    paragraph = document.add_paragraph(style=f"{prefix}_Body")
    paragraph.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(paragraph.add_run(bold_prefix), 12, True, accent, font_name)
        set_run_font(paragraph.add_run(text[len(bold_prefix):]), 12, False, "000000", font_name)
    else:
        set_run_font(paragraph.add_run(text), 12, False, "000000", font_name)
    return paragraph


def add_paragraphs(document, heading, paragraphs, level=2, **kwargs):
    add_heading(document, heading, level=level, **kwargs)
    for text in paragraphs:
        add_body(document, text, **{key: value for key, value in kwargs.items() if key in {"prefix", "font_name", "accent"}})


def add_note(document, text, prefix="SCV2", font_name=BODY_FONT):
    paragraph = document.add_paragraph(style=f"{prefix}_Note")
    set_run_font(paragraph.add_run(text), 10.5, False, GRAY, font_name)
    return paragraph


def _omml_math_run(text, size=11, color="000000", font_name="Times New Roman"):
    run = OxmlElement("m:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), font_name)
    run_properties.append(fonts)
    for tag in ("w:sz", "w:szCs"):
        size_element = OxmlElement(tag)
        size_element.set(qn("w:val"), str(int(size * 2)))
        run_properties.append(size_element)
    if color:
        color_element = OxmlElement("w:color")
        color_element.set(qn("w:val"), color)
        run_properties.append(color_element)
    run.append(run_properties)
    text_element = OxmlElement("m:t")
    text_element.text = text
    run.append(text_element)
    return run


def _omml_subscript(base, subscript, size=11, color="000000", font_name="Times New Roman"):
    node = OxmlElement("m:sSub")
    node.append(OxmlElement("m:sSubPr"))
    base_element = OxmlElement("m:e")
    base_element.append(_omml_math_run(base, size, color, font_name))
    sub_element = OxmlElement("m:sub")
    sub_element.append(_omml_math_run(subscript, size * 0.78, color, font_name))
    node.extend((base_element, sub_element))
    return node


def _append_linear_omml(parent, equation, size=11, color="000000", font_name="Times New Roman"):
    pattern = re.compile(r"([A-Za-z]+)_\{?([A-Za-z0-9,]+)\}?")
    cursor = 0
    for match in pattern.finditer(equation):
        if match.start() > cursor:
            parent.append(_omml_math_run(equation[cursor:match.start()], size, color, font_name))
        parent.append(_omml_subscript(match.group(1), match.group(2), size, color, font_name))
        cursor = match.end()
    if cursor < len(equation):
        parent.append(_omml_math_run(equation[cursor:], size, color, font_name))


def add_omml_equation(
    document,
    equation,
    explanation=None,
    size=11,
    color="000000",
    font_name="Times New Roman",
    prefix="SCV2",
):
    """Insert an editable Word equation with Times New Roman math runs.

    Use ASCII identifiers with optional ``_{subscript}`` notation and Unicode
    operators supplied by verified calculations. Never paste rendered formula
    screenshots or use plain-text formulas for a lifting calculation appendix.
    """
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    math_paragraph = OxmlElement("m:oMathPara")
    math_properties = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "center")
    math_properties.append(justification)
    math_paragraph.append(math_properties)
    math = OxmlElement("m:oMath")
    _append_linear_omml(math, equation, size, color, font_name)
    math_paragraph.append(math)
    paragraph._p.append(math_paragraph)
    if explanation:
        note = document.add_paragraph(style=f"{prefix}_Note")
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(note.add_run(explanation), 10.5, False, GRAY, BODY_FONT)
    return paragraph


def add_caption(document, text, prefix="SCV2", font_name=BODY_FONT):
    paragraph = document.add_paragraph(style=f"{prefix}_Caption")
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), 10.5, False, GRAY, font_name)
    return paragraph


def add_picture(document, path, width_inches, caption=None, prefix="SCV2", font_name=BODY_FONT):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = bool(caption)
    paragraph.add_run().add_picture(str(Path(path)), width=Inches(width_inches))
    if caption:
        add_caption(document, caption, prefix, font_name)
    return paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="7A8792", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def dxa(cm_value):
    return int(cm_value * 1440 / 2.54)


def set_table_geometry(table, widths_cm: Sequence[float]):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(dxa(value) for value in widths_cm)))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(dxa(width)))
        grid.append(column)
    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)
        for cell in table.columns[index].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(dxa(width)))


def add_styled_table(
    document,
    title,
    headers,
    rows: Iterable[Sequence[object]],
    widths_cm,
    font_size=9.0,
    header_fill="245A86",
    alternate_fill=ALT_FILL,
    prefix="SCV2",
    font_name=BODY_FONT,
):
    caption = document.add_paragraph(style=f"{prefix}_NoIndent")
    caption.paragraph_format.keep_with_next = True
    set_run_font(caption.add_run(title), 10.5, True, header_fill, font_name)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    set_table_geometry(table, widths_cm)
    set_table_borders(table)
    repeat_header(table.rows[0])
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(str(text)), font_size, True, "FFFFFF", font_name)
    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell, value in zip(row.cells, values):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 0:
                set_cell_shading(cell, alternate_fill)
            paragraph = cell.paragraphs[0]
            set_run_font(paragraph.add_run("" if value is None else str(value)), font_size, False, "000000", font_name)
    return table


def add_continuous_body_section(document, top=1.5, bottom=1.5, left=1.6, right=1.6):
    """Add after the first body heading so frozen front-matter block count is stable."""
    section = document.add_section(WD_SECTION.CONTINUOUS)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    return section


def add_landscape_section(document, top=0.8, bottom=0.8, left=0.8, right=0.8):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    return section


def set_core_properties(document, title=None, subject=None, author=None):
    properties = document.core_properties
    if title is not None:
        properties.title = title
    if subject is not None:
        properties.subject = subject
    if author is not None:
        properties.author = author


def copy_template(template_path, output_path):
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(Path(template_path).read_bytes())
    return Document(output)


def clone_section_properties(document):
    """Return a defensive copy for callers that need before/after assertions."""
    return [copy.deepcopy(section._sectPr) for section in document.sections]
