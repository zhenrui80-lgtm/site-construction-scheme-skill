from pathlib import Path
import fitz
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(__file__).with_name("热轧厂操控中心景观提升项目施工组织设计_武汉兴达建筑安装工程有限公司_排版修订版.docx")
BASE = Path(__file__).resolve().parent
DRAWING_ROOT = BASE / "drawings"
FIGURE_DIR = BASE / "figures"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text_left(cell, text, bold=False, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_paragraph_border(paragraph, fill="1F4E79"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_doc_watermark(section, text="武钢兴达"):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + "    " + text + "    " + text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(210, 210, 210)


def setup_header_footer(section):
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("热轧厂操控中心景观提升项目施工组织设计")
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run("- ")
    rr.font.size = Pt(9)
    add_field(fp, "PAGE")
    rr = fp.add_run(" -")
    rr.font.size = Pt(9)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    if level == 1:
        set_paragraph_border(p, "EAF3F8")
    r = p.add_run(text)
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    if level == 1:
        r.font.color.rgb = RGBColor(31, 78, 121)
    r.font.size = Pt(14 if level == 1 else 12)
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    return p


def add_list(doc, items):
    for text in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run("（" + text.split("、", 1)[0] + "）" + text.split("、", 1)[1] if "、" in text[:4] else text)
        r.font.name = "宋体"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        r.font.size = Pt(10.5)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    return table


def find_pdf(name_part):
    matches = [p for p in DRAWING_ROOT.rglob("*.pdf") if name_part in p.name]
    if not matches:
        raise FileNotFoundError(name_part)
    return matches[0]


def render_figure(name, pdf_name_part, page_no, crop_mode="main"):
    FIGURE_DIR.mkdir(exist_ok=True)
    pdf = find_pdf(pdf_name_part)
    doc = fitz.open(pdf)
    page = doc[page_no]
    rect = page.rect
    if crop_mode == "main":
        # CAD PDFs usually carry title blocks and frames around the drawing.
        # This crop keeps the engineering content and removes the busiest margins.
        clip = fitz.Rect(rect.x0 + 70, rect.y0 + 60, rect.x1 - 330, rect.y1 - 130)
    elif crop_mode == "wide":
        clip = fitz.Rect(rect.x0 + 90, rect.y0 + 60, rect.x1 - 360, rect.y1 - 130)
    elif crop_mode == "full":
        clip = fitz.Rect(rect.x0 + 20, rect.y0 + 20, rect.x1 - 20, rect.y1 - 35)
    else:
        clip = rect
    pix = page.get_pixmap(matrix=fitz.Matrix(2.2, 2.2), clip=clip, alpha=False)
    out = FIGURE_DIR / f"{name}.png"
    if out.exists():
        out.unlink()
    pix.save(out)
    return out


def add_figure(doc, image_path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(15.4))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(9)
    r.italic = True
    return p


figures = {
    "plan": render_figure("01_园建平面图摘录", "02 LP-01平面图", 0, "main"),
    "water": render_figure("02_水景详图摘录", "03 LD-1.1~1.5 水景详图", 0, "main"),
    "planting": render_figure("03_绿化种植图摘录", "D-tree", 0, "main"),
    "drainage": render_figure("04_给排水图摘录", "01给排水", 0, "main"),
    "electric": render_figure("05_电气图摘录", "热轧电气ZONG", 0, "main"),
}


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.9)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.0)
sec.different_first_page_header_footer = True
setup_header_footer(sec)

style = doc.styles["Normal"]
style.font.name = "宋体"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.font.size = Pt(10.5)

for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("热轧厂操控中心景观提升项目")
r.bold = True
r.font.name = "黑体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
r.font.size = Pt(22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("施工组织设计")
r.bold = True
r.font.name = "黑体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
r.font.size = Pt(26)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("（报审版）")
r.font.name = "宋体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
r.font.size = Pt(14)

for _ in range(7):
    doc.add_paragraph()
cover_table = doc.add_table(rows=5, cols=2)
cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
cover_table.style = "Table Grid"
cover_rows = [
    ("工程名称", "热轧厂操控中心景观提升项目"),
    ("施工单位", "武汉兴达建筑安装工程有限公司"),
    ("计划工期", "2026年5月18日开工，2026年6月30日完工"),
    ("编制日期", "2026年5月"),
    ("编制/审核/批准", "编制：              审核：              批准："),
]
for row, (k, v) in zip(cover_table.rows, cover_rows):
    set_cell_text(row.cells[0], k, True)
    set_cell_text_left(row.cells[1], v, False, 11)
    row.cells[0].width = Cm(3.2)
    row.cells[1].width = Cm(10.8)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n单位名称（盖章）：武汉兴达建筑安装工程有限公司")
r.font.name = "宋体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
r.font.size = Pt(12)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("产权厂审批意见")
r.bold = True
r.font.name = "黑体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
r.font.size = Pt(18)
doc.add_paragraph()
approval = doc.add_table(rows=1, cols=1)
approval.alignment = WD_TABLE_ALIGNMENT.CENTER
approval.style = "Table Grid"
cell = approval.cell(0, 0)
cell.height = Cm(15)
cell.text = "\n\n\n\n\n\n\n\n\n\n\n签字：\n\n（盖章）\n\n年    月    日"
for para in cell.paragraphs:
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT if "签字" in para.text or "盖章" in para.text or "年" in para.text else WD_ALIGN_PARAGRAPH.LEFT
    for run in para.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(12)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("目  录")
r.bold = True
r.font.name = "黑体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
r.font.size = Pt(18)
toc_items = [
    "一、编制说明",
    "二、工程概况",
    "三、图纸摘录",
    "四、施工部署",
    "五、施工准备",
    "六、主要施工方法",
    "七、进度计划",
    "八、质量保证措施",
    "九、安全文明施工措施",
    "十、环境保护与成品保护",
    "十一、雨季及高温施工措施",
    "十二、资料管理与验收",
    "十三、劳动力与主要资源计划",
    "十四、应急与协调措施",
    "十五、专项技术控制要点",
    "十六、结束语",
]
for item in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(3.2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(item)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(11)
doc.add_page_break()

add_heading(doc, "一、编制说明", 1)
add_para(doc, "本施工方案依据热轧厂操控中心景观提升项目施工图CAD及PDF文件、工程量清单、现行施工验收规范、厂区施工管理要求以及武汉兴达建筑安装工程有限公司类似工业厂区景观、道路、给排水、电气和绿化工程施工经验编制。方案重点围绕园林建筑、给排水、电气、绿化四个专业组织施工，并将水景工程作为本项目的关键控制部位进行专项安排。")
add_para(doc, "本工程位于热轧厂操控中心周边，施工内容既包含景观展示界面的提升，又包含道路铺装、草阶、LOGO与宣传设施、水景结构及饰面、雨水与给水系统、景观照明、苗木移栽与新植等多专业交叉作业。项目体量不宜按单一园林绿化项目组织，应按厂区综合改造工程管理，重视既有道路、管线、建筑外立面、生产通行和现场文明形象的协调。")
add_para(doc, "施工组织坚持“先地下、后地上；先结构、后装饰；先管线、后铺装；先乔木、后灌木地被；先样板、后展开”的原则。拆除、沟槽、基础、水景结构、防水闭水、机电管线、铺装收边、绿化种植、灯具调试和成品保护应形成闭合流程，避免重复开挖、污染饰面和影响厂区正常通行。")

add_heading(doc, "二、工程概况", 1)
add_para(doc, "工程名称为热轧厂操控中心景观提升项目，施工单位按武汉兴达建筑安装工程有限公司名义组织。计划于2026年5月18日开工，2026年6月30日完工，总工期44日历天。施工范围包括绿化工程、园建工程、给排水工程、电气工程四个部分，其中园建工程内含新建人行道、草阶、字体及排水沟、水景工程、拆除工程、景石安装以及给排水和电气配套土建恢复等内容。")
add_para(doc, "清单显示，绿化工程包括现状树木移除39株、原有竹子移除64平方米，苏铁、桂花、圆柏球等移栽，新植罗汉松、点景朴树、朴树、乌桕、香橼、樱花、桂花、杨梅、紫薇、红梅、红枫及球类灌木，法国冬青、南天竹、小叶栀子、金森女贞、春鹃、龟甲冬青、金丝桃、大叶栀子、黄金菊、海桐、花叶络石、常夏石竹、红继木、麦冬等地被灌木，以及铺种草皮1400平方米、整理绿化用地2340平方米、土壤改良702立方米、绿化微地形257.10立方米、清杂130平方米。")
add_para(doc, "园建工程包括生态地铺石人行道面层214.53平方米、生态地铺石人行道356.48平方米，花岗岩路缘石314.66米、花岗岩开孔路缘石74.32米、草阶缘石171.30米，墙面LOGO金属字、斜面LOGO金属字、发光字、宣传牌、垃圾桶、景石等小品设施。水景工程包括土方开挖49.45立方米、回填10.96立方米、余方弃置38.49立方米，碎石垫层14.24立方米、C20垫层14.19立方米、C30P6水池底12.10立方米、C30P6水池壁31.80立方米、钢筋3.512吨、砂浆找平244.62平方米、聚氨酯涂膜防水247.05平方米、水景块料铺设247.05平方米和61.70平方米、不规则阶梯48.58米、不锈钢穿孔板沟盖板及玫瑰不锈钢排水沟24.02平方米、景石安装1块。")
add_para(doc, "给排水工程包括PE给水管DN25至DN100合计470米，UPVC排水管DN150、DN200各12米，快速取水器10个，水平螺翼式水表1台，手动闸阀、球阀、旋转散射喷头及水景潜水泵等。园建清单中另列给排水土建配套内容，包括沟槽土方、砂垫层、雨水口、砖地沟明沟、沟道盖板、排水沟抹灰、车行道路面恢复等。")
add_para(doc, "电气工程包括LED投光灯14套、LED开关电源箱14台、太阳能庭院灯1套、COB灯带260米成套、池底灯7套、CPVC管DN32 170米、CPVC管DN25 270米、镀锌钢管SC110 0.738吨、电力电缆合计320米、接地极15根、接地母线150米，并包含灯基础、砖砌电缆手孔井8座、等电位端子箱基础及箱体、道路开挖恢复等配套土建内容。")

add_heading(doc, "三、图纸摘录", 1)
add_para(doc, "为便于现场交底和施工组织，本方案从施工图PDF中截取主要控制图面作为方案附图。附图仅用于施工组织说明，最终尺寸、标高、材料和节点做法仍以正式施工图、设计变更和现场确认记录为准。")
add_figure(doc, figures["plan"], "图3-1  园建总平面图摘录")
add_figure(doc, figures["water"], "图3-2  水景详图摘录")
add_figure(doc, figures["planting"], "图3-3  绿化种植图摘录")
add_figure(doc, figures["drainage"], "图3-4  给排水图摘录")
add_figure(doc, figures["electric"], "图3-5  电气图摘录")

add_heading(doc, "四、施工部署", 1)
add_para(doc, "项目部实行项目经理负责制，设技术负责人、施工员、质量员、安全员、材料员、资料员、测量员和各专业班组长。施工高峰期按园建土建班组、管道安装班组、电气安装班组、铺装及石材班组、水景专项班组、绿化种植班组、杂工及成品保护班组配置。各班组在项目部统一测量控制、材料计划和安全交底下施工，避免同一区域多工种无序交叉。")
add_para(doc, "现场按施工区段划分为操控中心前场景观区、水景及草阶小品区、道路与人行道铺装区、管线沟槽区、绿化种植区和材料临时堆放加工区。施工前完成原始地形、既有设施、地下管线、乔灌木位置、道路通行边界和临电临水接驳点复核，形成施工交底记录。重要控制点包括水景池体定位标高、排水坡向、路缘石线形、LOGO基础及预埋件位置、灯具布点、给水压力试验、池体闭水试验、苗木土球规格和种植完成面。")
add_para(doc, "施工顺序总体安排为：进场准备及围挡警示、测量放线、保护性拆除和非保护性拆除、地下管线及沟槽施工、水景土方及结构施工、道路基层及路缘石施工、机电预埋与手孔井施工、水景防水及闭水、石材铺装及不锈钢沟盖板安装、LOGO及小品安装、灯具电缆敷设及调试、绿化土壤改良及微地形、乔灌木栽植、地被草皮施工、冲洗保洁、联合调试和竣工验收。")
add_para(doc, "本工程工期较紧，应以水景结构、防水闭水和管线隐蔽为主线组织穿插。水景池体混凝土达到防水基层施工条件前，不安排面层大面积铺贴；给排水和电气沟槽经验收、试压或绝缘测试合格后方可回填及路面恢复；铺装完成区域设硬质围护和通行导向，绿化种植与灯具调试交叉时应保护灯管、电缆井和喷灌设施。")

add_heading(doc, "五、施工准备", 1)
add_para(doc, "技术准备方面，项目部进场后组织图纸会审，重点核对LP、LD、LN、LS、SS、LE等图纸编号与清单范围是否一致，核对水景详图中池体尺寸、泵坑、排水沟、不规则阶梯、角钢预埋件、不锈钢沟盖板、景石及灯具位置；核对给排水管径、埋深、阀门、水表、快速取水器、喷头和潜水泵型号；核对电气管线、灯具、开关电源箱、接地、等电位端子箱和手孔井布置。对广告公司二次深化的LOGO、发光字、宣传牌应在加工前完成尺寸、颜色、字体、固定方式和电源接口确认。")
add_para(doc, "材料准备方面，石材、路缘石、开孔路缘石、草阶缘石、水景石材、不锈钢沟盖板、镀锌钢板、不锈钢发光字、灯具、管材、阀门、潜水泵、喷头、苗木和草皮应按进度分批进场。石材应进行规格、色差、边角、厚度和防污染检查；C30P6混凝土、防水涂料、钢筋、水泥砂浆等关键材料应具备合格证、检测报告或复验资料；苗木应选冠形完整、土球饱满、无病虫害、规格符合清单要求的苗源。")
add_para(doc, "现场准备方面，施工区域设置围挡、警示标志、夜间警示灯和临时通道。靠近厂区道路、操控中心出入口、既有管线和生产设施的部位，应提前与业主或厂区管理方确认作业时间、车辆路线、吊装点、土方外运路线、临时停水停电条件以及动火审批要求。雨季施工期间准备排水泵、彩条布、砂袋、集水坑和防滑措施，水景基坑、沟槽及材料堆放区不得积水。")
add_para(doc, "施工机械和器具包括小型挖机、装载机、翻斗车、切割机、振动夯、平板振动器、砂浆搅拌机、钢筋切断弯曲设备、电焊机、热熔机、绝缘摇表、接地电阻测试仪、压力试验泵、水准仪、全站仪或经纬仪、吊装带、绿化支撑材料、洒水车及高压冲洗设备。所有机具进场前应检查状态，电动工具配备漏电保护器，焊机执行一机一闸一箱一保护。")

add_heading(doc, "六、主要施工方法", 1)
add_heading(doc, "1. 拆除及保护工程", 2)
add_para(doc, "拆除范围包括车行道、彩色沥青路面、现状操控中心墙面LOGO、宣传牌保护性拆除、旧路缘石、部分乔灌木和竹子等。拆除前按图纸Z-01及现场定位标识拆除边线，采用切割先行、分块破除、分类堆放、及时清运的方式施工。靠近保留道路、建筑立面、管线井和绿化保留区域的拆除，应采用人工配合小型机械，严禁盲目大面积破碎。")
add_para(doc, "保护性拆除的宣传牌、可利用设施和业主要求保留的材料，拆除前拍照编号，拆除后集中堆放并采取防雨、防碰撞措施。非保护性拆除产生的建筑垃圾及时装车外运，车辆出场前冲洗轮胎，运输过程中覆盖严密。拆除完成后复核基层标高、管线位置和边界尺寸，作为后续沟槽、基层和铺装施工依据。")
add_heading(doc, "2. 测量放线", 2)
add_para(doc, "测量放线以施工图总平面、园建平面图、给排水图、电气图和现场业主移交控制点为依据。项目部设置临时水准点和轴线控制桩，对人行道边线、路缘石线、草阶、水景池体、排水沟、手孔井、灯具基础、快速取水器、喷头、乔木种植点等进行综合放样。放线后由施工员、质量员和专业班组共同复核，涉及景观效果的弧线、草阶、景石和LOGO位置应邀请业主或设计代表确认。")
add_para(doc, "水景工程对标高敏感，需重点控制池底、池壁顶标高、泵坑、排水沟、溢水及排水方向、石材完成面和不锈钢沟盖板完成面。铺装区域应依据道路和排水方向设置坡度，确保雨水口、开孔路缘石和排水沟形成顺畅排水系统，不得出现倒坡、积水和局部沉陷。")
add_heading(doc, "3. 土方、沟槽及基层施工", 2)
add_para(doc, "土方施工包括水景基坑、给排水管沟、电气管沟、井坑、灯基础、道路恢复和绿化微地形。开挖前先探明地下管线，必要时采用人工探沟确认。沟槽开挖按设计宽度和埋深控制，槽底不得扰动，局部超挖采用级配碎石或混凝土处理。水景基坑深度约0.8米，应设置排水沟和集水坑，雨天停止基底扰动性作业，已开挖基坑及时覆盖或排水。")
add_para(doc, "回填前隐蔽工程必须验收合格。给水管道试压、排水管道灌水或通水、沟槽砂垫层、管线保护层、电缆管敷设、接地测试等完成后方可分层回填。回填土应控制含水率，分层厚度不宜过厚，小型夯实机械分层夯实，路面及铺装基层范围压实度按设计和规范执行；绿化种植范围不得使用建筑垃圾、砖块、沥青块和污染土回填。")
add_heading(doc, "4. 水景工程专项施工", 2)
add_para(doc, "水景工程是本项目的质量和观感控制重点。施工流程为测量放线、基坑开挖、基底验槽、碎石垫层、C20混凝土垫层、钢筋绑扎及预埋件安装、模板支设、C30P6池底及池壁混凝土浇筑、养护、基层修补、砂浆找平、聚氨酯涂膜防水、闭水试验、保护层及块料铺设、不规则阶梯安装、不锈钢沟盖板与排水沟安装、景石吊装、给排水和灯具联合调试。")
add_para(doc, "水池钢筋按图纸LD-1.3及清单要求采用Φ12@200等规格，钢筋保护层垫块布置均匀，池壁钢筋与底板钢筋连接牢固，预埋套管、角钢、泵坑、排水沟和灯具线管位置一次复核。模板应有足够刚度和稳定性，池壁模板接缝严密，防止跑模、胀模和漏浆。C30P6混凝土连续浇筑，振捣密实，施工缝留设、止水处理和养护必须符合防水混凝土要求。")
add_para(doc, "混凝土达到基层施工条件后，对蜂窝、麻面、螺杆孔、阴阳角进行修补，阴阳角做圆弧或钝角处理，20厚1:3水泥砂浆找平层应平整、坚实、无空鼓。聚氨酯涂膜防水分遍成膜，厚度和搭接满足设计要求，管根、阴阳角、泵坑、排水沟、灯具出线口和预埋件周边先做附加层。防水完成后进行闭水试验，闭水时间和水位标识按规范及业主要求执行，检查无渗漏后方可进行石材铺贴。")
add_para(doc, "水景块料铺设应先做样板，确认石材颜色、缝宽、排版、坡度和收边。600×600×30厚块料与100×100×100厚芝麻灰小料石分区施工，不规则阶梯采用弧形整石定制，安装前核对编号和加工尺寸。铺贴砂浆饱满，面层平整，接缝顺直，边角处不得出现小于合理宽度的碎拼。石材表面随铺随清，防止水泥浆污染。")
add_para(doc, "景石A规格约1.0×1.0×2.0米，铁矿石一面切割成平面，饰面10mm热轧钢板并设置镂空字，吊装前编制吊装交底，检查吊点、吊带、地基承载和周边障碍物。景石就位后校正朝向和标高，嵌缝稳固，严禁硬碰硬冲击已完成石材和防水构造。不锈钢穿孔板沟盖板及玫瑰不锈钢排水沟安装应保证线形顺直、拼缝紧密、可开启检修、排水通畅。")
add_heading(doc, "5. 园建铺装及小品施工", 2)
add_para(doc, "生态地铺石人行道施工按基层处理、碎石垫层、C20混凝土基层、30厚水泥砂浆结合层、600×300×18厚生态地铺石铺设、勾缝清理、成品保护流程进行。基层压实和平整度是铺装质量基础，混凝土基层浇筑后按坡度和标高找平，铺贴前对石材进行选色、预排和湿润处理。铺贴时拉线控制标高和缝宽，边部切割整齐，与路缘石、雨水口、井盖、建筑边界交接处顺接。")
add_para(doc, "花岗岩路缘石及开孔路缘石按线形安砌，600×150×300高芝麻灰花岗岩安装前复核基础宽度、C20混凝土垫层和排水开孔方向。路缘石背后靠背混凝土密实，接口平顺，顶面标高一致。开孔路缘石应与排水沟、雨水口和道路坡向协调，安装完成后进行通水检查。草阶缘石600×300×200高，线形和高差控制需兼顾绿化微地形和坐凳式景观效果。")
add_para(doc, "LOGO金属字、斜面LOGO金属字、发光字和宣传牌安装前完成深化确认。墙面LOGO基层为热轧钢板，字体为镀锌钢板，固定方式采用焊接，面饰黄漆；斜面LOGO金属字采用不锈钢材质，面饰哑光白漆；发光字采用不锈钢基层及亚克力发光面，与电气电源接口同步检查。焊接和打磨完成后进行防锈、防腐和面漆施工，字体排列、水平线、间距和固定强度应一次验收。")
add_para(doc, "景石安装、垃圾桶安置及宣传栏利旧安装均应服从整体景观轴线和通行净距。小品设施安装前完成基础隐蔽验收，安装后检查垂直度、稳定性、边角安全和表面观感。利旧宣传栏拆装过程中不得损坏面板和支架，重新安装后补刷防护涂料。")
add_heading(doc, "6. 给排水工程施工", 2)
add_para(doc, "给水管道采用PE管，规格包含DN25、DN32、DN40、DN50、DN65、DN80、DN100，室外埋地深度按图纸及清单要求控制，连接形式为热熔。管材、管件进场应检查壁厚、外观、合格证和规格标识。热熔连接前清洁端面，控制加热温度、时间和插入深度，接口冷却期间不得扭动。管道敷设完成后进行压力试验、冲洗和隐蔽验收，合格后方可回填。")
add_para(doc, "UPVC排水管DN150、DN200埋地深度约1米，施工中控制坡度、接口密封和基础垫层。排水管道安装后进行通水、灌水或闭水检查，确保无渗漏、无倒坡、无堵塞。快速取水器、旋转散射喷头、水表、闸阀、球阀和潜水泵安装位置应便于检修，阀门启闭灵活，方向正确，水景潜水泵与电气防水电源、等电位和漏电保护配合调试。")
add_para(doc, "园建配套给排水土建包括雨水口8座、砖地沟明沟95米、沟道盖板95米、排水沟内外抹灰83.60平方米、砖地沟加高70米、砂垫层、满铺砂卵石、道路开挖及恢复等。雨水口和排水沟应按设计坡向设置，抹灰密实平整，盖板安装稳定、便于开启，沟内不得遗留砂浆和杂物。道路恢复按原结构或清单要求分层施工，恢复后与既有路面顺接。")
add_heading(doc, "7. 电气工程施工", 2)
add_para(doc, "电气工程按管线预埋、手孔井及基础、接地系统、电缆敷设、灯具及电源箱安装、绝缘测试、通电调试流程施工。CPVC管DN32、DN25和镀锌钢管SC110敷设前复核路由、埋深和交叉管线，管口封堵，弯曲半径符合穿线要求。电缆手孔井8座、开关电源及等电位端子箱等土建设施应与园建完成面协调，井盖标高与铺装或绿化面一致。")
add_para(doc, "LED投光灯14套、开关电源箱14台、太阳能庭院灯1套、COB灯带260米、池底灯7套安装前检查型号、色温、防护等级和外观。投光灯和庭院灯基础螺栓位置准确，灯具安装牢固，角度按景观效果调整；灯带安装线形顺直，固定可靠，电源容量与分路匹配；池底灯必须满足IP68防护要求，出线密封、接头防水、等电位连接和漏电保护必须验收合格。")
add_para(doc, "电缆敷设前用绝缘摇表测试电缆绝缘，敷设过程中不得硬拉、扭曲和损伤外皮，电缆两端挂牌编号。接地极15根、接地母线150米施工后进行接地电阻测试，达不到要求时增设接地措施。通电调试按单回路、单设备、分区域、整体联动顺序进行，记录电压、电流、绝缘、接地、灯具启闭、控制箱运行和水景设备联动情况。")
add_heading(doc, "8. 绿化工程施工", 2)
add_para(doc, "绿化施工前完成清杂、场地整形、土壤改良和微地形塑造。整理绿化用地2340平方米、土壤改良702立方米、绿化微地形257.10立方米是绿化质量基础，应清除建筑垃圾、石块、沥青块和污染土，按设计标高和坡度塑造自然地形，与草阶、道路、水景和排水口顺接。种植土应疏松、透水、肥力适宜，必要时掺入有机肥、营养土和改良材料。")
add_para(doc, "现状树木移除和竹子清理应先确认保留与移除范围。苏铁、桂花、圆柏球等移栽苗木起挖时保持土球完整，草绳包扎，运输中防晒防风，栽植后设置支撑、浇透定根水、草绳绕干、刷白和修剪。新植乔木如点景朴树、朴树、乌桕、香橼、樱花、桂花、杨梅、紫薇、红梅、红枫等，应按“先大乔木、后小乔木和球类、再灌木地被、最后草皮”的顺序施工，避免后续吊装和运输破坏已完成地被。")
add_para(doc, "乔木栽植穴尺寸应大于土球，穴底回填种植土并夯实适度，苗木朝向和观赏面按景观效果调整。栽植后分层回土、踏实、围堰、浇透水，支撑牢固美观，支撑点加软垫防止损伤树皮。球类和灌木栽植要求冠形饱满、密度均匀、不露土；法国冬青、南天竹、小叶栀子、金森女贞、春鹃、龟甲冬青、金丝桃、大叶栀子、黄金菊、海桐、花叶络石、常夏石竹、红继木、麦冬等地被按株行距满栽，边线顺直，色块交接清晰。")
add_para(doc, "草皮铺设1400平方米，施工前整平细耙，清除杂物，铺设时错缝紧密，边缘切齐，铺后滚压并浇足水。绿化养护期按清单要求执行绿化二级养护1年，其中成活养护3个月、保存养护9个月。养护内容包括浇水、排涝、扶正、修剪、除草、施肥、病虫害防治、补植和冬季防护。")

add_heading(doc, "七、进度计划", 1)
add_para(doc, "本工程计划2026年5月18日开工，2026年6月30日完工，总工期44日历天。进度安排采用平行流水与专业穿插相结合方式，水景结构及防水闭水、地下管线隐蔽、铺装面层和绿化种植为关键线路。若遇连续降雨、厂区生产协调限制或材料深化确认延迟，项目部应通过增加作业面、延长日作业时间、提前加工石材和小品、分区验收等方式追回工期。")
schedule = [
    ["序号", "施工阶段", "计划时间", "主要工作内容"],
    ["1", "进场准备、图纸会审、测量放线", "5月18日-5月20日", "项目部进场、围挡警示、临水临电、控制点复核、施工交底、材料计划确认"],
    ["2", "拆除及清杂", "5月20日-5月24日", "路面、旧LOGO、路缘石、竹子及需移除树木清理，宣传牌保护性拆除，垃圾外运"],
    ["3", "地下管线及沟槽施工", "5月23日-6月2日", "给排水、电气沟槽、管道砂垫层、管线敷设、手孔井及雨水口、试压和隐蔽验收"],
    ["4", "水景结构施工", "5月25日-6月8日", "水景土方、垫层、钢筋、模板、C30P6池底池壁、预埋件、养护"],
    ["5", "道路基层、路缘石及草阶基础", "5月28日-6月12日", "人行道基层、路缘石、草阶缘石、LOGO和宣传牌基础、道路恢复基层"],
    ["6", "水景防水、闭水及饰面", "6月9日-6月19日", "找平层、聚氨酯防水、闭水试验、块料铺设、不锈钢沟盖板、景石安装"],
    ["7", "铺装、小品及金属字安装", "6月10日-6月22日", "生态地铺石、路面修复、LOGO金属字、发光字、宣传牌、垃圾桶安装"],
    ["8", "电气设备安装及调试", "6月15日-6月25日", "灯具、电源箱、接地、电缆敷设、绝缘测试、景观照明和水景联动调试"],
    ["9", "绿化种植及草皮铺设", "6月12日-6月27日", "土壤改良、微地形、乔木、球类、灌木地被、草皮、支撑、定根水"],
    ["10", "清理、整改、联合验收", "6月26日-6月30日", "成品保洁、缺陷整改、资料整理、通水通电复查、竣工移交"],
]
add_table(doc, schedule[0], schedule[1:], widths=[1.0, 3.2, 3.0, 8.0])
add_para(doc, "进度控制以每日计划和周计划结合执行。每天收工前由施工员统计各专业完成量、材料到场情况、次日作业面和需业主协调事项；每周由项目经理组织进度例会，对水景闭水、管线隐蔽、石材到场、苗木进场、电气调试等关键节点进行预警。绿化种植安排在主要土建湿作业、石材切割和大件吊装完成后集中展开，局部不影响的区域可提前施工并加强保护。")

add_heading(doc, "八、质量保证措施", 1)
add_para(doc, "质量目标为一次验收合格，观感质量满足操控中心景观提升效果要求。项目部建立材料进场验收、样板先行、工序交接、隐蔽验收、实测实量和质量整改闭环制度。关键工序包括水景池体钢筋模板、防水基层、防水闭水、给水试压、排水通水、电气绝缘与接地、石材铺装、路缘石线形、苗木规格和成活养护。")
add_para(doc, "水景质量控制重点是抗渗混凝土、防水节点和闭水试验。池底池壁混凝土不得漏振、欠振或随意留施工缝；找平层不得空鼓开裂；聚氨酯防水应分遍涂刷，管根、阴阳角、泵坑、灯具线管和预埋件周边做附加层；闭水试验合格后办理验收记录。面层铺装不得破坏防水层，沟盖板安装不得刺穿或压坏防水构造。")
add_para(doc, "铺装质量控制重点是基层压实、标高坡度、石材色差、缝宽和平整度。生态地铺石、花岗岩路缘石、水景块料、不规则阶梯等进场后先做外观筛选，破损、缺角、色差明显的材料不得用于显眼部位。铺装完成后及时覆盖保护，禁止砂浆车、钢筋、石材二次搬运直接碾压。")
add_para(doc, "给排水质量控制重点是管材规格、接口质量、埋深、坡度、试压和冲洗。PE管热熔接口应形成均匀翻边，UPVC排水管接口严密，阀门方向正确，快速取水器和喷头标高与绿化完成面协调。沟槽回填前办理隐蔽验收，回填后道路和绿化区域不得沉陷。")
add_para(doc, "电气质量控制重点是防水、防漏电、接地和调试。灯具、开关电源箱、池底灯、电缆、管材等应满足设计防护等级，池底灯和水景设备必须严格进行绝缘、接地和漏电保护检查。接线端子压接牢固，编号清晰，手孔井内无积水和杂物。通电前完成单回路测试，通电后检查色温、亮度、角度和控制逻辑。")
add_para(doc, "绿化质量控制重点是苗木规格、土球、种植土、支撑、浇水和养护。苗木到场验收应对品种、胸径、冠幅、高度、土球、病虫害和树形进行检查，不合格苗木不得栽植。乔木栽植后及时支撑和浇透定根水，地被灌木栽植密度满足设计，草皮铺设后滚压、浇水并封闭养护。")

add_heading(doc, "九、安全文明施工措施", 1)
add_para(doc, "本项目位于厂区操控中心周边，安全管理必须服从厂区统一规定。进场人员进行三级安全教育和专项安全技术交底，佩戴安全帽、反光背心等防护用品。涉及动火、临电、吊装、夜间作业、沟槽开挖、有限空间或靠近生产设施作业时，按厂区审批流程办理作业票，专人监护。")
add_para(doc, "沟槽和基坑施工设置警戒线、硬质围护和夜间警示灯，深度较大或土质松散部位采取放坡或支护措施。雨天加强排水和边坡检查，禁止人员在不稳定沟槽边停留。吊装景石、钢板、灯杆和较大苗木时，起吊范围内严禁站人，吊装指挥信号统一，吊具使用前检查。")
add_para(doc, "临时用电执行三级配电、两级保护和一机一闸一漏一箱，电缆架空或埋地保护，严禁私拉乱接。电焊、切割、打磨作业清理周边易燃物，配备灭火器，作业结束后检查无余火。水景电气和池底灯调试前确认漏电保护和接地有效，严禁带水带电违规操作。")
add_para(doc, "文明施工方面，现场材料按石材、管材、钢筋、灯具、苗木、砂石、水泥、防水材料分类堆放，设置标识，保持通道畅通。切割、破除、土方外运采取洒水降尘和覆盖措施，施工废水不得直接排入绿化区或雨水口，建筑垃圾及时外运。已完成铺装、水景饰面、灯具、苗木和草皮设围护保护，交叉施工人员按指定通道通行。")

add_heading(doc, "十、环境保护与成品保护", 1)
add_para(doc, "环境保护重点控制扬尘、噪声、污水、固废和绿化损伤。破除和切割作业采用湿法降尘，运输车辆覆盖，出入口清扫冲洗。噪声较大的作业尽量避开厂区敏感时段，确需连续作业时提前报备。防水涂料、油漆、稀释剂等材料单独存放，空桶和废弃物分类收集，不得随意倾倒。")
add_para(doc, "成品保护贯穿全过程。管线隐蔽验收后回填夯实并设置标识，防止二次开挖损坏；水景防水完成后设专人看护，未做保护层前不得上人、堆料、拖拉工具；石材铺装完成后覆盖保护，砂浆强度未达到要求前禁止通行；灯具和电源箱安装后防雨防碰撞；苗木栽植后设置支撑和围护，严禁踩踏地被和草皮。")
add_para(doc, "交叉施工中实行工序移交制度。园建向电气、给排水、绿化移交时明确完成面、标高和保护要求；电气和给排水回填后向铺装移交隐蔽验收资料；绿化进场前确认土建湿作业基本结束，避免水泥浆、油漆、焊渣污染植物和土壤。")

add_heading(doc, "十一、雨季及高温施工措施", 1)
add_para(doc, "施工期处于5月中旬至6月底，武汉地区降雨和高温均需重点考虑。雨前检查沟槽、基坑、材料堆场和临时排水系统，水景基坑、管沟和道路基层不得长时间泡水。雨后复核基底承载、回填土含水率、基层压实度和铺装砂浆粘结条件，未达到要求不得强行进入下道工序。")
add_para(doc, "防水、油漆、石材铺贴和电气接线应避开雨天或采取可靠防雨措施。水景闭水试验期间应记录自然降雨影响，试验水位标识清楚。绿化种植避开正午高温，苗木运输和假植采取遮阴保湿，栽后及时浇水、喷雾和支撑。高温时段合理调整作业时间，现场配备饮用水、防暑药品和休息点。")

add_heading(doc, "十二、资料管理与验收", 1)
add_para(doc, "资料员随工程同步收集图纸会审记录、施工技术交底、安全交底、材料合格证、复验报告、隐蔽验收记录、试压记录、闭水记录、绝缘测试记录、接地测试记录、灯具调试记录、苗木进场验收记录、养护记录、分部分项验收记录和竣工影像资料。资料不得后补造假，应与现场实际施工时间、部位和验收结果一致。")
add_para(doc, "分项验收按专业组织。水景工程验收包括结构尺寸、混凝土观感、防水闭水、石材铺装、沟盖板、景石、潜水泵和池底灯；园建验收包括基层、路缘石、铺装、小品、金属字和道路恢复；给排水验收包括管道试压、排水通水、阀门、水表、喷头、雨水口和排水沟；电气验收包括绝缘、接地、灯具、电源箱、手孔井和联动调试；绿化验收包括苗木规格、种植密度、支撑、草皮平整度和养护移交。")
add_para(doc, "竣工前项目部组织自检，自检内容包括工程实体、观感质量、功能试验、资料完整性、现场清洁和成品保护。对发现的积水、空鼓、松动、色差、灯具不亮、阀门渗漏、苗木歪斜、草皮起翘等问题建立整改清单，明确责任人和完成时间。整改完成后申请业主、监理或相关管理单位进行竣工验收。")

add_heading(doc, "十三、劳动力与主要资源计划", 1)
resources = [
    ["类别", "计划配置", "说明"],
    ["管理人员", "项目经理1人、技术负责人1人、施工员2人、质量员1人、安全员1人、材料员1人、资料员1人", "按现场进度动态驻场"],
    ["土建园建班组", "12-18人", "拆除、基层、混凝土、路缘石、铺装、水景结构"],
    ["给排水班组", "4-6人", "PE管热熔、UPVC管、阀门、水表、喷头、潜水泵"],
    ["电气班组", "5-8人", "管线、手孔井配合、电缆、灯具、电源箱、接地调试"],
    ["绿化班组", "12-20人", "苗木移栽、新植、地被、草皮、养护"],
    ["机械设备", "小挖机、夯机、切割机、热熔机、电焊机、试压泵、洒水设备等", "根据作业面分批进退场"],
]
add_table(doc, resources[0], resources[1:], widths=[3.0, 5.5, 6.0])
add_para(doc, "资源投入按关键线路优先保障。5月下旬重点保障拆除、沟槽、水景结构和管线班组；6月上旬重点保障水景结构、道路基层、路缘石和机电预埋；6月中旬重点保障防水、铺装、小品、电气安装和绿化土方；6月下旬重点保障绿化种植、灯具调试、成品保护和整改验收。")

add_heading(doc, "十四、应急与协调措施", 1)
add_para(doc, "项目施工中可能出现连续降雨、地下管线位置不明、材料深化确认滞后、水景闭水发现渗漏、厂区通行限制、苗木到场质量不稳定等风险。项目部建立风险预控机制，对地下管线采用人工探沟和业主确认，对水景防水设置样板和节点旁站，对石材和小品提前下单加工，对苗木提前看苗锁定，对雨季施工提前准备排水和覆盖材料。")
add_para(doc, "若闭水试验发现渗漏，应排水后查找渗漏点，凿除空鼓或薄弱部位，按防水修补工艺处理并重新闭水，未经复验合格不得铺贴面层。若管线试压不合格，应分段排查热熔接口、阀门、管件和堵头，修复后重新试压。若灯具调试异常，应按电源、回路、接头、灯具、控制箱顺序排查，禁止带故障投入使用。")
add_para(doc, "与业主和厂区管理单位协调事项包括施工通道、材料卸车、垃圾外运、临时停水停电、动火审批、夜间施工、成品移交和养护责任边界。项目部指定专人负责沟通，每日形成需协调事项清单，重大事项提前24小时报备，确保施工组织与厂区生产秩序相协调。")

add_heading(doc, "十五、专项技术控制要点", 1)
add_para(doc, "本工程虽为景观提升项目，但施工界面紧邻操控中心，且包含水景、照明、给排水、道路恢复、绿化种植和金属小品，专业交叉密集。项目部应把“节点深化、标高统一、功能先行、观感收口”作为专项技术控制主线。各专业施工前由技术负责人组织一次综合标高复核，将水景池口、铺装完成面、路缘石顶面、草阶完成面、排水沟盖板、雨水口、手孔井、灯基础、绿化种植土完成面统一到同一控制体系内，避免出现井盖高低不平、铺装倒坡、草坪积水、灯具基础外露和水景收边不顺等通病。")
add_para(doc, "水景区域应实行样板引路和节点旁站制度。防水施工前先完成一处阴角、管根、灯具出线口、排水沟转角和预埋件周边的节点样板，经项目部、监理或业主确认后再大面积展开。防水施工期间保持基层干燥洁净，严禁在未固化防水层上堆放石材、砂浆桶和金属构件。闭水试验期间设专人观察池壁外侧、泵坑、排水沟、灯具出线口和施工缝位置，发现湿痕、渗点或水位异常时立即停止下道工序，查明原因后按修补方案处理。")
add_para(doc, "铺装及石材工程的观感质量直接影响操控中心前场形象。施工前应对生态地铺石、水景石材、花岗岩路缘石、开孔路缘石、草阶缘石和不规则阶梯进行编号排版，控制色差、纹理方向、边线弧度和接缝宽度。弧形部位不得简单以碎块拼凑，应优先采用工厂切割或现场精切，保证弧线连续。与建筑墙面、LOGO基础、景石、不锈钢排水沟、雨水口和井盖交接处应进行二次放样，收口材料提前加工，避免现场临时切割造成崩边和尺寸不协调。")
add_para(doc, "绿化种植与园建完成面之间应形成自然过渡。微地形塑造不得遮挡开孔路缘石和排水路径，不得使种植土冲刷进入水景池、排水沟和铺装缝。乔木栽植完成后应从操控中心主视角、道路入口视角和水景观赏视角检查树冠朝向，重点景观树应保证树形舒展、观赏面朝向合理。地被色块边界应与铺装边线、草阶边线和景石位置协调，草皮铺设后及时封闭养护，待根系稳定后再开放通行。")
add_para(doc, "电气和水景给排水联动调试应在土建和防水验收合格后进行。调试前清理手孔井、检查电源箱防水密封、复核接地和等电位连接，确认潜水泵、池底灯、灯带和投光灯的回路编号一致。水景运行试验应检查水泵启停、出水效果、排水回流、池底灯照明、防溅水和漏电保护动作情况。夜景照明调试应结合现场实际观感调整投光角度和灯带亮度，避免眩光直射操控中心窗口、道路行人和车辆驾驶方向。")
add_para(doc, "竣工移交前应组织一次全专业联合巡查，从白天观感、夜间照明、雨后排水、水景运行、绿化成活、道路通行和设施检修七个角度复核。对铺装污染、边角破损、灯具角度偏差、苗木支撑松动、排水沟杂物、井盖异响、水景水位异常等问题逐项整改，整改完成后再进行最终清洗、资料移交和养护交底，确保工程形象完整稳定。")

add_heading(doc, "十六、结束语", 1)
add_para(doc, "本施工方案以热轧厂操控中心景观提升项目的施工图和工程量清单为基础，围绕园林建筑、给排水、电气、绿化及水景工程的实际施工特点进行组织。武汉兴达建筑安装工程有限公司将按计划工期、质量标准和厂区安全文明施工要求组织实施，做到施工有计划、工序有交接、隐蔽有验收、质量有记录、安全有管控、成品有保护，确保项目于2026年6月30日前完成并具备验收移交条件。")

doc.save(OUT)
print(OUT)
