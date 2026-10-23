"""Isolated regressions for the read-only latest-format auditor."""
import hashlib
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pypdf import PdfWriter

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "audit_latest_format.py"
SPEC = importlib.util.spec_from_file_location("audit_latest_format", SCRIPT)
AUDITOR = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(AUDITOR)


def chinese_font(style_or_run, name="宋体"):
    style_or_run.font.name = name
    style_or_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


class LatestFormatTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.docx = self.root / "scheme.docx"
        self.pdf = self.root / "rendered.pdf"
        self.profile = self.root / "profile.json"
        self.spec = {
            "chapters": [{"number": i, "title": f"施工主题{i}"} for i in range(1, 16)],
            "appendices": [{"number": i, "title": f"附表主题{i}", "chapter": 15} for i in range(1, 7)],
        }
        self.spec["appendices"][0]["columns"] = ["序号", "设备名称", "数量"]
        self.spec["appendices"][2].update(column_count=8, first_header="工种", merged_header="按施工阶段投入劳动力情况")
        self.profile.write_text(json.dumps(self.spec, ensure_ascii=False), encoding="utf-8")

    def document(self, *, body=True, inherited=False, tables=True):
        document = Document()
        document.add_paragraph("封面可以保留原有字号和字体", "Title")
        body_style = document.styles.add_style("BiddingBody", WD_STYLE_TYPE.PARAGRAPH)
        target = body_style
        if inherited:
            base = document.styles.add_style("BodyBase", WD_STYLE_TYPE.PARAGRAPH)
            body_style.base_style = base
            target = base
        chinese_font(target)
        target.font.size = Pt(12)
        target.paragraph_format.line_spacing = 1.5
        heading = document.styles.add_style("BiddingHeading1", WD_STYLE_TYPE.PARAGRAPH)
        chinese_font(heading, "黑体")
        heading.font.size = Pt(16)
        appendix = document.styles.add_style("BiddingAppendix", WD_STYLE_TYPE.PARAGRAPH)
        chinese_font(appendix, "黑体")
        appendix.font.size = Pt(14)
        for chapter in self.spec["chapters"]:
            document.add_paragraph(f"第{chapter['number']}章 {chapter['title']}", "BiddingHeading1")
            if body:
                document.add_paragraph("本节施工工艺及质量检查控制要点。", "BiddingBody")
        for entry in self.spec["appendices"]:
            document.add_paragraph(f"附表{entry['number']}：{entry['title']}", "BiddingAppendix")
            if tables and (entry.get("columns") or entry.get("merged_header")):
                table = document.add_table(rows=2, cols=entry.get("column_count", len(entry.get("columns", []))))
                if entry["number"] == 3:
                    table.cell(0, 0).add_paragraph("工种")
                    table.cell(0, 1).merge(table.cell(0, 7)).add_paragraph("按施工阶段投入劳动力情况")
                    for i in range(1, 8):
                        table.cell(1, i).add_paragraph(f"第{i}阶段")
                else:
                    for i, label in enumerate(entry["columns"]):
                        table.cell(0, i).add_paragraph(label)
                # Tables deliberately retain default fonts/spacing; do not flag them.
        return document

    def render_fixture(self, pages=200):
        writer = PdfWriter()
        for _ in range(pages):
            writer.add_blank_page(width=595, height=842)
        writer.write(self.pdf)

    def run_audit(self, document=None, pages=200, **options):
        (document or self.document()).save(self.docx)
        if pages is not None:
            self.render_fixture(pages)
        return AUDITOR.audit(self.docx, pdf=self.pdf if pages is not None else None,
                             profile=self.profile, **options)

    def test_200_pages_and_inherited_body_pass(self):
        report = self.run_audit(self.document(inherited=True))
        self.assertEqual(report["failures"], [])
        self.assertEqual(report["pages"], 200)
        self.assertTrue(report["release_ready"])
        self.assertTrue(report["clean_compliant"])
        self.assertEqual(report["body_paragraphs_checked"], 15)
        self.assertEqual(len(report["appendix_table_checks"]), 2)

    def test_document_defaults_resolve_when_style_omits_properties(self):
        document = self.document()
        style = document.styles["BiddingBody"]
        for tag in ("w:rPr", "w:pPr"):
            node = style.element.find(qn(tag))
            if node is not None:
                style.element.remove(node)
        defaults = document.styles.element.find(qn("w:docDefaults"))
        rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
        for child in list(rpr):
            rpr.remove(child)
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:eastAsia"), "宋体")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "24")
        rpr.extend([fonts, size])
        ppr = defaults.find(qn("w:pPrDefault")).find(qn("w:pPr"))
        for child in list(ppr):
            ppr.remove(child)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        report = self.run_audit(document)
        self.assertEqual(report["failures"], [])

    def test_201_pages_need_review(self):
        report = self.run_audit(pages=201)
        self.assertEqual(report["status"], "needs_review")
        self.assertFalse(report["release_ready"])

    def test_201_page_exception_is_not_clean_compliance(self):
        report = self.run_audit(pages=201, page_limit_exception="用户确认保留必要计算书")
        self.assertEqual(report["status"], "passed_with_exception")
        self.assertTrue(report["release_ready"])
        self.assertFalse(report["clean_compliant"])

    def test_pdf_is_required_for_release(self):
        report = self.run_audit(pages=None)
        self.assertTrue(report["failures"])
        self.assertFalse(report["release_ready"])
        report = self.run_audit(pages=None, allow_unrendered=True)
        self.assertEqual(report["failures"], [])
        self.assertEqual(report["status"], "template_qa_only")
        self.assertFalse(report["release_ready"])

    def test_no_body_fails_even_in_template_mode(self):
        report = self.run_audit(self.document(body=False), pages=None, allow_unrendered=True)
        self.assertTrue(any("No body paragraphs" in f for f in report["failures"]))

    def test_substantive_untagged_body_cannot_evade_audit(self):
        document = self.document()
        document.add_paragraph("未标记的长段施工正文不能因有一段正确正文而绕过格式检查。")
        report = self.run_audit(document)
        self.assertEqual(len(report["uncovered_substantive_paragraphs"]), 1)
        self.assertTrue(any("outside audited styles" in f for f in report["failures"]))

    def test_heading_and_caption_styles_are_excluded(self):
        document = self.document()
        paragraph = document.add_paragraph("这是说明工艺流程图的长图注，无须强制为正文格式。", "Caption")
        paragraph.runs[0].font.size = Pt(10)
        document.add_paragraph("这是较长的小节标题，应保持标题自身的格式。", "Heading 2")
        report = self.run_audit(document)
        self.assertEqual(report["failures"], [])

    def test_appendix_chapter_placement_is_checked(self):
        document = self.document()
        appendix = next(p for p in document.paragraphs if p.text.startswith("附表4"))
        chapter = next(p for p in document.paragraphs if p.text.startswith("第9章"))
        chapter._p.addprevious(appendix._p)
        report = self.run_audit(document)
        self.assertTrue(any("Appendix placement mismatch" in f for f in report["failures"]))

    def test_wrong_direct_font_size_and_line_rule_fail(self):
        document = self.document(inherited=True)
        paragraph = next(p for p in document.paragraphs if p.style.name == "BiddingBody")
        chinese_font(paragraph.runs[0], "黑体")
        paragraph.runs[0].font.size = Pt(11)
        paragraph.paragraph_format.line_spacing = Pt(18)  # not 1.5 multiple
        report = self.run_audit(document)
        self.assertEqual({d["property"] for d in report["format_defects"]},
                         {"east_asian_font", "font_size", "line_spacing"})

    def test_character_style_override_is_checked(self):
        document = self.document()
        character = document.styles.add_style("BadEmphasis", WD_STYLE_TYPE.CHARACTER)
        chinese_font(character, "楷体")
        paragraph = next(p for p in document.paragraphs if p.style.name == "BiddingBody")
        paragraph.runs[0].style = character
        report = self.run_audit(document)
        self.assertTrue(any(d["actual"] == "楷体" for d in report["format_defects"]))

    def test_placeholder_override_always_nonrelease(self):
        document = self.document()
        document.add_paragraph("【填写：工艺】", "BiddingBody")
        report = self.run_audit(document)
        self.assertTrue(any("Unfilled placeholders" in f for f in report["failures"]))
        report = self.run_audit(document, allow_placeholders=True)
        self.assertEqual(report["failures"], [])
        self.assertFalse(report["release_ready"])

    def test_toc_alone_does_not_satisfy_chapter_coverage(self):
        document = self.document()
        toc = document.styles.add_style("TOC1", WD_STYLE_TYPE.PARAGRAPH)
        paragraph = next(p for p in document.paragraphs if p.text.startswith("第2章"))
        paragraph.style = toc
        report = self.run_audit(document)
        self.assertTrue(any("Missing body chapter heading: 施工主题2" in f for f in report["failures"]))

    def test_nested_pageref_does_not_end_toc_exclusion(self):
        document = self.document()
        style_xml = AUDITOR.ET.fromstring(document.styles.element.xml)
        xml = f'''<w:body xmlns:w="{AUDITOR.W[1:-1]}">
          <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> TOC </w:instrText></w:r></w:p>
          <w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGEREF </w:instrText></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
          <w:p><w:r><w:t>仍在目录字段中的条目</w:t></w:r></w:p>
          <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
          <w:p><w:r><w:t>正文开始</w:t></w:r></w:p>
        </w:body>'''
        result = AUDITOR.toc_indices(list(AUDITOR.ET.fromstring(xml)), AUDITOR.Styles(style_xml))
        self.assertEqual(result, {0, 1, 2, 3})

    def test_missing_appendix_table_and_wrong_header_fail(self):
        report = self.run_audit(self.document(tables=False))
        self.assertTrue(any("Missing appendix table" in f for f in report["failures"]))
        document = self.document()
        document.tables[0].cell(0, 1).paragraphs[-1].runs[0].text = "错误列"
        report = self.run_audit(document)
        self.assertTrue(any("table header/schema mismatch" in f for f in report["failures"]))

    def test_labor_merged_header_is_required(self):
        document = self.document()
        document.tables[1].cell(0, 1).paragraphs[-1].runs[0].text = "错误合并表头"
        report = self.run_audit(document)
        labor = report["appendix_table_checks"][1]
        self.assertFalse(labor["merged_header_passed"])
        self.assertTrue(report["failures"])

    def test_explicit_normal_scope_excludes_cover(self):
        document = self.document()
        normal = document.styles["Normal"]
        chinese_font(normal)
        normal.font.size = Pt(12)
        normal.paragraph_format.line_spacing = 1.5
        document.paragraphs[0].style = normal
        document.paragraphs[0].runs[0].font.size = Pt(24)
        for paragraph in document.paragraphs:
            if paragraph.style.name == "BiddingBody":
                paragraph.style = normal
        report = self.run_audit(document, body_styles=["Normal"])
        self.assertEqual(report["failures"], [])
        self.assertEqual(report["body_paragraphs_checked"], 15)

    def test_input_hashes_unchanged_and_cli_rejects_input_overwrite(self):
        self.document().save(self.docx)
        self.render_fixture()
        before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.docx, self.pdf)]
        report = AUDITOR.audit(self.docx, pdf=self.pdf, profile=self.profile)
        self.assertTrue(report["release_ready"])
        after = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.docx, self.pdf)]
        self.assertEqual(before, after)
        result = subprocess.run([sys.executable, str(SCRIPT), str(self.docx), "--profile", str(self.profile),
                                 "--pdf", str(self.pdf), "--output", str(self.docx)], capture_output=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(before[0], hashlib.sha256(self.docx.read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
