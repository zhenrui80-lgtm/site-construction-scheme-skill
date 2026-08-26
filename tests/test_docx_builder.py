import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts import docx_builder as builder


class DocxBuilderTests(unittest.TestCase):
    def test_builder_creates_reopenable_body_and_landscape_section(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scheme.docx"
            document = Document()
            builder.configure_body_styles(document)
            builder.add_heading(document, "一、编制依据及工程概况")
            builder.add_continuous_body_section(document)
            builder.add_body(document, "项目专属正文。")
            builder.add_styled_table(document, "表1  测试", ["项目", "内容"], [("A", "B")], [6, 10])
            builder.add_landscape_section(document)
            document.save(path)
            reopened = Document(path)
            self.assertEqual(len(reopened.tables), 1)
            self.assertEqual(len(reopened.sections), 3)
            self.assertGreater(reopened.sections[-1].page_width, reopened.sections[-1].page_height)

    def test_run_level_replacement_preserves_run_count(self):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("工程名称：")
        paragraph.add_run("旧名称")
        before = len(paragraph.runs)
        builder.replace_paragraph_text_preserve_runs(paragraph, "工程名称：新项目")
        self.assertEqual(len(paragraph.runs), before)
        self.assertEqual(paragraph.text, "工程名称：新项目")


if __name__ == "__main__":
    unittest.main()
